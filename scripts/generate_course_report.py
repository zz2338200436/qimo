from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / ".report_assets"
OUTPUT = ROOT / "智能学习辅助系统-分布式框架技术项目设计文档.docx"
LOGO = ROOT / ".firecrawl" / "template.converted.files" / "image002.png"

FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_KAI = "楷体"
FONT_EN = "Times New Roman"
FONT_CODE = "Consolas"


def font_path(name: str) -> str:
    candidates = {
        "hei": [
            r"C:\Windows\Fonts\simhei.ttf",
            r"C:\Windows\Fonts\msyh.ttc",
        ],
        "song": [
            r"C:\Windows\Fonts\simsun.ttc",
            r"C:\Windows\Fonts\simsunb.ttf",
            r"C:\Windows\Fonts\SimsunExtG.ttf",
        ],
        "kai": [
            r"C:\Windows\Fonts\simkai.ttf",
            r"C:\Windows\Fonts\STKAITI.TTF",
        ],
    }
    for path in candidates[name]:
        if Path(path).exists():
            return path
    return r"C:\Windows\Fonts\simhei.ttf"


def set_run_font(run, cn=FONT_SONG, en=FONT_EN, size: Pt | None = None, bold=None, color=None):
    run.font.name = en
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), cn)
    r_fonts.set(qn("w:ascii"), en)
    r_fonts.set(qn("w:hAnsi"), en)
    r_fonts.set(qn("w:cs"), en)


def set_paragraph_body(paragraph, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(24) if indent else Pt(0)
    for run in paragraph.runs:
        set_run_font(run, FONT_SONG, FONT_EN, Pt(12))


def add_body(doc: Document, text: str, indent=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, FONT_EN, Pt(12))
    set_paragraph_body(p, indent=indent)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, FONT_EN, Pt(10.5), bold=False)
    return p


def add_table_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, FONT_EN, Pt(10.5), bold=False)
    return p


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(10.5)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, FONT_SONG, FONT_EN, size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color="808080", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_matrix_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "8A8A8A")
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Cm(widths_cm[idx])
        set_cell_shading(cell, "D9EAF7")
        set_cell_text(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(10.5))
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].width = Cm(widths_cm[idx])
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(text) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], text, align=align, size=Pt(10.5))
    doc.add_paragraph()
    return table


def add_code_block(doc: Document, title: str, code: str):
    add_table_caption(doc, title)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "B0B0B0")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    cell.text = ""
    for line in code.strip().splitlines():
        p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        set_run_font(run, FONT_CODE, FONT_CODE, Pt(9), bold=False)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    doc.add_paragraph()


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Pt(24)

    h1 = styles["Heading 1"]
    h1.font.name = FONT_HEI
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEI)
    h1._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.first_line_indent = Pt(0)

    h2 = styles["Heading 2"]
    h2.font.name = FONT_HEI
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEI)
    h2._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Pt(0)

    h3 = styles["Heading 3"]
    h3.font.name = FONT_HEI
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEI)
    h3._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(6)
    h3.paragraph_format.first_line_indent = Pt(24)


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(0)
    elif level == 3:
        p.paragraph_format.first_line_indent = Pt(24)
    for run in p.runs:
        set_run_font(run, FONT_HEI, FONT_EN, Pt(16 if level == 1 else 14 if level == 2 else 12), bold=True)
    return p


def add_toc(doc: Document):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("目    录")
    set_run_font(run, FONT_HEI, FONT_EN, Pt(16), bold=True)
    title.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "请在 Word 中更新目录"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_sep)
    r._r.append(fld_text)
    r._r.append(fld_end)


def add_cover(doc: Document):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(12.5))
    doc.add_paragraph()
    title1 = doc.add_paragraph()
    title1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title1.add_run("《分布式框架技术》")
    set_run_font(r, FONT_HEI, FONT_EN, Pt(22), bold=True)
    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title2.add_run("项目设计文档")
    set_run_font(r, FONT_HEI, FONT_EN, Pt(22), bold=True)
    doc.add_paragraph()
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "000000", "8")
    cell = table.cell(0, 0)
    set_cell_text(cell, "项目名称：基于 Spring Cloud 的智能学习辅助系统", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(14))
    cell.width = Cm(13.5)
    doc.add_paragraph()

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta, "FFFFFF", "0")
    rows = [
        ("学    院", "计算机学院"),
        ("班    级", "23计算机科学与技术12班"),
        ("组    长", "詹俊英"),
        ("组    员", "詹俊英"),
        ("指导教师", "唐盛平"),
    ]
    for idx, (k, v) in enumerate(rows):
        set_cell_text(meta.rows[idx].cells[0], k, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(12))
        set_cell_text(meta.rows[idx].cells[1], v, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(12))
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = date.add_run("二〇二六 年 六 月")
    set_run_font(r, FONT_SONG, FONT_EN, Pt(12))
    doc.add_page_break()


def wrap_cn(text: str, max_chars: int) -> list[str]:
    lines: list[str] = []
    for segment in text.split("\n"):
        if len(segment) <= max_chars:
            lines.append(segment)
        else:
            lines.extend([segment[i:i + max_chars] for i in range(0, len(segment), max_chars)])
    return lines


def draw_centered(draw: ImageDraw.ImageDraw, box, lines: list[str], font, fill=(35, 35, 35), line_gap=6):
    x1, y1, x2, y2 = box
    heights = []
    widths = []
    for line in lines:
        b = draw.textbbox((0, 0), line, font=font)
        widths.append(b[2] - b[0])
        heights.append(b[3] - b[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + line_gap


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(70, 70, 70), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    import math

    angle = math.atan2(y2 - y1, x2 - x1)
    size = 14
    p1 = (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6))
    p2 = (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, p1, p2], fill=fill)


def box(draw: ImageDraw.ImageDraw, xy, title, subtitle="", fill="#FFFFFF", outline="#2F5597"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=4)
    title_font = ImageFont.truetype(font_path("hei"), 27)
    small_font = ImageFont.truetype(font_path("song"), 21)
    lines = [title]
    if subtitle:
        lines.extend(wrap_cn(subtitle, 14))
    draw_centered(draw, xy, lines, title_font if not subtitle else small_font, fill=(30, 30, 30), line_gap=8)


def make_architecture_diagram(path: Path):
    im = Image.new("RGB", (1600, 1120), "white")
    d = ImageDraw.Draw(im)
    title_font = ImageFont.truetype(font_path("hei"), 50)
    label_font = ImageFont.truetype(font_path("hei"), 30)
    box_title = ImageFont.truetype(font_path("hei"), 30)
    box_text = ImageFont.truetype(font_path("song"), 25)
    note_font = ImageFont.truetype(font_path("song"), 24)

    def layer_label(text, y, fill):
        d.rounded_rectangle((55, y, 185, y + 66), radius=16, fill=fill, outline="#2F5597", width=3)
        draw_centered(d, (55, y, 185, y + 66), [text], label_font, fill=(20, 40, 70))

    def arch_box(xy, title, subtitle="", fill="#FFFFFF", outline="#2F5597"):
        x1, y1, x2, y2 = xy
        d.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=4)
        if subtitle:
            draw_centered(d, (x1 + 15, y1 + 14, x2 - 15, y1 + 58), wrap_cn(title, 14), box_title, fill=(20, 40, 70), line_gap=4)
            draw_centered(d, (x1 + 18, y1 + 62, x2 - 18, y2 - 12), wrap_cn(subtitle, 18), box_text, fill=(35, 35, 35), line_gap=7)
        else:
            draw_centered(d, xy, wrap_cn(title, 14), box_title, fill=(20, 40, 70), line_gap=8)

    d.text((60, 38), "智能学习辅助系统微服务总体架构", font=title_font, fill=(20, 40, 70))

    layer_label("访问层", 145, "#EAF4FF")
    arch_box((250, 125, 520, 245), "前端页面", "教师端 / 学生端 / 登录页", "#EAF4FF")
    arch_box((670, 105, 1000, 265), "网关服务", "统一入口、JWT鉴权、路由、限流、熔断", "#E2F0D9")
    arch_box((1130, 105, 1450, 190), "Eureka 注册中心", "服务注册与发现", "#FFF2CC")
    arch_box((1130, 220, 1450, 305), "Redis", "验证码、Token 黑名单、限流计数", "#FCE4D6")

    arrow(d, (520, 185), (670, 185), fill=(50, 80, 130), width=5)
    arrow(d, (1000, 160), (1130, 148), fill=(50, 80, 130), width=4)
    arrow(d, (1000, 225), (1130, 260), fill=(50, 80, 130), width=4)

    d.line((120, 350, 1500, 350), fill=(220, 226, 236), width=4)
    layer_label("服务层", 430, "#F8FBFF")

    services = [
        ("认证服务", "登录 / Token", 250, 390),
        ("用户服务", "用户 / 角色", 560, 390),
        ("课程服务", "课程 / 班级", 870, 390),
        ("作业服务", "作业 / 批改", 1180, 390),
        ("考试服务", "考试 / 成绩", 250, 560),
        ("分析服务", "趋势 / 预警", 560, 560),
        ("通知服务", "通知 / 已读", 870, 560),
        ("AI服务", "题目 / 建议", 1180, 560),
    ]
    for name, sub, x, y in services:
        arch_box((x, y, x + 240, y + 120), name, sub, "#F8FBFF")

    arrow(d, (835, 265), (835, 360), fill=(50, 80, 130), width=5)
    d.rounded_rectangle((720, 345, 950, 382), radius=14, fill="#FFFFFF", outline="#9EADCC", width=2)
    draw_centered(d, (720, 345, 950, 382), ["按路由转发到业务服务"], note_font, fill=(40, 50, 80))
    for x in (370, 680, 990, 1300):
        arrow(d, (835, 382), (x, 390), fill=(50, 80, 130), width=3)

    d.line((120, 745, 1500, 745), fill=(220, 226, 236), width=4)
    layer_label("基础层", 855, "#EFEFEF")
    arch_box((250, 815, 540, 940), "MySQL", "各服务独立 schema\n本地事务保存业务事实", "#EFEFEF")
    arch_box((660, 815, 950, 940), "RabbitMQ", "Outbox 事件投递\n异步分析与通知", "#FDE9D9")
    arch_box((1070, 815, 1450, 940), "监控组件", "指标采集、健康检查、运行仪表盘", "#E4DFEC")

    arrow(d, (835, 680), (395, 815), fill=(105, 105, 105), width=4)
    arrow(d, (1300, 680), (805, 815), fill=(50, 80, 130), width=4)
    arrow(d, (680, 680), (805, 815), fill=(50, 80, 130), width=4)
    arrow(d, (835, 680), (1260, 815), fill=(90, 70, 130), width=4)

    d.text((250, 990), "说明：实线表示同步 API/网关调用，蓝色箭头表示领域事件流，灰色箭头表示数据库持久化。", font=note_font, fill=(70, 70, 70))

    im.save(path)


def make_sequence_diagram(path: Path):
    im = Image.new("RGB", (1250, 900), "white")
    d = ImageDraw.Draw(im)
    title_font = ImageFont.truetype(font_path("hei"), 42)
    small = ImageFont.truetype(font_path("song"), 24)
    tiny = ImageFont.truetype(font_path("song"), 21)
    d.text((45, 35), "作业提交后异步学情分析时序图", font=title_font, fill=(20, 40, 70))
    actors = ["学生页面", "网关", "作业服务", "Outbox表", "RabbitMQ", "分析服务", "通知服务"]
    xs = [80, 255, 455, 640, 800, 985, 1160]
    for x, actor in zip(xs, actors):
        box(d, (x - 68, 125, x + 68, 190), actor, fill="#EAF4FF")
        d.line([(x, 190), (x, 805)], fill=(190, 190, 190), width=3)
    steps = [
        (0, 1, 250, "提交作业"),
        (1, 2, 330, "验签并转发"),
        (2, 3, 410, "本地事务写提交与 Outbox"),
        (3, 4, 490, "发布作业提交事件"),
        (4, 5, 570, "消费作业提交事件"),
        (5, 5, 650, "幂等去重并更新掌握度"),
        (5, 6, 730, "生成预警或通知事件"),
    ]
    for a, b, y, text in steps:
        if a == b:
            d.rounded_rectangle((xs[a] + 22, y - 26, xs[a] + 205, y + 42), radius=12, outline=(95, 95, 95), width=3)
            d.text((xs[a] + 34, y - 15), text, font=tiny, fill=(30, 30, 30))
        else:
            arrow(d, (xs[a] + 55, y), (xs[b] - 55, y), fill=(50, 80, 130), width=4)
            d.text((min(xs[a], xs[b]) + 10, y - 34), text, font=small, fill=(30, 30, 30))
    im.save(path)


def make_uml_diagram(path: Path):
    im = Image.new("RGB", (1250, 1000), "white")
    d = ImageDraw.Draw(im)
    title_font = ImageFont.truetype(font_path("hei"), 42)
    d.text((45, 38), "核心类与模块关系 UML 简图", font=title_font, fill=(20, 40, 70))
    groups = [
        ("Gateway安全链", [
            "JWT认证过滤器",
            "Token校验器",
            "身份头透传过滤器",
            "网关安全配置",
        ], 45, 130, "#EAF4FF"),
        ("认证域", [
            "认证控制器",
            "认证应用服务",
            "JWT签发服务",
            "凭证仓储",
        ], 435, 130, "#E2F0D9"),
        ("作业域", [
            "教师作业控制器",
            "学生作业控制器",
            "作业应用服务",
            "作业仓储",
        ], 825, 130, "#FFF2CC"),
        ("分析域", [
            "分析查询控制器",
            "作业提交事件处理器",
            "分析仓储",
            "掌握度记录",
        ], 235, 560, "#FCE4D6"),
        ("事件基础设施", [
            "Outbox发布任务",
            "Outbox事件仓储",
            "消息发布器",
            "幂等事件处理",
        ], 625, 560, "#F2F2F2"),
        ("AI能力域", [
            "AI控制器",
            "AI生成服务",
            "模型调用接口",
            "本地模型适配",
        ], 860, 560, "#E4DFEC"),
    ]
    for title, items, x, y, fill in groups:
        d.rounded_rectangle((x, y, x + 330, y + 280), radius=18, fill=fill, outline="#2F5597", width=4)
        f_title = ImageFont.truetype(font_path("hei"), 28)
        f_item = ImageFont.truetype(font_path("song"), 22)
        d.text((x + 20, y + 20), title, font=f_title, fill=(20, 40, 70))
        for idx, item in enumerate(items):
            d.rectangle((x + 25, y + 70 + idx * 48, x + 305, y + 108 + idx * 48), fill="white", outline="#AAAAAA", width=2)
            d.text((x + 38, y + 77 + idx * 48), item, font=f_item, fill=(30, 30, 30))
    arrow(d, (375, 270), (435, 270))
    arrow(d, (765, 270), (825, 270))
    arrow(d, (990, 410), (425, 560))
    arrow(d, (990, 410), (735, 560))
    arrow(d, (600, 410), (1010, 560))
    im.save(path)


def add_figure(doc: Document, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(14.6))
    add_caption(doc, caption)


def add_screenshot_figure(doc: Document, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(14.2))
    add_caption(doc, caption)


def chapter(doc: Document, title: str, first=False):
    heading = add_heading(doc, title, 1)
    if not first:
        heading.paragraph_format.page_break_before = True


def build_report():
    ASSET_DIR.mkdir(exist_ok=True)
    arch = ASSET_DIR / "architecture.png"
    seq = ASSET_DIR / "sequence.png"
    uml = ASSET_DIR / "uml.png"
    screen_teacher = ASSET_DIR / "screen_teacher_dashboard.png"
    screen_student = ASSET_DIR / "screen_student_assignments.png"
    screen_ai = ASSET_DIR / "screen_ai_questions.png"
    make_architecture_diagram(arch)
    make_sequence_diagram(seq)
    make_uml_diagram(uml)

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc)

    chapter(doc, "1. 项目概述", first=False)
    add_heading(doc, "1.1. 项目背景与意义", 2)
    add_body(doc, "随着高校课程教学逐渐转向线上线下混合模式，课程、作业、考试、学情统计、预警通知等数据分散在不同业务环节中。传统单体系统在功能扩展时容易出现接口耦合、数据库表责任不清、页面与后端联调困难等问题，难以支撑持续演进。")
    add_body(doc, "本项目以“智能学习辅助系统”为题，围绕教师教学管理与学生学习闭环，综合运用 Spring Boot、Spring Cloud Gateway、Eureka、OpenFeign、RabbitMQ、Redis、MySQL、Resilience4j 等分布式框架技术，将原有教学系统拆分为多个独立业务服务，并通过统一网关提供对外入口。项目既能体现课程要求中的页面设计、系统功能实现和文档编写能力，也能展示服务注册发现、网关路由、鉴权透传、熔断限流、事件驱动、数据所有权等核心知识点。")
    add_body(doc, "系统面向教师、学生两类用户。教师可以管理课程、班级、知识点、作业、考试、预警和通知；学生可以查看课程、提交作业、参加考试、查看成绩、学习统计和通知；平台还提供 AI 题目生成、模拟试卷生成和学习建议生成能力，为教学过程提供智能辅助。")
    add_heading(doc, "1.2. 项目目标", 2)
    add_body(doc, "项目目标是构建一个能够稳定运行、功能完整、结构清晰、便于扩展的分布式学习辅助系统。具体目标包括：第一，完成教师端和学生端主要业务页面，保证页面布局美观、交互完整；第二，完成网关、认证、用户、课程、作业、考试、分析、通知、AI 等服务的独立拆分；第三，通过 Eureka 完成服务注册发现，通过 Gateway 统一鉴权、路由、限流和熔断；第四，通过 RabbitMQ 与 outbox 机制实现作业提交、考试完成、预警通知等跨服务异步协作；第五，形成结构完整、格式统一、图表充分的系统设计说明书。")
    add_heading(doc, "1.3. 项目整体功能简介", 2)
    add_body(doc, "系统功能覆盖教学管理、学习过程、评价反馈和智能辅助四条主线。教学管理侧包括课程 CRUD、班级管理、课程分配、学生加入班级、知识点维护、作业发布、考试发布、批改评分和通知发布。学习过程侧包括学生课程列表、作业提交、考试作答、成绩查询、学习数据统计、知识点掌握度查询和通知已读管理。评价反馈侧包括教师仪表盘、成绩趋势、知识点统计、学生学习画像和学情预警。智能辅助侧包括 AI 生成题目、生成模拟试卷和生成学习建议。")
    add_table_caption(doc, "表1-1 项目功能总览")
    add_matrix_table(
        doc,
        ["角色/模块", "主要功能", "对应服务", "得分支撑点"],
        [
            ["教师端", "课程、班级、知识点、作业、考试、预警、通知、AI 工具", "course / assignment / exam / analysis / notification / ai", "页面功能丰富，覆盖教学主流程"],
            ["学生端", "课程查看、作业提交、考试、成绩、学习统计、通知、个人资料", "course / assignment / exam / analysis / notification / user", "学生学习闭环完整"],
            ["平台能力", "登录鉴权、角色切换、网关路由、限流熔断、TraceId、注册发现", "gateway / auth / registry / common", "体现分布式框架技术运用"],
            ["数据与事件", "多 schema 数据拆分、Flyway 迁移、outbox、幂等消费、RabbitMQ", "common-events / common / analysis", "体现高内聚、低耦合与最终一致性"],
        ],
        [2.6, 4.8, 4.0, 3.6],
    )
    add_heading(doc, "1.4. 选用的讯飞开放平台能力", 2)
    add_body(doc, "模板要求说明可选的开放平台能力。本项目在 AI 服务中设计了统一的 AiModelClient 接口，当前为了保证本地课堂演示稳定，采用 LocalMockAiModelClient 作为无外网、无密钥环境下的本地模型适配实现，能够返回题目、试卷和学习建议的结构化结果。")
    add_body(doc, "若接入讯飞开放平台，可将 AiModelClient 的实现替换为讯飞星火大模型接口，用于“按课程与知识点生成题目”“生成模拟试卷”“根据学情数据生成学习建议”。该设计使业务层只依赖抽象接口，不直接绑定第三方 SDK，符合面向接口编程与低耦合原则。")

    chapter(doc, "2. 需求分析")
    add_heading(doc, "2.1. 功能需求", 2)
    add_body(doc, "系统需求从用户角色出发划分为教师端、学生端和平台支撑三类。教师端强调教学资源和评价过程管理，学生端强调学习任务完成与反馈查看，平台支撑强调认证授权、服务治理、数据持久化、事件一致性和运行监控。")
    add_table_caption(doc, "表2-1 功能需求说明")
    add_matrix_table(
        doc,
        ["需求编号", "需求名称", "需求描述", "优先级"],
        [
            ["R1", "统一登录与角色识别", "支持验证码、用户名密码登录、JWT 签发、教师/学生身份识别和 /api/auth/me 上下文获取。", "高"],
            ["R2", "教师课程与班级管理", "教师可维护课程、班级、专业、课程分配关系，并查看班级学生。", "高"],
            ["R3", "作业与考试管理", "教师可发布作业和考试，学生可提交作业、参加考试，教师可批改评分。", "高"],
            ["R4", "学情分析与预警", "系统根据作业、考试和知识点数据形成成绩趋势、知识点掌握度和早期预警。", "高"],
            ["R5", "通知中心", "教师可发布通知，学生可查看通知、标记已读、统计未读数。", "中"],
            ["R6", "AI 教学辅助", "教师可生成题目和试卷，学生或教师可获取学习建议。", "中"],
            ["R7", "服务治理", "所有业务通过 Gateway 进入，服务注册到 Eureka，并提供熔断、限流、TraceId 和健康检查。", "高"],
        ],
        [1.5, 3.0, 7.0, 1.5],
    )
    add_heading(doc, "2.2. 基础功能清单", 2)
    add_table_caption(doc, "表2-2 基础功能清单")
    add_matrix_table(
        doc,
        ["序号", "基础功能", "完成情况", "说明"],
        [
            ["1", "教师登录与学生登录", "已完成", "auth-service 提供登录、验证码、刷新、登出、身份查询接口。"],
            ["2", "课程列表与课程 CRUD", "已完成", "course-service 承接教师课程和学生课程查询。"],
            ["3", "班级与学生关系维护", "已完成", "支持班级新增、编辑、删除、课程分配和学生加入班级。"],
            ["4", "作业发布、提交与批改", "已完成", "assignment-service 拆分教师端与学生端控制器。"],
            ["5", "考试发布、提交与成绩", "已完成", "exam-service 提供考试 CRUD、学生考试和成绩查询。"],
            ["6", "通知查看与已读", "已完成", "notification-service 管理通知投递与阅读状态。"],
            ["7", "前端页面与浏览器烟测", "已完成", "frontend/dist 包含教师端和学生端页面，并配套 Playwright 运行时验证脚本。"],
        ],
        [1.1, 3.8, 1.7, 7.1],
    )
    add_heading(doc, "2.3. 扩展功能清单", 2)
    add_table_caption(doc, "表2-3 扩展功能清单")
    add_matrix_table(
        doc,
        ["序号", "扩展功能", "技术实现", "价值"],
        [
            ["1", "统一网关鉴权", "Gateway GlobalFilter 验证 JWT，并重写 X-User-Id、X-Roles、X-Active-Role。", "避免业务服务重复解析 Token。"],
            ["2", "服务注册发现", "registry-server 提供 Eureka Server，各业务服务通过 lb://service-name 调用。", "服务地址解耦，便于扩缩容。"],
            ["3", "限流与熔断", "Gateway 路由配置 CircuitBreaker，考试提交和 AI 接口设置差异化限流。", "提升高并发和下游异常时的稳定性。"],
            ["4", "事件驱动分析", "作业和考试事件通过 outbox、RabbitMQ、幂等消费驱动 Analysis 服务更新 read model。", "保证业务写入与分析计算低耦合。"],
            ["5", "AI 教学辅助", "AiModelClient 抽象模型调用，当前本地模拟，后续可替换讯飞星火。", "便于扩展智能题库和学习建议。"],
            ["6", "可观测性", "接入 Actuator、Prometheus、Grafana、TraceId/MDC。", "便于演示、排障和运维。"],
        ],
        [1.1, 3.4, 6.0, 3.2],
    )
    add_heading(doc, "2.4. 运行环境与依赖", 2)
    add_table_caption(doc, "表2-4 运行环境与依赖")
    add_matrix_table(
        doc,
        ["类型", "版本/工具", "用途", "说明"],
        [
            ["JDK", "Java 17", "后端服务运行", "parent-pom 使用 Maven Enforcer 约束 Java 版本。"],
            ["框架", "Spring Boot 3.5.3 / Spring Cloud 2025.0.0", "微服务与分布式治理", "统一由 parent-pom BOM 管理。"],
            ["数据库", "MySQL 8.4", "业务数据持久化", "按 sc_auth、sc_user 等 schema 拆分。"],
            ["缓存", "Redis 7.4", "验证码、Token 黑名单、限流计数", "Gateway 与 Auth 服务使用。"],
            ["消息", "RabbitMQ 3.13", "领域事件投递", "支持 outbox relay 和异步分析。"],
            ["构建", "Maven 3.9+", "多模块构建", "根 pom 聚合 24 个模块。"],
            ["前端", "HTML / CSS / JavaScript / Bootstrap", "教师端和学生端页面", "静态资源通过 5500 或 Spring Boot 访问。"],
            ["部署", "Docker Compose", "本地完整环境", "包含 MySQL、Redis、RabbitMQ、Eureka、Gateway 和业务服务。"],
        ],
        [2.0, 3.6, 4.6, 3.5],
    )

    chapter(doc, "3. 系统总体设计")
    add_heading(doc, "3.1. 设计目标与原则", 2)
    add_body(doc, "总体设计遵循“服务自治、数据归属清晰、接口契约稳定、故障可降级、演示可验证”的原则。服务边界以业务领域划分，避免一个服务直接读写另一个服务的数据库；跨服务读操作通过 API 或 Feign 完成，跨服务写协作优先通过领域事件完成。")
    add_body(doc, "在面向对象设计上，系统将 Controller、Application Service、Repository、DTO、Event Handler 分层，体现封装、接口抽象和职责单一原则。例如 AI 服务通过 AiModelClient 抽象第三方模型能力，分析服务通过 AssignmentSubmittedAnalysisHandler 消费事件并更新 read model，网关通过 JwtAuthenticationFilter 与 HeaderEnrichFilter 分别负责认证和身份头透传。")
    add_heading(doc, "3.2. 系统架构设计", 2)
    add_body(doc, "系统采用“前端页面 + API Gateway + 微服务集合 + 独立数据库 schema + 消息中间件 + 可观测性组件”的架构。所有外部请求先进入 Gateway，Gateway 根据路径和方法转发到对应服务，同时完成 JWT 验证、CORS、限流、熔断和 TraceId 透传。业务服务启动后注册到 Eureka，Gateway 通过 lb://service-name 完成负载均衡调用。")
    add_figure(doc, arch, "图3-1 智能学习辅助系统 Spring Cloud 总体架构")
    doc.add_page_break()
    add_table_caption(doc, "表3-1 服务端口与职责")
    add_matrix_table(
        doc,
        ["模块", "端口", "核心职责", "数据库"],
        [
            ["registry-server", "8761", "Eureka 注册中心", "无"],
            ["gateway", "8080", "统一 API 入口、JWT 鉴权、路由、限流、熔断", "Redis"],
            ["auth-service", "8081", "登录、验证码、Token、角色切换", "sc_auth"],
            ["user-service", "8082", "用户档案、角色、学生/教师资料", "sc_user"],
            ["course-service", "8083", "课程、班级、专业、知识点", "sc_course"],
            ["assignment-service", "8084", "作业、提交、批改、作业知识点", "sc_assignment"],
            ["exam-service", "8085", "考试、试题、提交、成绩", "sc_exam"],
            ["analysis-service", "8086", "学情分析、趋势、预警、掌握度", "sc_analysis"],
            ["notification-service", "8087", "通知、未读数、已读状态", "sc_notification"],
            ["ai-service", "8088", "AI 生成题目、试卷和学习建议", "sc_ai"],
        ],
        [2.9, 1.2, 6.0, 3.3],
    )
    add_heading(doc, "3.3. 模块划分与职责", 2)
    add_heading(doc, "3.3.1. 模块 1：认证与网关模块", 3)
    add_body(doc, "认证与网关模块由 gateway、auth-service、user-service 共同组成。auth-service 负责验证码、登录、Access Token、Refresh Token、登出、角色切换等认证事实；user-service 负责用户资料和角色关系；gateway 负责校验 Token，并将可信身份写入请求头，业务服务只读取 X-User-Id、X-Roles、X-Active-Role。")
    add_heading(doc, "3.3.2. 模块 2：课程与教学资源模块", 3)
    add_body(doc, "课程与教学资源模块由 course-service 承接，负责课程、班级、班级学生关系、课程分配、专业字典和教师知识点目录。作业、考试和分析服务只保存课程、班级、知识点等外部 ID，通过 Course API 读取摘要信息，不直接访问 course-service 的数据库。")
    add_heading(doc, "3.3.3. 模块 3：作业与考试模块", 3)
    add_body(doc, "作业与考试模块分别由 assignment-service 和 exam-service 承接。作业服务负责作业发布、作业提交、批改和作业知识点关联；考试服务负责考试发布、试题、考试提交、答案和成绩。两个服务都面向教师端和学生端提供不同的控制器入口，降低权限判断和接口语义混杂。")
    add_heading(doc, "3.3.4. 模块 4：分析、通知与 AI 模块", 3)
    add_body(doc, "analysis-service 负责成绩趋势、知识点掌握度和学情预警；notification-service 负责通知投递和已读状态；ai-service 负责生成题目、试卷和学习建议。它们与核心业务服务之间通过 API 和事件协作，既能保证功能完整，也能保持业务事实的单一归属。")
    add_heading(doc, "3.3.5. 明确模块间调用关系", 3)
    add_body(doc, "同步调用主要发生在页面读路径、权限校验和下拉数据加载中，例如 Gateway 转发教师课程接口到 course-service，assignment-service 通过 Course API 校验课程知识点。异步调用主要发生在作业提交、考试完成和预警通知中，生产者先写本地业务表和 outbox_event，再由 OutboxRelayJob 投递消息，消费者通过 processed_event 做幂等控制。")

    chapter(doc, "4. 详细设计")
    add_heading(doc, "4.1. 核心类设计", 2)
    add_body(doc, "系统核心类按“入口控制器、应用服务、仓储适配、事件处理、公共基础设施”组织。Controller 只处理 HTTP 参数、身份上下文和响应封装；Service 承接业务规则；Repository 隔离 JPA/JDBC 持久化；Event Handler 处理跨服务事件；common 模块提供异常、响应、Feign、MDC、outbox 和幂等消费等公共能力。")
    add_table_caption(doc, "表4-1 代表性核心类设计")
    add_matrix_table(
        doc,
        ["类/接口", "所属模块", "职责", "设计要点"],
        [
            ["JwtAuthenticationFilter", "gateway", "校验 Bearer Token，解析用户身份。", "实现 GlobalFilter，白名单和 OPTIONS 请求直接放行。"],
            ["HeaderEnrichFilter", "gateway", "清理客户端身份头并写入可信身份头。", "防止客户端伪造 X-User-Id 等内部身份字段。"],
            ["JwtTokenService", "auth-service", "签发和校验 JWT。", "使用 RS256、kid、jti 和过期时间控制。"],
            ["TeacherAssignmentController", "assignment-service", "教师作业发布、编辑、删除和查询。", "教师入口与学生入口分离。"],
            ["AssignmentSubmittedAnalysisHandler", "analysis-service", "消费作业提交事件并更新掌握度。", "结合 IdempotentEventHandler 防重复消费。"],
            ["OutboxRelayJob", "common", "定时发布待投递 outbox 事件。", "失败时记录重试次数和下次重试时间。"],
            ["AiModelClient", "ai-service", "AI 模型调用抽象接口。", "本地实现可替换为讯飞星火等模型实现。"],
        ],
        [4.0, 3.0, 5.0, 3.2],
    )
    add_heading(doc, "4.2. 接口 / API 调用设计", 2)
    add_body(doc, "接口设计统一以 /api 开头，由 Gateway 根据路径和 HTTP 方法转发到后端服务。Gateway 中每条路由都配置 routeId、目标服务、Path Predicate、Method Predicate 和 CircuitBreaker Filter。Context7 查询的 Spring Cloud Gateway 文档说明，Gateway 支持在路由过滤器中配置 CircuitBreaker 和 fallbackUri，也支持按路由配置限流，这与项目 application.yml 中的实现一致。")
    add_table_caption(doc, "表4-2 关键网关路由与接口设计")
    add_matrix_table(
        doc,
        ["路由", "路径", "目标服务", "说明"],
        [
            ["auth-route", "/api/auth/login / refresh / captcha", "auth-service", "公开认证入口。"],
            ["course-route", "/api/teacher/courses/** / api/student/courses/**", "course-service", "教师和学生课程查询、课程管理。"],
            ["assignment-write-route", "/api/teacher/assignments/**", "assignment-service", "作业发布、编辑、删除。"],
            ["student-assignment-submit-route", "/api/student/assignments/*/submit", "assignment-service", "学生作业提交。"],
            ["student-exam-submit-route", "/api/student/exams/*/submit", "exam-service", "学生考试提交，使用更高突发限流和更严格熔断阈值。"],
            ["analysis-route", "/api/teacher/dashboard / api/student/stats 等", "analysis-service", "教师仪表盘、学生统计和预警查询。"],
            ["notification-route", "/api/notifications/**", "notification-service", "通知查询、创建、已读、删除。"],
            ["ai-route", "/api/ai/**", "ai-service", "题目、试卷、学习建议生成，限流更严格。"],
        ],
        [3.4, 5.2, 3.0, 3.6],
    )
    add_code_block(
        doc,
        "代码4-1 Gateway 路由与熔断配置片段",
        """
spring:
  cloud:
    gateway:
      routes:
        - id: student-exam-submit-route
          uri: lb://exam-service
          predicates:
            - Path=/api/student/exams/*/submit
            - Method=POST
          filters:
            - name: CircuitBreaker
              args:
                name: exam-service-student-submit
                fallbackUri: forward:/_fallback/exam-service
""",
    )
    add_heading(doc, "4.3. 数据存储设计", 2)
    add_body(doc, "数据存储采用“每个业务服务一个 schema”的拆分方式，保证每张业务表只有一个写入所有者。服务之间可以保存外部业务 ID，例如 user_id、course_id、knowledge_point_id，但不创建跨 schema 外键，也不允许通过 JDBC/JPA 直接读写其他服务的业务表。")
    add_table_caption(doc, "表4-3 数据所有权设计")
    add_matrix_table(
        doc,
        ["服务", "Schema", "主要表/数据", "跨服务共享方式"],
        [
            ["auth-service", "sc_auth", "auth_credentials、验证码和 Token 状态", "Gateway 通过认证接口和 JWT 获取身份。"],
            ["user-service", "sc_user", "users、roles、user_roles", "其他服务通过 User API 读取用户摘要。"],
            ["course-service", "sc_course", "courses、course_classes、class_students、teacher_knowledge_points", "作业、考试和分析保存课程/班级/知识点 ID。"],
            ["assignment-service", "sc_assignment", "assignments、assignment_submissions、assignment_knowledge_points", "通过 Assignment 事件驱动分析。"],
            ["exam-service", "sc_exam", "exams、questions、exam_submissions、exam_answers", "通过 ExamFinishedEvent 驱动分析和通知。"],
            ["analysis-service", "sc_analysis", "kp_mastery、score_trends、early_warnings", "面向页面提供 read model 查询。"],
            ["notification-service", "sc_notification", "notifications", "只保存通知事实和阅读状态。"],
            ["ai-service", "sc_ai", "ai_generations", "记录 AI 请求与生成结果。"],
        ],
        [3.0, 2.4, 5.2, 4.6],
    )
    add_heading(doc, "4.4. 异常处理设计", 2)
    add_body(doc, "异常处理分为网关层、服务层和前端层。网关层对缺少 Token、Token 过期、黑名单命中等情况统一返回未授权响应；服务层通过公共异常基类和 GlobalExceptionHandler 统一封装业务异常、资源不存在、权限不足和远端调用失败；前端页面根据 HTTP 状态展示登录过期、无权限、数据加载失败或可重试提示。")
    add_table_caption(doc, "表4-4 异常与降级策略")
    add_matrix_table(
        doc,
        ["异常场景", "处理位置", "处理方式", "效果"],
        [
            ["未携带 Bearer Token", "Gateway", "JwtAuthenticationFilter 返回 401。", "阻止未登录请求进入业务服务。"],
            ["下游服务不可用", "Gateway", "CircuitBreaker 转发到 fallback。", "返回统一错误结构，避免页面长时间等待。"],
            ["跨服务事件重复投递", "Analysis 等消费者", "processed_event 幂等表去重。", "保证 read model 不重复累加。"],
            ["消息发布失败", "OutboxRelayJob", "记录失败原因并按退避时间重试。", "提高事件最终投递成功率。"],
            ["前端接口失败", "页面 JS", "显示真实错误态或空态，并保留刷新/重试入口。", "提升演示和使用体验。"],
        ],
        [3.0, 3.2, 5.0, 3.8],
    )

    chapter(doc, "5. UML 模型图")
    add_heading(doc, "5.1. UML 类图", 2)
    add_body(doc, "图5-1 以核心类关系展示系统中网关安全链、认证域、作业域、分析域、事件基础设施和 AI 能力域的协作关系。图中每个方框代表一个类或接口，箭头表示调用、依赖或事件流转关系。")
    add_figure(doc, uml, "图5-1 核心类与模块关系 UML 简图")
    add_heading(doc, "5.2. 关键业务时序图", 2)
    add_body(doc, "作业提交后的学情分析流程是本项目最能体现分布式框架技术的链路之一。学生提交作业时，assignment-service 在同一个本地事务中写入提交记录和 outbox_event；OutboxRelayJob 定时发布事件到 RabbitMQ；analysis-service 消费事件后通过幂等处理器更新知识点掌握度，并可继续触发预警和通知。")
    add_figure(doc, seq, "图5-2 作业提交后异步学情分析时序图")

    chapter(doc, "6. 关键技术与实现说明")
    add_body(doc, "本项目使用的关键技术覆盖课程中的分布式框架核心知识点。Spring Boot 用于构建每个独立服务；Spring Cloud Gateway 用于统一入口和路由；Eureka 用于服务注册发现；OpenFeign 用于声明式服务间调用；Resilience4j 用于熔断；Redis 用于验证码、黑名单和限流计数；RabbitMQ 用于异步事件；Flyway 用于数据库迁移；Docker Compose 用于本地多服务编排；Actuator、Prometheus 和 Grafana 用于可观测性。")
    add_table_caption(doc, "表6-1 关键技术应用说明")
    add_matrix_table(
        doc,
        ["技术点", "项目应用", "对应代码/配置", "课程知识点体现"],
        [
            ["Spring Cloud Gateway", "统一 /api 入口、路由转发、CORS、限流熔断。", "gateway/application.yml", "网关、过滤器、路由谓词。"],
            ["Eureka", "所有业务服务注册到 registry-server。", "registry-server、各服务 application.yml", "服务注册与发现。"],
            ["OpenFeign", "跨服务读取用户、课程、作业知识点等摘要数据。", "*-service-api 模块", "声明式远程调用。"],
            ["JWT + RSA", "Access Token 使用 RS256，Gateway 统一验签。", "auth-service、gateway filter", "统一认证与授权。"],
            ["RabbitMQ + Outbox", "本地事务写 outbox，定时 relay 发布事件。", "common/event/outbox", "异步消息与最终一致性。"],
            ["Flyway", "每个服务独立维护迁移脚本。", "src/main/resources/db/migration", "数据库版本治理。"],
            ["Resilience4j", "考试提交和 AI 接口差异化熔断限流。", "gateway/application.yml", "容错与稳定性设计。"],
            ["Playwright/脚本烟测", "页面级和接口级验证教师、学生链路。", "scripts/verify-*.js", "测试与演示支撑。"],
        ],
        [3.2, 4.5, 4.2, 3.1],
    )
    add_code_block(
        doc,
        "代码6-1 Outbox 事件发布核心逻辑摘要",
        """
@Scheduled(fixedDelayString = "${platform.outbox.relay.fixed-delay:5000}")
public void run() {
    List<OutboxEventEntity> events = repository.findDuePendingEvents(batchSize, now);
    for (OutboxEventEntity event : events) {
        if (publisher.publish(OutboxEventMessage.from(event, bindingName))) {
            repository.markPublished(event.id(), now);
        } else {
            repository.markFailed(event.id(), retryCount, nextRetryAt, reason, maxRetries);
        }
    }
}
""",
    )
    add_body(doc, "该实现将“业务写入”和“消息投递”拆成两个阶段。业务服务只负责在本地事务中写业务表和 outbox 表，消息投递由定时任务异步完成；如果消息中间件短时不可用，事件仍保存在数据库中，可按退避策略重试，从而避免业务提交成功但事件丢失。")
    add_body(doc, "前端页面方面，系统采用教师端和学生端分开的页面结构。教师端页面包括 teacher-dashboard、teacher-courses、teacher-assignments、teacher-knowledge、teacher-warning、teacher-student-dashboard、teacher-notifications、teacher-settings、teacher-ai-tools 等；学生端页面包括 student-dashboard、student-courses、student-assignments、student-notifications、student-settings、student-stats、student-ai-assistant 等。页面通过 api.js 与 Gateway 通信，运行时脚本验证了核心页面在 JWT 环境下可以正常加载和操作。")
    add_table_caption(doc, "表6-2 分布式框架技术亮点映射")
    add_matrix_table(
        doc,
        ["课程知识点", "项目实现亮点", "证明材料", "评分价值"],
        [
            ["服务注册发现", "registry-server 作为 Eureka 注册中心，Gateway 和业务服务以服务名注册。", "docker-compose.yml、各服务 application.yml", "体现微服务治理基础。"],
            ["统一网关", "Gateway 统一 /api 入口，承担鉴权、路由、CORS、限流、熔断和 TraceId。", "gateway 路由配置、JwtAuthenticationFilter", "体现分布式入口与安全控制。"],
            ["远程调用", "各 *-service-api 模块定义 DTO 与 Feign 契约，业务服务通过接口读取摘要。", "course-service-api、user-service-api 等", "体现服务间接口解耦。"],
            ["消息驱动", "作业提交和考试完成事件通过 outbox 与 RabbitMQ 异步投递。", "common-events、OutboxRelayJob", "体现最终一致性和异步协作。"],
            ["容错稳定", "考试提交和 AI 路由配置差异化限流与熔断 fallback。", "gateway/application.yml", "体现高并发和故障降级意识。"],
            ["数据自治", "按 sc_auth、sc_course、sc_assignment 等 schema 划分数据所有权。", "docs/data-ownership.md、Flyway 脚本", "体现服务边界清晰。"],
        ],
        [2.6, 5.0, 4.2, 3.0],
    )

    chapter(doc, "7. 运行说明")
    add_heading(doc, "7.1. 运行步骤", 2)
    add_body(doc, "本地运行建议先启动基础设施，再构建后端服务，最后启动前端预览和烟测脚本。以下步骤适合 Windows PowerShell 环境，端口保持项目默认配置。")
    add_table_caption(doc, "表7-1 系统运行步骤")
    add_matrix_table(
        doc,
        ["步骤", "命令/操作", "说明"],
        [
            ["1", "docker compose -f docker-compose.dev.yml up -d", "启动 MySQL、Redis、RabbitMQ 等基础设施。"],
            ["2", "mvn -T 1 -DskipTests package", "构建根聚合工程和各业务服务。"],
            ["3", "powershell -ExecutionPolicy Bypass -File scripts/start-runtime-smoke-stack.ps1", "启动 registry、gateway 和业务服务运行时烟测栈。"],
            ["4", "powershell -ExecutionPolicy Bypass -File scripts/get-dev-auth-session.ps1 -Role teacher", "获取教师端 JWT 会话。"],
            ["5", "powershell -ExecutionPolicy Bypass -File scripts/get-dev-auth-session.ps1 -Role student", "获取学生端 JWT 会话。"],
            ["6", "python scripts/frontend_dev_server.py", "启动前端 5500 预览服务，页面请求转发到 Gateway。"],
            ["7", "node scripts/verify-gateway-api-smoke.js", "运行统一网关 API 烟测。"],
            ["8", "node scripts/verify-teacher-jwt-pages.js / node scripts/verify-student-jwt-pages.js", "运行教师端和学生端页面烟测。"],
        ],
        [1.2, 8.8, 4.2],
    )
    add_heading(doc, "7.2. 运行与测试结果", 2)
    add_body(doc, "项目文档记录的最终交付状态显示，Eureka 注册表中 AI、Analysis、Assignment、Auth、Course、Exam、Gateway、Notification、User 等 9 个应用均为 UP；Gateway 统一接口烟测覆盖学生和教师链路，包含认证、课程、作业、考试、分析、通知和 AI；页面级烟测覆盖教师端 9 个受保护页面与学生端高频页面。")
    add_table_caption(doc, "表7-2 验证与演示覆盖")
    add_matrix_table(
        doc,
        ["验证类型", "覆盖范围", "结果", "说明"],
        [
            ["服务注册", "registry-server、gateway、8 个业务服务", "通过", "Eureka 中核心服务状态为 UP。"],
            ["API 烟测", "登录、课程、作业、考试、分析、通知、AI", "通过", "通过 Gateway 调用，不依赖旧单体兜底。"],
            ["教师页面", "dashboard、courses、assignments、knowledge、warning 等", "通过", "验证 JWT 页面链和浏览器 CRUD。"],
            ["学生页面", "dashboard、courses、assignments、notifications、settings、stats", "通过", "验证学生学习闭环。"],
            ["数据迁移", "各服务 Flyway V1/V2/V3/V21 脚本", "通过", "多 schema 初始化与迁移脚本已纳入服务资源目录。"],
        ],
        [2.4, 5.2, 1.8, 4.8],
    )
    add_table_caption(doc, "表7-3 关键烟测用例与通过标准")
    add_matrix_table(
        doc,
        ["用例编号", "验证场景", "通过标准", "对应评分项"],
        [
            ["T1", "教师登录后进入仪表盘", "统计卡片、图表、最近活动可正常渲染。", "页面设计、课堂演示"],
            ["T2", "学生查看作业列表", "作业状态、课程、截止时间、提交入口均可见。", "系统功能实现"],
            ["T3", "学生提交或重新提交作业", "提交接口返回成功后列表刷新，页面显示成功提示。", "系统功能实现"],
            ["T4", "教师调用 AI 题目生成", "生成题目、答案、解析和知识点标签。", "扩展功能、课堂演示"],
            ["T5", "Gateway 统一转发 API", "所有页面请求走 /api，由 JWT 与角色上下文控制。", "分布式框架技术"],
            ["T6", "作业提交事件进入分析链路", "outbox_event 被 relay 发布，分析服务幂等更新掌握度。", "系统功能实现"],
        ],
        [1.8, 4.2, 5.2, 3.0],
    )
    add_heading(doc, "7.3. 系统运行界面截图", 2)
    add_body(doc, "为增强课堂演示和报告说服力，报告补充三张运行界面截图。截图覆盖教师端仪表盘、学生端作业提交入口和 AI 题目生成结果，分别对应页面设计、基础功能和扩展功能三个评分维度。")
    add_screenshot_figure(doc, screen_teacher, "图7-1 教师端仪表盘运行界面")
    add_screenshot_figure(doc, screen_student, "图7-2 学生端作业列表与提交入口")
    add_screenshot_figure(doc, screen_ai, "图7-3 教师端 AI 题目生成结果")
    add_body(doc, "演示时建议按“教师发布课程与作业、学生提交作业、教师查看学情分析、系统生成通知或学习建议”的路线进行。该路线覆盖页面设计、功能实现、接口调用、数据库写入、事件驱动和分析读模型，能够充分展示项目工作量和课程知识点。")

    chapter(doc, "8. 项目总结与展望")
    add_heading(doc, "8.1. 项目完成情况", 2)
    add_body(doc, "本项目已经完成从单体教学系统向 Spring Cloud 多服务架构的主要迁移和功能实现。系统具备较完整的教师端和学生端页面，后端拆分出网关、注册中心、认证、用户、课程、作业、考试、分析、通知和 AI 服务；数据层完成多 schema 所有权划分；运行层提供 Docker Compose、启动脚本和多类烟测脚本；文档层保留架构、认证、数据所有权、部署、可观测性和最终交付说明。")
    add_heading(doc, "8.2. 遇到的问题与解决方法", 2)
    add_table_caption(doc, "表8-1 问题与解决方法")
    add_matrix_table(
        doc,
        ["问题", "原因", "解决方法", "效果"],
        [
            ["单体遗留接口多", "旧页面依赖路径复杂，直接删除风险高。", "保留 legacy-adapter 和兼容路由，按模块逐步切流。", "降低迁移风险。"],
            ["跨服务数据边界不清", "课程、作业、考试、分析之间存在共享数据。", "制定数据所有权矩阵，只保存外部 ID，通过 API 或事件共享。", "避免跨库直接读写。"],
            ["事件可能重复消费", "消息系统天然至少一次投递。", "消费者先写 processed_event 做幂等判断。", "避免分析结果重复累计。"],
            ["前端与后端能力不同步", "页面早期存在模拟数据和旧提示。", "增加 capability 接口、契约脚本和运行时烟测。", "页面语义与真实接口对齐。"],
            ["本地多服务启动复杂", "服务数量多、端口和依赖多。", "使用 Docker Compose 和 start-runtime-smoke-stack.ps1 编排。", "提升演示可重复性。"],
        ],
        [3.4, 4.0, 5.0, 2.0],
    )
    add_heading(doc, "8.3. 可优化与扩展方向", 2)
    add_body(doc, "后续可继续从四个方向优化。第一，接入真实讯飞星火大模型，将本地 LocalMockAiModelClient 替换为正式模型适配实现，并增加提示词版本管理。第二，补全 Helm chart 的 Deployment、Service、ConfigMap 和 Secret 模板，使系统可以部署到 Kubernetes。第三，继续收敛旧单体 legacy-route，在所有页面与接口稳定后下线过渡服务。第四，增加更细粒度的权限模型和审计日志，例如教师只能访问自己授课班级、学生只能访问自己的提交与成绩。")
    add_body(doc, "整体来看，项目已经覆盖课程考核的主要观测点：页面布局完整且具有实际业务主题；系统功能覆盖教师、学生和平台治理；分布式框架知识点在网关、注册发现、Feign、消息、熔断限流、数据拆分和可观测性中得到综合运用；报告结构完整，配有表格、架构图、UML 图和时序图，适合课程期末大作业提交和现场演示。")

    chapter(doc, "附 录")
    add_heading(doc, "附录1. 小组成员分工说明", 2)
    add_table_caption(doc, "表A-1 小组成员分工")
    add_matrix_table(
        doc,
        ["成员", "主要分工", "完成内容"],
        [
            ["詹俊英", "系统设计、后端微服务、前端页面、测试与文档", "完成 Spring Cloud 多服务拆分、Gateway 配置、主要页面联调、烟测脚本和系统设计报告。"],
        ],
        [2.5, 5.0, 7.0],
    )
    add_heading(doc, "附录2. 项目目录结构说明", 2)
    add_table_caption(doc, "表A-2 项目目录结构说明")
    add_matrix_table(
        doc,
        ["目录/模块", "类型", "主要作用", "说明"],
        [
            ["gateway", "Spring Cloud Gateway", "统一 API 入口、JWT 鉴权、路由、限流和熔断。", "所有前端请求先进入该模块。"],
            ["registry-server", "Eureka", "服务注册与发现。", "业务服务启动后注册为服务实例。"],
            ["auth-service / user-service", "业务服务", "认证、Token、用户资料和角色。", "支撑教师/学生身份识别。"],
            ["course-service", "业务服务", "课程、班级、知识点、学生加入班级。", "教学资源数据所有者。"],
            ["assignment-service / exam-service", "业务服务", "作业、提交、批改、考试和成绩。", "教学任务与评价主流程。"],
            ["analysis-service / notification-service", "业务服务", "学情分析、预警、通知和已读状态。", "通过事件和 read model 支撑页面查询。"],
            ["ai-service", "业务服务", "题目、试卷和学习建议生成。", "当前本地模型适配，可替换讯飞星火。"],
            ["common / common-events", "公共模块", "统一响应、异常、Feign、outbox、事件定义。", "减少跨服务重复代码。"],
            ["frontend/dist", "前端页面", "教师端、学生端、登录页和静态资源。", "报告截图来自该目录页面。"],
            ["docs / scripts", "文档与脚本", "架构说明、部署说明、烟测和辅助脚本。", "支撑提交、运行和验收。"],
        ],
        [3.4, 2.7, 5.0, 3.4],
    )
    add_heading(doc, "附录3. 评分点自查表", 2)
    add_table_caption(doc, "表A-3 评分点自查")
    add_matrix_table(
        doc,
        ["评分项", "分值", "报告与项目对应内容", "自查结论"],
        [
            ["页面设计", "20", "教师端和学生端多页面、统一风格、仪表盘与 CRUD 页面完整。", "满足"],
            ["系统功能实现", "50", "系统可通过 Gateway 运行，功能覆盖课程、作业、考试、分析、通知、AI 和认证。", "满足"],
            ["系统报告编写", "20", "报告按模板组织，包含自动目录、章节分页、表格、架构图、UML 图和时序图。", "满足"],
            ["课堂演示", "10", "可按教师发布任务、学生提交、分析预警、通知/AI 建议路线演示。", "满足"],
        ],
        [3.0, 1.3, 8.6, 1.6],
    )

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_report()
    print(path)
