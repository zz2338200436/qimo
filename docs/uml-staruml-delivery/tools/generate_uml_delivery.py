from __future__ import annotations

import json
import math
import re
import shutil
import textwrap
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
IMAGE_DIR = ROOT / "images"
SOURCE_DIR = ROOT / "sources"
STARUML_DIR = ROOT / "staruml-extension" / "smart-learning-uml-generator"
REPORT_DIR = ROOT / "report"
MODEL_PATH = SOURCE_DIR / "smart-learning-system.mdj"
SPEC_PATH = SOURCE_DIR / "uml-model-spec.json"

MAVEN_MODULES = [
    "parent-pom",
    "major_assignment",
    "common",
    "common-events",
    "registry-server",
    "gateway",
    "auth-service-api",
    "auth-service",
    "user-service-api",
    "user-service",
    "course-service-api",
    "course-service",
    "assignment-service-api",
    "assignment-service",
    "exam-service-api",
    "exam-service",
    "analysis-service-api",
    "analysis-service",
    "notification-service-api",
    "notification-service",
    "ai-service-api",
    "ai-service",
    "legacy-adapter",
]

FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simhei.ttf"),
    Path("C:/Windows/Fonts/simsun.ttc"),
    Path("C:/Windows/Fonts/arial.ttf"),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc") if bold else Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf") if bold else Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/arialbd.ttf") if bold else Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


FONT_TITLE = font(34, True)
FONT_SUBTITLE = font(21, True)
FONT_TEXT = font(19)
FONT_SMALL = font(16)
FONT_TINY = font(14)
FONT_BOLD = font(19, True)


INK = "#1f2937"
MUTED = "#6b7280"
LINE = "#4b5563"
BLUE = "#dbeafe"
GREEN = "#dcfce7"
YELLOW = "#fef3c7"
RED = "#fee2e2"
PURPLE = "#ede9fe"
GRAY = "#f8fafc"
WHITE = "#ffffff"
ORANGE = "#ffedd5"


@dataclass
class Node:
    id: str
    label: str
    x: int
    y: int
    w: int
    h: int
    kind: str = "class"
    fill: str = WHITE
    stereotype: str | None = None
    members: list[str] = field(default_factory=list)


@dataclass
class Edge:
    source: str
    target: str
    label: str = ""
    kind: str = "association"
    dashed: bool = False
    sequence: int | None = None


@dataclass
class Diagram:
    key: str
    title: str
    uml_type: str
    filename: str
    description: str
    evidence: list[str]
    plantuml: str
    nodes: list[Node] = field(default_factory=list)
    edges: list[Edge] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


def sequence_nodes_and_edges(lifelines: list[str], messages: list[dict[str, Any]]) -> tuple[list[Node], list[Edge]]:
    margin = 90
    x_step = 250
    nodes = [
        Node(
            f"seq-{idx}",
            label,
            margin + idx * x_step,
            180,
            210,
            860,
            kind="lifeline",
            fill=BLUE,
            members=["lifeline"],
        )
        for idx, label in enumerate(lifelines)
    ]
    name_to_id = {label: f"seq-{idx}" for idx, label in enumerate(lifelines)}
    edges = [
        Edge(
            name_to_id[msg["from"]],
            name_to_id[msg["to"]],
            msg["label"],
            "message",
            bool(msg.get("dashed", False)),
            sequence=i,
        )
        for i, msg in enumerate(messages, 1)
    ]
    return nodes, edges


def reset_dirs() -> None:
    for path in [IMAGE_DIR, SOURCE_DIR, REPORT_DIR, STARUML_DIR]:
        if path.exists():
            shutil.rmtree(path)
        path.mkdir(parents=True, exist_ok=True)
    (STARUML_DIR / "menus").mkdir(parents=True, exist_ok=True)


def display_width(text: str) -> int:
    return sum(2 if ord(c) > 127 else 1 for c in text)


def split_long_token(token: str, max_chars: int) -> list[str]:
    expanded = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", token.replace("/", "/ "))
    if expanded != token:
        return [part for part in expanded.split(" ") if part]
    parts: list[str] = []
    current = ""
    for ch in token:
        if display_width(current + ch) > max_chars and current:
            parts.append(current)
            current = ch
        else:
            current += ch
    if current:
        parts.append(current)
    return parts


def wrap_text(text: str, max_chars: int) -> list[str]:
    if not text:
        return []
    text = str(text).replace("\\n", "\n")
    lines: list[str] = []
    max_chars = max(1, max_chars)
    for raw_line in text.split("\n"):
        current = ""
        tokens = raw_line.split()
        if not tokens:
            if current:
                lines.append(current)
            continue
        for token in tokens:
            units = split_long_token(token, max_chars) if display_width(token) > max_chars else [token]
            for unit in units:
                candidate = unit if not current else f"{current} {unit}"
                if display_width(candidate) <= max_chars:
                    current = candidate
                else:
                    if current:
                        lines.append(current)
                    current = unit
        if current:
            lines.append(current)
    return lines


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    if not text:
        return 0, 0
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def draw_text_center(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
                     fnt: ImageFont.FreeTypeFont, fill: str = INK) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_text(text, max(4, int((x2 - x1) / (fnt.size * 0.62))))
    total_h = len(lines) * (fnt.size + 6)
    y = y1 + (y2 - y1 - total_h) // 2
    for line in lines:
        w, h = text_size(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - w) // 2, y), line, font=fnt, fill=fill)
        y += fnt.size + 6


def dashed_line(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
                fill: str = LINE, width: int = 2, dash: int = 12) -> None:
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    dist = math.hypot(dx, dy)
    if dist == 0:
        return
    steps = int(dist // dash)
    for i in range(steps + 1):
        if i % 2 == 0:
            a = i / max(1, steps)
            b = min(1, (i + 1) / max(1, steps))
            draw.line((x1 + dx * a, y1 + dy * a, x1 + dx * b, y1 + dy * b), fill=fill, width=width)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], *,
          fill: str = LINE, width: int = 2, dashed: bool = False, open_head: bool = False) -> None:
    if dashed:
        dashed_line(draw, start, end, fill, width)
    else:
        draw.line((*start, *end), fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 13
    pts = [
        (x2, y2),
        (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6)),
        (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6)),
    ]
    if open_head:
        draw.line((pts[1][0], pts[1][1], pts[0][0], pts[0][1], pts[2][0], pts[2][1]), fill=fill, width=width)
    else:
        draw.polygon(pts, fill=fill)


def polyline_arrow(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]], *,
                   fill: str = LINE, width: int = 2, dashed: bool = False,
                   open_head: bool = False) -> None:
    if len(points) < 2:
        return
    for start, end in zip(points, points[1:]):
        if dashed:
            dashed_line(draw, start, end, fill=fill, width=width)
        else:
            draw.line((*start, *end), fill=fill, width=width)
    x1, y1 = points[-2]
    x2, y2 = points[-1]
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 13
    pts = [
        (x2, y2),
        (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6)),
        (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6)),
    ]
    if open_head:
        draw.line((pts[1][0], pts[1][1], pts[0][0], pts[0][1], pts[2][0], pts[2][1]), fill=fill, width=width)
    else:
        draw.polygon(pts, fill=fill)


def node_anchor(src: Node, dst: Node) -> tuple[tuple[int, int], tuple[int, int]]:
    sx, sy = src.x + src.w // 2, src.y + src.h // 2
    tx, ty = dst.x + dst.w // 2, dst.y + dst.h // 2
    dx, dy = tx - sx, ty - sy
    if abs(dx) / max(1, src.w) > abs(dy) / max(1, src.h):
        start = (src.x + (src.w if dx > 0 else 0), sy)
    else:
        start = (sx, src.y + (src.h if dy > 0 else 0))
    if abs(dx) / max(1, dst.w) > abs(dy) / max(1, dst.h):
        end = (dst.x + (0 if dx > 0 else dst.w), ty)
    else:
        end = (tx, dst.y + (0 if dy > 0 else dst.h))
    return start, end


def route_edge(src: Node, dst: Node) -> list[tuple[int, int]]:
    start, end = node_anchor(src, dst)
    dx, dy = end[0] - start[0], end[1] - start[1]
    if abs(dx) < 40 or abs(dy) < 40:
        return [start, end]
    if abs(dx) >= abs(dy):
        mid_x = (start[0] + end[0]) // 2
        return [start, (mid_x, start[1]), (mid_x, end[1]), end]
    mid_y = (start[1] + end[1]) // 2
    return [start, (start[0], mid_y), (end[0], mid_y), end]


def offset_polyline(points: list[tuple[int, int]], offset: int) -> list[tuple[int, int]]:
    if not offset or len(points) < 2:
        return points
    first, last = points[0], points[-1]
    dx, dy = last[0] - first[0], last[1] - first[1]
    if abs(dx) >= abs(dy):
        return [(x, y + offset) for x, y in points]
    return [(x + offset, y) for x, y in points]


def label_point(points: list[tuple[int, int]]) -> tuple[int, int]:
    if len(points) == 2:
        return ((points[0][0] + points[1][0]) // 2, (points[0][1] + points[1][1]) // 2)
    longest = (points[0], points[1])
    longest_len = -1.0
    for a, b in zip(points, points[1:]):
        dist = math.hypot(b[0] - a[0], b[1] - a[1])
        if dist > longest_len:
            longest = (a, b)
            longest_len = dist
    return ((longest[0][0] + longest[1][0]) // 2, (longest[0][1] + longest[1][1]) // 2)


def draw_stick_actor(draw: ImageDraw.ImageDraw, x: int, y: int, label: str) -> None:
    draw.ellipse((x + 34, y, x + 70, y + 36), outline=INK, width=3)
    draw.line((x + 52, y + 36, x + 52, y + 92), fill=INK, width=3)
    draw.line((x + 20, y + 56, x + 84, y + 56), fill=INK, width=3)
    draw.line((x + 52, y + 92, x + 24, y + 132), fill=INK, width=3)
    draw.line((x + 52, y + 92, x + 82, y + 132), fill=INK, width=3)
    draw_text_center(draw, (x, y + 138, x + 104, y + 178), label, FONT_SMALL)


def draw_node(draw: ImageDraw.ImageDraw, node: Node) -> None:
    box = (node.x, node.y, node.x + node.w, node.y + node.h)
    if node.kind == "usecase":
        draw.ellipse(box, fill=node.fill, outline=INK, width=2)
        draw_text_center(draw, box, node.label, FONT_TEXT)
        return
    if node.kind == "actor":
        draw_stick_actor(draw, node.x, node.y, node.label)
        return
    if node.kind == "state":
        draw.rounded_rectangle(box, radius=24, fill=node.fill, outline=INK, width=2)
        draw_text_center(draw, box, node.label, FONT_TEXT)
        return
    if node.kind == "initial":
        draw.ellipse(box, fill=INK, outline=INK)
        return
    if node.kind == "final":
        draw.ellipse(box, fill=WHITE, outline=INK, width=3)
        pad = 8
        draw.ellipse((node.x + pad, node.y + pad, node.x + node.w - pad, node.y + node.h - pad), fill=INK, outline=INK)
        return
    if node.kind == "activity":
        draw.rounded_rectangle(box, radius=26, fill=node.fill, outline=INK, width=2)
        draw_text_center(draw, box, node.label, FONT_TEXT)
        return
    if node.kind == "decision":
        cx, cy = node.x + node.w // 2, node.y + node.h // 2
        pts = [(cx, node.y), (node.x + node.w, cy), (cx, node.y + node.h), (node.x, cy)]
        draw.polygon(pts, fill=node.fill, outline=INK)
        draw.line((*pts[0], *pts[1], *pts[2], *pts[3], *pts[0]), fill=INK, width=2)
        draw_text_center(draw, box, node.label, FONT_SMALL)
        return
    if node.kind == "component":
        draw.rectangle(box, fill=node.fill, outline=INK, width=2)
        draw.rectangle((node.x + node.w - 58, node.y + 20, node.x + node.w - 20, node.y + 42), fill=WHITE, outline=INK, width=2)
        draw.rectangle((node.x + node.w - 58, node.y + 58, node.x + node.w - 20, node.y + 80), fill=WHITE, outline=INK, width=2)
        draw_text_center(draw, (node.x + 10, node.y + 8, node.x + node.w - 68, node.y + node.h - 8), node.label, FONT_TEXT)
        return
    if node.kind == "node3d":
        x, y, w, h = node.x, node.y, node.w, node.h
        d = 22
        draw.rectangle((x, y + d, x + w - d, y + h), fill=node.fill, outline=INK, width=2)
        draw.polygon([(x, y + d), (x + d, y), (x + w, y), (x + w - d, y + d)], fill="#eef2ff", outline=INK)
        draw.polygon([(x + w - d, y + d), (x + w, y), (x + w, y + h - d), (x + w - d, y + h)], fill="#e0e7ff", outline=INK)
        draw_text_center(draw, (x + 8, y + d + 4, x + w - d - 6, y + h - 6), node.label, FONT_TEXT)
        return
    # default class/object box
    draw.rounded_rectangle(box, radius=4, fill=node.fill, outline=INK, width=2)
    y = node.y
    if node.stereotype:
        draw_text_center(draw, (node.x, y + 5, node.x + node.w, y + 30), f"<<{node.stereotype}>>", FONT_TINY, MUTED)
        y += 24
    title_lines = wrap_text(node.label, max(8, int((node.w - 18) / (FONT_BOLD.size * 0.58))))
    title_height = max(40, len(title_lines) * (FONT_BOLD.size + 6))
    header_bottom = y + 18 + title_height
    draw_text_center(draw, (node.x + 8, y + 8, node.x + node.w - 8, header_bottom - 8), node.label, FONT_BOLD)
    draw.line((node.x, header_bottom, node.x + node.w, header_bottom), fill=INK, width=1)
    ty = header_bottom + 9
    for member in node.members:
        for line in wrap_text(member, max(10, int((node.w - 22) / 9.2))):
            if ty + 18 > node.y + node.h - 8:
                draw.text((node.x + 10, ty), "...", font=FONT_SMALL, fill=INK)
                return
            draw.text((node.x + 10, ty), line, font=FONT_SMALL, fill=INK)
            ty += 22


def draw_header(draw: ImageDraw.ImageDraw, title: str, subtitle: str, width: int) -> None:
    draw.rectangle((0, 0, width, 72), fill="#0f172a")
    draw.text((36, 18), title, font=FONT_TITLE, fill=WHITE)
    draw.text((36, 78), subtitle, font=FONT_SMALL, fill=MUTED)


def render_graph_diagram(diagram: Diagram, size: tuple[int, int] = (1800, 1180)) -> Path:
    path = IMAGE_DIR / diagram.filename
    img = Image.new("RGB", size, "#ffffff")
    draw = ImageDraw.Draw(img)
    draw_header(draw, diagram.title, diagram.uml_type, size[0])
    node_map = {n.id: n for n in diagram.nodes}
    seen_pairs: dict[tuple[str, str], int] = {}
    for edge in diagram.edges:
        src, dst = node_map[edge.source], node_map[edge.target]
        pair = tuple(sorted((edge.source, edge.target)))
        seen_pairs[pair] = seen_pairs.get(pair, 0) + 1
        points = route_edge(src, dst)
        if seen_pairs[pair] > 1:
            points = offset_polyline(points, 34 * (seen_pairs[pair] - 1))
        open_head = edge.kind in {"inheritance", "dependency"}
        polyline_arrow(draw, points, dashed=edge.dashed or edge.kind == "dependency", open_head=open_head)
        if edge.label:
            mx, my = label_point(points)
            lines = wrap_text(edge.label, 18)
            text_w = max(text_size(draw, line, FONT_TINY)[0] for line in lines)
            text_h = len(lines) * 18 + 8
            draw.rounded_rectangle((mx - text_w // 2 - 8, my - text_h // 2, mx + text_w // 2 + 8, my + text_h // 2),
                                   radius=6, fill=WHITE, outline="#e5e7eb")
            ty = my - text_h // 2 + 4
            for line in lines:
                tw, _ = text_size(draw, line, FONT_TINY)
                draw.text((mx - tw // 2, ty), line, font=FONT_TINY, fill=INK)
                ty += 18
    for node in diagram.nodes:
        draw_node(draw, node)
    img.save(path)
    return path


def render_sequence(diagram: Diagram) -> Path:
    path = IMAGE_DIR / diagram.filename
    lifelines = diagram.metadata["lifelines"]
    messages = diagram.metadata["messages"]
    step = 88
    top = 150
    y_start = 255
    bottom = max(1220, y_start + step * len(messages) + 160)
    size = (2800, bottom + 90)
    img = Image.new("RGB", size, WHITE)
    draw = ImageDraw.Draw(img)
    draw_header(draw, diagram.title, diagram.uml_type, size[0])
    margin = 150
    spacing = (size[0] - margin * 2) / max(1, len(lifelines) - 1)
    xs = [int(margin + i * spacing) for i in range(len(lifelines))]
    half_box = min(136, int(spacing / 2) - 18)
    for x, label in zip(xs, lifelines):
        draw.rounded_rectangle((x - half_box, top, x + half_box, top + 64), radius=5, fill=BLUE, outline=INK, width=2)
        draw_text_center(draw, (x - half_box + 6, top + 5, x + half_box - 6, top + 59), label, FONT_TINY)
        dashed_line(draw, (x, top + 64), (x, bottom), fill="#9ca3af", width=2, dash=16)
    y = y_start
    for i, msg in enumerate(messages, 1):
        src = lifelines.index(msg["from"])
        dst = lifelines.index(msg["to"])
        x1, x2 = xs[src], xs[dst]
        if src == dst:
            loop_w = 90
            direction = -1 if x1 + loop_w + 260 > size[0] - 24 else 1
            elbow_x = x1 + direction * loop_w
            points = [(x1, y), (elbow_x, y), (elbow_x, y + 30), (x1 + direction * 10, y + 30)]
            polyline_arrow(draw, points, dashed=msg.get("dashed", False), open_head=msg.get("return", False))
            tx = elbow_x + 34 if direction > 0 else elbow_x - 34
            ty = y - 12
        else:
            arrow(draw, (x1, y), (x2, y), dashed=msg.get("dashed", False), open_head=msg.get("return", False))
            tx = min(x1, x2) + abs(x2 - x1) // 2
            ty = y - 42
        label = f"{i}. {msg['label']}"
        max_chars = 34 if abs(x2 - x1) > 360 else 24
        lines = wrap_text(label, max_chars)
        text_w = max(text_size(draw, line, FONT_TINY)[0] for line in lines)
        text_h = len(lines) * 18 + 8
        if src == dst:
            left = tx if tx > x1 else tx - text_w - 14
            left = max(18, min(left, size[0] - text_w - 32))
            label_box = (left, ty, left + text_w + 14, ty + text_h)
        else:
            left = tx - text_w // 2 - 7
            right = tx + text_w // 2 + 7
            if left < 18:
                right += 18 - left
                left = 18
            if right > size[0] - 18:
                left -= right - (size[0] - 18)
                right = size[0] - 18
            label_box = (left, ty, right, ty + text_h)
        draw.rounded_rectangle(label_box, radius=5, fill=WHITE, outline="#e5e7eb")
        text_x = label_box[0] + 7
        ty = label_box[1] + 4
        for line in lines:
            draw.text((text_x, ty), line, font=FONT_TINY, fill=INK)
            ty += 18
        y += step
    img.save(path)
    return path


def render_communication(diagram: Diagram) -> Path:
    path = IMAGE_DIR / diagram.filename
    img = Image.new("RGB", (1800, 1120), WHITE)
    draw = ImageDraw.Draw(img)
    draw_header(draw, diagram.title, diagram.uml_type, 1800)
    node_map = {n.id: n for n in diagram.nodes}
    for edge in diagram.edges:
        src, dst = node_map[edge.source], node_map[edge.target]
        points = route_edge(src, dst)
        polyline_arrow(draw, points, fill=LINE, width=2, open_head=False)
        mx, my = label_point(points)
        for idx, line in enumerate(wrap_text(edge.label, 34)):
            tw, _ = text_size(draw, line, FONT_TINY)
            draw.rectangle((mx - tw // 2 - 4, my + idx * 18 - 10, mx + tw // 2 + 4, my + idx * 18 + 9), fill=WHITE)
            draw.text((mx - tw // 2, my + idx * 18 - 10), line, font=FONT_TINY, fill=INK)
    for node in diagram.nodes:
        draw_node(draw, node)
    img.save(path)
    return path


def render_object(diagram: Diagram) -> Path:
    return render_graph_diagram(diagram, (1800, 1080))


def make_diagrams() -> list[Diagram]:
    diagrams: list[Diagram] = []

    diagrams.append(Diagram(
        key="use_case",
        title="图1 智能学习辅助系统用例图",
        uml_type="Use Case Diagram",
        filename="01-use-case.png",
        description=(
            "用例图以学生、教师、系统管理员和外部 AI/消息服务为参与者，概括平台对外提供的核心业务目标。"
            "学生侧重点是登录、学习课程、提交作业和考试、查看成绩预警与接收通知；教师侧重点是课程班级管理、作业考试发布、批改、分析学情、处理预警和使用 AI 生成资源；管理员负责基础数据和系统运行配置。"
        ),
        evidence=[
            "gateway/src/main/resources/application.yml 中的 /api/student/**、/api/teacher/**、/api/auth/**、/api/ai/** 路由",
            "docs/service-boundary-and-monolith-shrink-plan.md 中的服务边界总表",
        ],
        plantuml="""@startuml
left to right direction
actor 学生
actor 教师
actor 系统管理员
actor "AI 服务" as AI
actor "消息服务" as MQ
rectangle 智能学习辅助系统 {
  usecase "登录与角色切换" as UC1
  usecase "浏览课程与班级" as UC2
  usecase "提交作业" as UC3
  usecase "参加考试并提交" as UC4
  usecase "查看成绩与知识点掌握" as UC5
  usecase "接收通知与预警" as UC6
  usecase "维护课程/班级/知识点" as UC7
  usecase "发布作业/考试" as UC8
  usecase "批改与查看提交" as UC9
  usecase "触发学情分析" as UC10
  usecase "处理学情预警" as UC11
  usecase "AI 生成题目/建议" as UC12
  usecase "管理基础数据与服务" as UC13
}
学生 --> UC1
学生 --> UC2
学生 --> UC3
学生 --> UC4
学生 --> UC5
学生 --> UC6
教师 --> UC1
教师 --> UC7
教师 --> UC8
教师 --> UC9
教师 --> UC10
教师 --> UC11
教师 --> UC12
系统管理员 --> UC13
UC12 --> AI
UC6 --> MQ
@enduml""",
        nodes=[
            Node("student", "学生", 70, 250, 110, 180, "actor"),
            Node("teacher", "教师", 70, 610, 110, 180, "actor"),
            Node("admin", "系统管理员", 70, 900, 130, 180, "actor"),
            Node("ai", "AI 服务", 1590, 320, 110, 180, "actor"),
            Node("mq", "消息服务", 1590, 700, 110, 180, "actor"),
            Node("uc1", "登录与角色切换", 360, 150, 210, 90, "usecase", BLUE),
            Node("uc2", "浏览课程与班级", 650, 150, 210, 90, "usecase", GREEN),
            Node("uc3", "提交作业", 940, 150, 210, 90, "usecase", GREEN),
            Node("uc4", "参加考试并提交", 1230, 150, 230, 90, "usecase", GREEN),
            Node("uc5", "查看成绩与知识点掌握", 650, 330, 260, 90, "usecase", YELLOW),
            Node("uc6", "接收通知与预警", 1000, 330, 230, 90, "usecase", YELLOW),
            Node("uc7", "维护课程/班级/知识点", 360, 570, 260, 90, "usecase", BLUE),
            Node("uc8", "发布作业/考试", 700, 570, 230, 90, "usecase", BLUE),
            Node("uc9", "批改与查看提交", 1010, 570, 230, 90, "usecase", BLUE),
            Node("uc10", "触发学情分析", 360, 750, 230, 90, "usecase", PURPLE),
            Node("uc11", "处理学情预警", 690, 750, 230, 90, "usecase", PURPLE),
            Node("uc12", "AI 生成题目/建议", 1030, 750, 240, 90, "usecase", ORANGE),
            Node("uc13", "管理基础数据与服务", 520, 940, 280, 90, "usecase", GRAY),
        ],
        edges=[
            Edge("student", "uc1"), Edge("student", "uc2"), Edge("student", "uc3"), Edge("student", "uc4"),
            Edge("student", "uc5"), Edge("student", "uc6"),
            Edge("teacher", "uc1"), Edge("teacher", "uc7"), Edge("teacher", "uc8"), Edge("teacher", "uc9"),
            Edge("teacher", "uc10"), Edge("teacher", "uc11"), Edge("teacher", "uc12"),
            Edge("admin", "uc13"),
            Edge("uc12", "ai", "<<include>>", "dependency", True),
            Edge("uc6", "mq", "异步投递", "dependency", True),
        ],
    ))

    class_nodes = [
        Node("AuthController", "AuthController", 70, 170, 245, 128, fill=BLUE, stereotype="controller",
             members=["login()", "refresh()", "switchRole()"]),
        Node("AuthService", "AuthApplicationService", 390, 170, 260, 128, fill=GREEN, stereotype="service",
             members=["authenticate()", "refreshToken()"]),
        Node("JwtService", "JwtTokenService", 710, 170, 245, 128, fill=GREEN, stereotype="service",
             members=["issueTokenPair()", "parseClaims()"]),
        Node("Credential", "AuthCredentialEntity", 1040, 170, 260, 128, fill=YELLOW, stereotype="entity",
             members=["userId", "username", "passwordHash", "enabled"]),
        Node("User", "UserEntity", 1390, 170, 245, 128, fill=YELLOW, stereotype="entity",
             members=["id", "username", "realName", "majorId"]),

        Node("CourseController", "TeacherCourseAdminController", 70, 380, 245, 142, fill=BLUE, stereotype="controller",
             members=["listCourses()", "saveCourse()", "assignClass()"]),
        Node("CourseService", "CourseApplicationService", 390, 380, 260, 142, fill=GREEN, stereotype="service",
             members=["queryTeacherCourses()", "upsertCourse()", "assignStudent()"]),
        Node("CourseRepo", "CourseRepository", 710, 380, 245, 142, fill=GRAY, stereotype="repository",
             members=["findCourses()", "saveCourse()", "findClasses()"]),
        Node("Course", "CourseEntity", 1040, 380, 260, 142, fill=YELLOW, stereotype="entity",
             members=["id", "courseName", "teacherId", "semester"]),
        Node("Class", "CourseClassEntity", 1390, 380, 245, 142, fill=YELLOW, stereotype="entity",
             members=["id", "className", "majorId", "studentCount"]),

        Node("AssignmentController", "StudentAssignmentController", 70, 610, 245, 142, fill=BLUE, stereotype="controller",
             members=["listAssignments()", "submitAssignment()"]),
        Node("AssignmentService", "AssignmentApplicationService", 390, 610, 260, 142, fill=GREEN, stereotype="service",
             members=["submitAssignment()", "publishEvent()"]),
        Node("AssignmentRepo", "AssignmentRepository", 710, 610, 245, 142, fill=GRAY, stereotype="repository",
             members=["findAssignment()", "insertSubmission()"]),
        Node("Assignment", "AssignmentRecord", 1040, 610, 260, 142, fill=YELLOW, stereotype="domain",
             members=["id", "courseId", "classId", "deadline"]),
        Node("AssignmentSubmission", "AssignmentSubmissionRecord", 1390, 610, 245, 156, fill=YELLOW, stereotype="domain",
             members=["assignmentId", "studentId", "content", "graded", "score"]),

        Node("ExamController", "StudentExamController", 70, 850, 245, 142, fill=BLUE, stereotype="controller",
             members=["listExams()", "submitExam()"]),
        Node("ExamService", "ExamApplicationService", 390, 850, 260, 142, fill=GREEN, stereotype="service",
             members=["submitExam()", "autoGradeIfPossible()", "persistExamFinishedEvent()"]),
        Node("ExamRepo", "ExamRepository", 710, 850, 245, 142, fill=GRAY, stereotype="repository",
             members=["findExam()", "insertSubmission()", "findQuestions()"]),
        Node("Exam", "ExamRecord", 1040, 850, 260, 142, fill=YELLOW, stereotype="domain",
             members=["id", "courseId", "startTime", "endTime", "status"]),
        Node("ExamSubmission", "ExamSubmissionRecord", 1390, 850, 245, 166, fill=YELLOW, stereotype="domain",
             members=["examId", "studentId", "timeTaken", "graded", "score"]),

        Node("AnalysisService", "AnalysisQueryService", 390, 1060, 260, 128, fill=PURPLE, stereotype="service",
             members=["queryTrend()", "queryMastery()", "listWarnings()"]),
        Node("Warning", "EarlyWarningEntity", 710, 1060, 245, 128, fill=RED, stereotype="entity",
             members=["studentId", "courseId", "warningLevel", "resolved"]),
        Node("Notification", "NotificationEntity", 1040, 1060, 260, 128, fill=ORANGE, stereotype="entity",
             members=["studentId", "type", "title", "read"]),
        Node("AI", "AiGenerationEntity", 1390, 1060, 245, 128, fill=ORANGE, stereotype="entity",
             members=["requestType", "inputSummary", "outputSummary", "modelName"]),
    ]
    class_edges = [
        Edge("AuthController", "AuthService"), Edge("AuthService", "JwtService"), Edge("AuthService", "Credential"),
        Edge("AuthService", "User", "", "dependency", True),
        Edge("CourseController", "CourseService"), Edge("CourseService", "CourseRepo"), Edge("CourseRepo", "Course"),
        Edge("CourseRepo", "Class"), Edge("AssignmentController", "AssignmentService"), Edge("AssignmentService", "AssignmentRepo"),
        Edge("AssignmentRepo", "Assignment"), Edge("AssignmentRepo", "AssignmentSubmission"),
        Edge("ExamController", "ExamService"), Edge("ExamService", "ExamRepo"), Edge("ExamRepo", "Exam"),
        Edge("ExamRepo", "ExamSubmission"), Edge("ExamService", "AnalysisService", "", "dependency", True),
        Edge("AnalysisService", "Warning"), Edge("AnalysisService", "Notification", "", "dependency", True),
    ]
    diagrams.append(Diagram(
        key="class",
        title="图2 智能学习辅助系统类图",
        uml_type="Class Diagram",
        filename="02-class.png",
        description=(
            "类图按微服务内部常见的 Controller、ApplicationService、Repository、Entity/Record 分层组织，"
            "同时展示 Auth、Course、Assignment、Exam、Analysis、Notification、AI 等核心领域对象的依赖关系。"
            "图中虚线表示跨服务 API 或事件依赖，实线表示同一服务内部调用或持久化依赖。"
        ),
        evidence=[
            "各服务 src/main/java 中的 Controller、Service、Repository、Entity/Record 类",
            "common-events 模块中的 ExamFinishedEvent、EarlyWarningRaisedEvent 等领域事件",
        ],
        plantuml="""@startuml
skinparam classAttributeIconSize 0
package Auth {
  class AuthController <<controller>>
  class AuthApplicationService <<service>>
  class JwtTokenService <<service>>
  class AuthCredentialEntity <<entity>>
}
package User {
  class UserEntity <<entity>>
  class RoleEntity <<entity>>
}
package Course {
  class TeacherCourseAdminController <<controller>>
  class CourseApplicationService <<service>>
  interface CourseRepository
  class CourseEntity <<entity>>
  class CourseClassEntity <<entity>>
}
package Assignment {
  class StudentAssignmentController <<controller>>
  class AssignmentApplicationService <<service>>
  interface AssignmentRepository
  class AssignmentRecord
  class AssignmentSubmissionRecord
}
package Exam {
  class StudentExamController <<controller>>
  class ExamApplicationService <<service>>
  interface ExamRepository
  class ExamRecord
  class ExamSubmissionRecord
}
package Analysis {
  class AnalysisQueryService <<service>>
  class EarlyWarningEntity <<entity>>
}
package Notification {
  class NotificationEntity <<entity>>
}
package AI {
  class AiGenerationEntity <<entity>>
}
AuthController --> AuthApplicationService
AuthApplicationService --> JwtTokenService
AuthApplicationService --> AuthCredentialEntity
AuthApplicationService ..> UserEntity : Feign 查询角色
TeacherCourseAdminController --> CourseApplicationService
CourseApplicationService --> CourseRepository
CourseRepository --> CourseEntity
CourseRepository --> CourseClassEntity
StudentAssignmentController --> AssignmentApplicationService
AssignmentApplicationService --> AssignmentRepository
AssignmentRepository --> AssignmentRecord
AssignmentRepository --> AssignmentSubmissionRecord
StudentExamController --> ExamApplicationService
ExamApplicationService --> ExamRepository
ExamRepository --> ExamRecord
ExamRepository --> ExamSubmissionRecord
ExamApplicationService ..> AnalysisQueryService : ExamFinishedEvent
AnalysisQueryService --> EarlyWarningEntity
AnalysisQueryService ..> NotificationEntity : EarlyWarningRaisedEvent
AiGenerationEntity ..> ExamApplicationService : 生成结果由业务服务落库
@enduml""",
        nodes=class_nodes,
        edges=class_edges,
    ))

    diagrams.append(Diagram(
        key="object",
        title="图3 考试提交场景对象图",
        uml_type="Object Diagram",
        filename="03-object.png",
        description=(
            "对象图描述一次具体考试提交后的运行时对象快照。示例中学生张三通过网关提交考试，"
            "ExamSubmissionRecord 保存提交事实，ExamFinishedEvent 进入 outbox 后由分析服务消费，最终形成掌握度、成绩趋势、预警和通知对象。"
        ),
        evidence=[
            "exam-service/domain/ExamSubmissionRecord.java 字段 examId、studentId、timeTaken、graded、score",
            "analysis-service/repository/EarlyWarningEntity.java 与 notification-service/repository/NotificationEntity.java",
        ],
        plantuml="""@startuml
object "student:UserEntity" as student {
  id = 10023
  realName = 张三
  activeRole = STUDENT
}
object "exam:ExamRecord" as exam {
  id = 301
  title = Java 基础测验
  courseId = 12
}
object "submission:ExamSubmissionRecord" as sub {
  examId = 301
  studentId = 10023
  graded = true
  score = 78
}
object "event:ExamFinishedEvent" as event {
  aggregate = ExamSubmission/301-10023
}
object "warning:EarlyWarningEntity" as warning {
  warningLevel = MEDIUM
  resolved = false
}
object "notice:NotificationEntity" as notice {
  type = EARLY_WARNING
  read = false
}
student --> sub
exam --> sub
sub --> event
event --> warning
warning --> notice
@enduml""",
        nodes=[
            Node("student", "student:UserEntity", 80, 190, 300, 150, fill=BLUE, members=["id = 10023", "realName = 张三", "activeRole = STUDENT"]),
            Node("exam", "exam:ExamRecord", 520, 190, 300, 150, fill=GREEN, members=["id = 301", "title = Java 基础测验", "courseId = 12"]),
            Node("sub", "submission:ExamSubmissionRecord", 960, 190, 360, 178, fill=YELLOW, members=["examId = 301", "studentId = 10023", "graded = true", "score = 78"]),
            Node("event", "event:ExamFinishedEvent", 960, 480, 360, 150, fill=PURPLE, members=["aggregate = ExamSubmission", "eventId = EXAM-301-10023"]),
            Node("warning", "warning:EarlyWarningEntity", 520, 720, 330, 160, fill=RED, members=["warningLevel = MEDIUM", "resolved = false", "studentId = 10023"]),
            Node("notice", "notice:NotificationEntity", 80, 720, 330, 160, fill=ORANGE, members=["type = EARLY_WARNING", "read = false", "title = 学情预警"]),
        ],
        edges=[
            Edge("student", "sub", "提交人"), Edge("exam", "sub", "所属考试"), Edge("sub", "event", "产生事件"),
            Edge("event", "warning", "分析生成"), Edge("warning", "notice", "推送通知"),
        ],
    ))

    sequence_lifelines = ["学生", "前端页面", "Gateway", "ExamController", "ExamApplicationService", "sc_exam", "RabbitMQ", "Analysis Service", "Notification Service"]
    sequence_messages = [
        {"from": "学生", "to": "前端页面", "label": "填写答案并提交"},
        {"from": "前端页面", "to": "Gateway", "label": "POST /api/student/exams/{id}/submit"},
        {"from": "Gateway", "to": "Gateway", "label": "JWT 校验、限流、熔断"},
        {"from": "Gateway", "to": "ExamController", "label": "转发提交请求"},
        {"from": "ExamController", "to": "ExamApplicationService", "label": "submitExam(studentId, examId, answers)"},
        {"from": "ExamApplicationService", "to": "sc_exam", "label": "保存 ExamSubmissionRecord"},
        {"from": "ExamApplicationService", "to": "ExamApplicationService", "label": "autoGradeIfPossible()"},
        {"from": "ExamApplicationService", "to": "sc_exam", "label": "写 outbox_event(ExamFinishedEvent)"},
        {"from": "ExamApplicationService", "to": "ExamController", "label": "返回提交与成绩", "return": True, "dashed": True},
        {"from": "ExamController", "to": "Gateway", "label": "ResponseResult", "return": True, "dashed": True},
        {"from": "Gateway", "to": "前端页面", "label": "提交成功", "return": True, "dashed": True},
        {"from": "sc_exam", "to": "RabbitMQ", "label": "OutboxRelayJob 发布事件"},
        {"from": "RabbitMQ", "to": "Analysis Service", "label": "消费 ExamFinishedEvent"},
        {"from": "Analysis Service", "to": "Analysis Service", "label": "更新掌握度、成绩趋势、预警"},
        {"from": "RabbitMQ", "to": "Notification Service", "label": "消费事件或预警事件"},
        {"from": "Notification Service", "to": "Notification Service", "label": "创建通知"},
    ]
    sequence_nodes, sequence_edges = sequence_nodes_and_edges(sequence_lifelines, sequence_messages)
    diagrams.append(Diagram(
        key="sequence",
        title="图4 学生考试提交顺序图",
        uml_type="Sequence Diagram",
        filename="04-sequence.png",
        description=(
            "顺序图选择“学生提交考试并触发分析通知”作为核心业务场景。"
            "浏览器请求先经 Gateway 鉴权、限流和熔断，再进入 exam-service 保存提交、自动阅卷和写 outbox；"
            "RabbitMQ 将 ExamFinishedEvent 分发给 analysis-service 与 notification-service，完成成绩趋势、知识点掌握、预警和通知更新。"
        ),
        evidence=[
            "gateway/application.yml 的 student-exam-submit-route",
            "ExamApplicationService 中 submitExam、autoGradeIfPossible、persistExamFinishedEvent 相关方法",
            "common-events 中的 ExamFinishedEvent",
        ],
        plantuml="""@startuml
actor 学生
participant "前端页面" as UI
participant Gateway
participant "ExamController" as C
participant "ExamApplicationService" as S
database "sc_exam" as DB
queue RabbitMQ
participant "Analysis Service" as A
participant "Notification Service" as N
学生 -> UI : 填写答案并提交
UI -> Gateway : POST /api/student/exams/{id}/submit
Gateway -> Gateway : JWT校验/限流/熔断
Gateway -> C : 转发提交请求
C -> S : submitExam(studentId, examId, answers)
S -> DB : 保存 ExamSubmissionRecord
S -> S : autoGradeIfPossible()
S -> DB : 写 outbox_event(ExamFinishedEvent)
S --> C : 返回提交与成绩
C --> Gateway : ResponseResult
Gateway --> UI : 提交成功
DB -> RabbitMQ : OutboxRelayJob 发布事件
RabbitMQ -> A : 消费 ExamFinishedEvent
A -> A : 更新掌握度/成绩趋势/预警
RabbitMQ -> N : 消费事件或预警事件
N -> N : 创建通知
@enduml""",
        nodes=sequence_nodes,
        edges=sequence_edges,
        metadata={"lifelines": sequence_lifelines, "messages": sequence_messages},
    ))

    diagrams.append(Diagram(
        key="communication",
        title="图5 学生考试提交协作图",
        uml_type="Communication Diagram",
        filename="05-communication.png",
        description=(
            "协作图从对象协同角度重新表达同一考试提交场景，重点不是时间轴，而是对象之间的链接和消息编号。"
            "可以看到前端、Gateway、考试服务、数据库、消息队列、分析服务和通知服务各自承担的协作责任。"
        ),
        evidence=[
            "Gateway 路由、ExamApplicationService、Outbox/RabbitMQ、Analysis/Notification 消费处理链",
        ],
        plantuml="""@startuml
object 学生
object 前端页面
object Gateway
object ExamController
object ExamApplicationService
database sc_exam
queue RabbitMQ
object AnalysisService
object NotificationService
学生 -- 前端页面 : 1 提交答案
前端页面 -- Gateway : 2 POST submit
Gateway -- ExamController : 3 转发
ExamController -- ExamApplicationService : 4 submitExam
ExamApplicationService -- sc_exam : 5 保存提交\\n6 写事件
sc_exam -- RabbitMQ : 7 发布事件
RabbitMQ -- AnalysisService : 8 分析事件
AnalysisService -- NotificationService : 9 预警通知
NotificationService -- 学生 : 10 通知可见
@enduml""",
        nodes=[
            Node("student", "学生", 80, 230, 220, 92, fill=BLUE),
            Node("ui", "前端页面", 420, 140, 240, 92, fill=BLUE),
            Node("gateway", "Gateway", 780, 140, 240, 92, fill=GREEN),
            Node("controller", "ExamController", 1140, 140, 260, 92, fill=GREEN),
            Node("service", "ExamApplicationService", 1140, 390, 300, 104, fill=GREEN),
            Node("db", "sc_exam", 790, 590, 240, 92, fill=YELLOW),
            Node("mq", "RabbitMQ", 440, 590, 240, 92, fill=PURPLE),
            Node("analysis", "AnalysisService", 430, 830, 260, 92, fill=PURPLE),
            Node("notification", "NotificationService", 80, 830, 280, 92, fill=ORANGE),
        ],
        edges=[
            Edge("student", "ui", "1: 提交答案"),
            Edge("ui", "gateway", "2: POST submit"),
            Edge("gateway", "controller", "3: 转发"),
            Edge("controller", "service", "4: submitExam"),
            Edge("service", "db", "5: 保存提交\\n6: 写 outbox"),
            Edge("db", "mq", "7: 发布事件"),
            Edge("mq", "analysis", "8: 分析事件"),
            Edge("analysis", "notification", "9: 预警通知"),
            Edge("notification", "student", "10: 通知可见"),
        ],
    ))

    diagrams.append(Diagram(
        key="state",
        title="图6 考试提交记录状态图",
        uml_type="Statechart Diagram",
        filename="06-state.png",
        description=(
            "状态图选择 ExamSubmissionRecord 作为生命周期对象。"
            "从未提交开始，学生提交后进入已提交；客观题可自动阅卷进入已评分，主观题或异常情况进入待人工批改；"
            "教师复核后完成归档，异常情况下可回退为需重提交。该对象有清晰的状态迁移，适合状态图建模。"
        ),
        evidence=[
            "ExamSubmissionRecord 中 graded、score、teacherComment 字段",
            "ExamApplicationService 的自动阅卷和教师批改相关路径",
        ],
        plantuml="""@startuml
[*] --> 未提交
未提交 --> 已提交 : 学生提交答案
已提交 --> 自动阅卷中 : 客观题可评分
已提交 --> 待人工批改 : 含主观题/需复核
自动阅卷中 --> 已评分 : 计算得分
自动阅卷中 --> 待人工批改 : 自动评分失败
待人工批改 --> 已评分 : 教师批改
已评分 --> 已归档 : 发布成绩/写事件
已评分 --> 待人工批改 : 教师重新批改
已提交 --> 需重提交 : 内容异常/超时
需重提交 --> 已提交 : 学生重新提交
已归档 --> [*]
@enduml""",
        nodes=[
            Node("start", "", 140, 210, 34, 34, "initial"),
            Node("unsubmitted", "未提交", 270, 185, 210, 84, "state", BLUE),
            Node("submitted", "已提交", 590, 185, 210, 84, "state", GREEN),
            Node("auto", "自动阅卷中", 900, 140, 230, 84, "state", YELLOW),
            Node("manual", "待人工批改", 900, 330, 230, 84, "state", YELLOW),
            Node("graded", "已评分", 1210, 235, 210, 84, "state", GREEN),
            Node("archived", "已归档", 1500, 235, 210, 84, "state", PURPLE),
            Node("resubmit", "需重提交", 590, 460, 210, 84, "state", RED),
            Node("end", "", 1650, 470, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "unsubmitted"), Edge("unsubmitted", "submitted", "学生提交答案"),
            Edge("submitted", "auto", "客观题可评分"), Edge("submitted", "manual", "含主观题/需复核"),
            Edge("auto", "graded", "计算得分"), Edge("auto", "manual", "自动评分失败"),
            Edge("manual", "graded", "教师批改"), Edge("graded", "archived", "发布成绩/写事件"),
            Edge("graded", "manual", "教师重新批改"), Edge("submitted", "resubmit", "内容异常/超时"),
            Edge("resubmit", "submitted", "重新提交"), Edge("archived", "end"),
        ],
    ))

    diagrams.append(Diagram(
        key="activity",
        title="图7 考试提交与分析活动图",
        uml_type="Activity Diagram",
        filename="07-activity.png",
        description=(
            "活动图展示从学生提交考试到分析通知完成的控制流。"
            "图中包含登录校验、考试开放校验、保存提交、自动/人工批改分支、事件发布、分析更新和通知生成，"
            "体现了同步请求与异步事件处理的结合。"
        ),
        evidence=[
            "Gateway 的鉴权与 student-exam-submit-route",
            "exam-service 保存提交和 outbox 事件发布",
            "analysis-service、notification-service 的事件消费职责",
        ],
        plantuml="""@startuml
start
:学生打开考试页面;
:填写答案并点击提交;
if (JWT 有效且未限流?) then (是)
  :Gateway 转发到 exam-service;
else (否)
  :返回登录或限流提示;
  stop
endif
if (考试开放且学生有权限?) then (是)
  :保存 ExamSubmissionRecord;
else (否)
  :返回不可提交原因;
  stop
endif
if (可自动阅卷?) then (是)
  :计算得分并标记 graded;
else (否)
  :进入待人工批改;
endif
:写入 ExamFinishedEvent outbox;
:返回提交结果;
:OutboxRelay 发布到 RabbitMQ;
fork
  :Analysis 更新成绩趋势和知识点掌握;
  if (达到预警条件?) then (是)
    :生成 EarlyWarning;
  endif
fork again
  :Notification 创建考试/预警通知;
end fork
stop
@enduml""",
        nodes=[
            Node("start", "", 920, 150, 34, 34, "initial"),
            Node("open", "学生打开考试页面", 780, 235, 320, 78, "activity", BLUE),
            Node("fill", "填写答案并点击提交", 760, 355, 360, 78, "activity", BLUE),
            Node("auth", "JWT 有效且未限流?", 815, 495, 250, 118, "decision", YELLOW),
            Node("deny", "返回登录或限流提示", 1235, 515, 300, 78, "activity", RED),
            Node("forward", "Gateway 转发到 exam-service", 740, 675, 400, 78, "activity", GREEN),
            Node("permission", "考试开放且学生有权限?", 800, 825, 280, 126, "decision", YELLOW),
            Node("reject", "返回不可提交原因", 360, 845, 300, 78, "activity", RED),
            Node("save", "保存 ExamSubmissionRecord", 740, 1020, 400, 78, "activity", GREEN),
            Node("auto", "可自动阅卷?", 815, 1165, 250, 118, "decision", YELLOW),
            Node("grade", "计算得分并标记 graded", 1325, 1165, 340, 78, "activity", GREEN),
            Node("manual", "进入待人工批改", 1325, 1295, 340, 78, "activity", ORANGE),
            Node("outbox", "写入 ExamFinishedEvent outbox", 740, 1435, 400, 78, "activity", PURPLE),
            Node("return", "返回提交结果", 360, 1435, 300, 78, "activity", BLUE),
            Node("publish", "OutboxRelay 发布到 RabbitMQ", 740, 1585, 400, 78, "activity", PURPLE),
            Node("analysis", "Analysis 更新趋势/掌握度/预警", 420, 1740, 410, 78, "activity", PURPLE),
            Node("notice", "Notification 创建通知", 1050, 1740, 360, 78, "activity", ORANGE),
            Node("end", "", 1500, 1760, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "open"), Edge("open", "fill"), Edge("fill", "auth"),
            Edge("auth", "deny"), Edge("auth", "forward"), Edge("forward", "permission"),
            Edge("permission", "reject"), Edge("permission", "save"), Edge("save", "auto"),
            Edge("auto", "grade"), Edge("auto", "manual"), Edge("grade", "outbox"),
            Edge("manual", "outbox"), Edge("outbox", "return"), Edge("outbox", "publish"),
            Edge("publish", "analysis"), Edge("publish", "notice"), Edge("analysis", "end"), Edge("notice", "end"),
        ],
    ))

    diagrams.append(Diagram(
        key="component",
        title="图8 智能学习辅助系统组件图",
        uml_type="Component Diagram",
        filename="08-component.png",
        description=(
            "组件图展示平台从前端到网关、注册中心、业务微服务、公共模块、事件总线和数据存储的静态组件依赖。"
            "每个业务服务只拥有自己的 schema，通过 Feign/API 契约或领域事件协作，旧单体和 legacy-adapter 作为迁移期兼容组件保留。"
        ),
        evidence=[
            "根 pom.xml 聚合模块列表",
            "docs/data-ownership.md 的服务与数据库映射",
            "docker-compose.yml 的服务编排",
        ],
        plantuml="""@startuml
component "Frontend 静态页面" as Frontend
component "Gateway\\n路由/鉴权/限流/熔断" as Gateway
component "Registry Server\\nEureka" as Registry
component "业务微服务层\\nAuth/User/Course\\nAssignment/Exam\\nAnalysis/Notification/AI" as Services
component "legacy-adapter\\n迁移期只读兼容" as Legacy
component "common\\n响应/异常/Feign/MDC/outbox" as Common
component "common-events\\n领域事件契约" as Events
database "MySQL\\nmajor_assignment + sc_*" as MySQL
database "Redis\\nJWT/JTI/验证码/限流" as Redis
queue "RabbitMQ\\n领域事件总线" as RabbitMQ
Frontend --> Gateway
Gateway ..> Registry
Gateway --> Services
Gateway --> Legacy
Services ..> Common
Services ..> Events
Services --> MySQL
Services --> Redis
Services --> RabbitMQ
Legacy --> MySQL
@enduml""",
        nodes=[
            Node("frontend", "Frontend\\n静态页面", 80, 190, 280, 105, "component", BLUE),
            Node("gateway", "Gateway\\n路由/鉴权/限流/熔断", 470, 180, 330, 125, "component", GREEN),
            Node("registry", "Registry Server\\nEureka", 910, 175, 300, 110, "component", GRAY),
            Node("common", "common\\n响应/异常/Feign/MDC/outbox", 1400, 140, 360, 110, "component", GRAY),
            Node("events", "common-events\\n领域事件契约", 1400, 340, 360, 110, "component", PURPLE),
            Node("services", "业务微服务层\\nAuth/User/Course\\nAssignment/Exam\\nAnalysis/Notification/AI", 470, 520, 760, 180, "component", BLUE),
            Node("legacy", "legacy-adapter\\n迁移期只读兼容", 1370, 535, 340, 120, "component", RED),
            Node("mysql", "MySQL\\nmajor_assignment + sc_*", 270, 900, 360, 120, "component", YELLOW),
            Node("redis", "Redis\\nJWT/JTI/验证码/限流", 820, 900, 360, 120, "component", YELLOW),
            Node("rabbit", "RabbitMQ\\n领域事件总线", 1360, 900, 380, 120, "component", PURPLE),
        ],
        edges=[
            Edge("frontend", "gateway"),
            Edge("gateway", "registry", "", "dependency", True),
            Edge("gateway", "services"),
            Edge("gateway", "legacy"),
            Edge("services", "common", "", "dependency", True),
            Edge("services", "events", "", "dependency", True),
            Edge("services", "mysql"),
            Edge("services", "redis"),
            Edge("services", "rabbit"),
            Edge("legacy", "mysql"),
        ],
    ))

    diagrams.append(Diagram(
        key="deployment",
        title="图9 Docker Compose 部署图",
        uml_type="Deployment Diagram",
        filename="09-deployment.png",
        description=(
            "部署图依据 docker-compose.yml 表示本地/测试环境的运行拓扑。"
            "浏览器访问前端和 Gateway，Gateway 通过 Eureka 发现业务服务；各服务连接 MySQL 多 schema，"
            "认证与网关使用 Redis，作业/考试/分析/通知通过 RabbitMQ 协作，Prometheus 与 Grafana 提供观测能力。"
        ),
        evidence=[
            "docker-compose.yml 中 qimo-platform 网络、端口、容器和依赖关系",
            "docs/deployment.md 中本地编排和中间件说明",
        ],
        plantuml="""@startuml
node "开发者电脑/浏览器" as client
node "qimo-platform Docker Network" {
  node "gateway:8080\\n统一入口" as gateway
  node "registry-server:8761\\nEureka" as registry
  node "业务服务容器组\\nauth/user/course\\nassignment/exam\\nanalysis/notification/ai" as services
  node "legacy-monolith:8090->8080\\n旧单体兼容" as mono
  database "mysql:3306\\nmajor_assignment + sc_*" as mysql
  node "redis:6379\\nJWT/JTI/限流" as redis
  queue "rabbitmq:5672/15672\\n事件总线" as rabbit
  node "prometheus:9090\\n指标采集" as prom
  node "grafana:3000\\n监控面板" as grafana
}
client --> gateway
gateway ..> registry
gateway --> services
gateway --> mono
gateway --> redis
services --> mysql
services --> redis
services --> rabbit
mono --> mysql
prom ..> gateway
prom ..> services
grafana --> prom
@enduml""",
        nodes=[
            Node("client", "开发者电脑/浏览器", 90, 220, 320, 120, "node3d", BLUE),
            Node("gateway", "gateway:8080\\n统一入口", 560, 200, 320, 130, "node3d", GREEN),
            Node("registry", "registry-server:8761\\nEureka", 1010, 205, 330, 120, "node3d", GRAY),
            Node("services", "业务服务容器组\\nauth/user/course\\nassignment/exam\\nanalysis/notification/ai", 555, 535, 580, 185, "node3d", BLUE),
            Node("mono", "legacy-monolith:8090->8080\\n旧单体兼容", 1350, 545, 380, 135, "node3d", RED),
            Node("mysql", "mysql:3306\\nmajor_assignment + sc_*", 230, 925, 390, 130, "node3d", YELLOW),
            Node("redis", "redis:6379\\nJWT/JTI/限流", 760, 925, 330, 120, "node3d", YELLOW),
            Node("rabbit", "rabbitmq:5672/15672\\n事件总线", 1280, 925, 390, 120, "node3d", PURPLE),
            Node("prom", "prometheus:9090\\n指标采集", 1485, 200, 330, 120, "node3d", GRAY),
            Node("grafana", "grafana:3000\\n监控面板", 1485, 385, 330, 120, "node3d", GRAY),
        ],
        edges=[
            Edge("client", "gateway"),
            Edge("gateway", "registry", "", "dependency", True),
            Edge("gateway", "services"),
            Edge("gateway", "mono"),
            Edge("gateway", "redis", "", "dependency", True),
            Edge("services", "mysql"),
            Edge("services", "redis"),
            Edge("services", "rabbit"),
            Edge("mono", "mysql"),
            Edge("prom", "gateway", "", "dependency", True),
            Edge("prom", "services", "", "dependency", True),
            Edge("grafana", "prom", "", "dependency", True),
        ],
    ))

    for diagram in diagrams:
        diagram.metadata.setdefault("group", "core")
    diagrams.extend(make_business_extension_diagrams(start_index=len(diagrams) + 1))
    return diagrams


def make_business_extension_diagrams(start_index: int) -> list[Diagram]:
    diagrams: list[Diagram] = []

    def filename(offset: int, key: str) -> str:
        return f"{start_index + offset:02d}-{key}.png"

    def meta(focus: str, size: tuple[int, int] | None = None) -> dict[str, Any]:
        data: dict[str, Any] = {"group": "extension", "focus": focus}
        if size:
            data["size"] = list(size)
        return data

    diagrams.append(Diagram(
        key="role_use_case",
        title="图10 角色分组用例图",
        uml_type="Use Case Diagram",
        filename=filename(0, "role-use-case"),
        description=(
            "角色分组用例图在系统总用例图基础上进一步拆分学生、教师和管理员的典型目标。"
            "学生关注学习、提交、成绩和通知；教师关注课程、作业、考试、分析和 AI 辅助；管理员关注账号、角色、字典和运行配置。"
        ),
        evidence=[
            "gateway 路由按 /api/student/**、/api/teacher/**、/api/admin/** 与 /api/auth/** 分组",
            "major_assignment 静态页面中 student/teacher 管理入口与迁移后各微服务接口职责",
        ],
        plantuml="""@startuml
left to right direction
actor 学生
actor 教师
actor 管理员
rectangle "学生端" {
  usecase "课程学习" as S1
  usecase "作业提交" as S2
  usecase "考试提交" as S3
  usecase "成绩与预警查看" as S4
}
rectangle "教师端" {
  usecase "课程班级维护" as T1
  usecase "作业发布批改" as T2
  usecase "考试发布阅卷" as T3
  usecase "学情分析与 AI 辅助" as T4
}
rectangle "管理端" {
  usecase "账号角色维护" as A1
  usecase "基础字典维护" as A2
  usecase "服务运行配置" as A3
}
学生 --> S1
学生 --> S2
学生 --> S3
学生 --> S4
教师 --> T1
教师 --> T2
教师 --> T3
教师 --> T4
管理员 --> A1
管理员 --> A2
管理员 --> A3
@enduml""",
        nodes=[
            Node("student", "学生", 70, 215, 110, 180, "actor"),
            Node("teacher", "教师", 70, 560, 110, 180, "actor"),
            Node("admin", "管理员", 70, 875, 110, 180, "actor"),
            Node("s1", "课程学习", 360, 150, 220, 90, "usecase", GREEN),
            Node("s2", "作业提交", 690, 150, 220, 90, "usecase", GREEN),
            Node("s3", "考试提交", 1020, 150, 220, 90, "usecase", GREEN),
            Node("s4", "成绩与预警查看", 1350, 150, 260, 90, "usecase", YELLOW),
            Node("t1", "课程班级维护", 360, 515, 250, 90, "usecase", BLUE),
            Node("t2", "作业发布批改", 700, 515, 250, 90, "usecase", BLUE),
            Node("t3", "考试发布阅卷", 1040, 515, 250, 90, "usecase", BLUE),
            Node("t4", "学情分析与 AI 辅助", 1370, 515, 280, 90, "usecase", PURPLE),
            Node("a1", "账号角色维护", 430, 875, 250, 90, "usecase", GRAY),
            Node("a2", "基础字典维护", 790, 875, 250, 90, "usecase", GRAY),
            Node("a3", "服务运行配置", 1150, 875, 250, 90, "usecase", GRAY),
        ],
        edges=[
            Edge("student", "s1"), Edge("student", "s2"), Edge("student", "s3"), Edge("student", "s4"),
            Edge("teacher", "t1"), Edge("teacher", "t2"), Edge("teacher", "t3"), Edge("teacher", "t4"),
            Edge("admin", "a1"), Edge("admin", "a2"), Edge("admin", "a3"),
        ],
        metadata=meta("分角色用户目标"),
    ))

    diagrams.append(Diagram(
        key="exam_assignment_class",
        title="图11 考试作业核心类图",
        uml_type="Class Diagram",
        filename=filename(1, "exam-assignment-class"),
        description=(
            "考试作业核心类图抽取 assignment-service 与 exam-service 的共同业务结构。"
            "作业和考试都围绕课程、班级发布范围、学生提交、教师批改、成绩事件展开，适合放在一张图中比较两条评价链路。"
        ),
        evidence=[
            "assignment-service-api 与 exam-service-api 中的发布、提交、批改 DTO",
            "ExamFinishedEvent、AssignmentGradedEvent 等评价结果事件",
        ],
        plantuml="""@startuml
skinparam classAttributeIconSize 0
class Course
class CourseClass
class Assignment
class AssignmentPublishScope
class AssignmentSubmission
class AssignmentGrade
class Exam
class ExamQuestion
class ExamSubmission
class ExamGrade
class LearningEvent
Course "1" -- "*" Assignment
Course "1" -- "*" Exam
CourseClass "1" -- "*" AssignmentPublishScope
Assignment "1" -- "*" AssignmentSubmission
AssignmentSubmission "1" -- "0..1" AssignmentGrade
Exam "1" -- "*" ExamQuestion
Exam "1" -- "*" ExamSubmission
ExamSubmission "1" -- "0..1" ExamGrade
AssignmentGrade ..> LearningEvent
ExamGrade ..> LearningEvent
@enduml""",
        nodes=[
            Node("course", "Course", 120, 170, 260, 118, fill=GREEN, members=["id", "name", "teacherId"]),
            Node("class", "CourseClass", 120, 390, 260, 118, fill=GREEN, members=["id", "className", "majorId"]),
            Node("assignment", "Assignment", 520, 170, 280, 136, fill=YELLOW, members=["title", "deadline", "scoreRule"]),
            Node("scope", "AssignmentPublishScope", 520, 390, 310, 136, fill=YELLOW, members=["assignmentId", "classId"]),
            Node("asub", "AssignmentSubmission", 900, 250, 320, 150, fill=ORANGE, members=["studentId", "content", "submittedAt", "status"]),
            Node("agrade", "AssignmentGrade", 1320, 250, 290, 136, fill=ORANGE, members=["score", "comment", "gradedBy"]),
            Node("exam", "Exam", 520, 670, 280, 136, fill=BLUE, members=["title", "startTime", "endTime"]),
            Node("question", "ExamQuestion", 900, 620, 300, 136, fill=BLUE, members=["type", "score", "answer"]),
            Node("esub", "ExamSubmission", 900, 820, 320, 150, fill=PURPLE, members=["answers", "timeTaken", "graded", "score"]),
            Node("egrade", "ExamGrade", 1320, 790, 290, 136, fill=PURPLE, members=["autoScore", "manualScore", "finalScore"]),
            Node("event", "LearningEvent", 1320, 560, 290, 118, fill=RED, members=["studentId", "courseId", "eventType"]),
        ],
        edges=[
            Edge("course", "assignment", "1..*"), Edge("class", "scope", "发布范围"), Edge("assignment", "scope", "1..*"),
            Edge("assignment", "asub", "提交"), Edge("asub", "agrade", "批改"), Edge("course", "exam", "1..*"),
            Edge("exam", "question", "包含题目"), Edge("exam", "esub", "考试提交"), Edge("esub", "egrade", "评分"),
            Edge("agrade", "event", "成绩事件", "dependency", True), Edge("egrade", "event", "成绩事件", "dependency", True),
        ],
        metadata=meta("考试与作业领域结构"),
    ))

    diagrams.append(Diagram(
        key="user_course_permission_class",
        title="图12 用户课程权限类图",
        uml_type="Class Diagram",
        filename=filename(2, "user-course-permission-class"),
        description=(
            "用户课程权限类图描述账号、角色、学生档案、教师授课、班级学生和课程之间的静态关系。"
            "该图用于说明 Gateway 与各服务鉴权后，业务服务如何根据课程、班级和角色关系判断数据访问范围。"
        ),
        evidence=[
            "auth-service 与 user-service 的用户、角色、认证凭证模型",
            "course-service 中 course、class、teacher-course、class-student 等关系表",
        ],
        plantuml="""@startuml
skinparam classAttributeIconSize 0
class User
class Role
class AuthCredential
class StudentProfile
class TeacherProfile
class Course
class TeacherCourse
class CourseClass
class ClassStudent
User "*" -- "*" Role
User "1" -- "0..1" AuthCredential
User "1" -- "0..1" StudentProfile
User "1" -- "0..1" TeacherProfile
TeacherProfile "1" -- "*" TeacherCourse
TeacherCourse "*" -- "1" Course
Course "1" -- "*" CourseClass
CourseClass "1" -- "*" ClassStudent
ClassStudent "*" -- "1" StudentProfile
@enduml""",
        nodes=[
            Node("user", "User", 120, 200, 260, 130, fill=BLUE, members=["id", "username", "realName", "status"]),
            Node("role", "Role", 500, 150, 240, 118, fill=GRAY, members=["code", "name", "permissions"]),
            Node("cred", "AuthCredential", 500, 330, 270, 130, fill=YELLOW, members=["passwordHash", "lastLogin", "enabled"]),
            Node("student", "StudentProfile", 880, 170, 280, 130, fill=GREEN, members=["studentNo", "majorId", "grade"]),
            Node("teacher", "TeacherProfile", 880, 380, 280, 130, fill=GREEN, members=["teacherNo", "department", "title"]),
            Node("tc", "TeacherCourse", 1240, 380, 280, 130, fill=PURPLE, members=["teacherId", "courseId", "semester"]),
            Node("course", "Course", 1240, 600, 280, 130, fill=PURPLE, members=["name", "credit", "semester"]),
            Node("class", "CourseClass", 880, 760, 280, 130, fill=ORANGE, members=["className", "majorId", "studentCount"]),
            Node("cs", "ClassStudent", 500, 760, 280, 130, fill=ORANGE, members=["classId", "studentId", "joinedAt"]),
        ],
        edges=[
            Edge("user", "role", "*..*"), Edge("user", "cred", "登录凭证"), Edge("user", "student", "学生档案"),
            Edge("user", "teacher", "教师档案"), Edge("teacher", "tc", "授课关系"), Edge("tc", "course", "授课课程"),
            Edge("course", "class", "开设班级"), Edge("class", "cs", "班级学生"), Edge("cs", "student", "学生归属"),
        ],
        metadata=meta("用户、角色与课程授权关系"),
    ))

    diagrams.append(Diagram(
        key="assignment_grading_object",
        title="图13 作业批改场景对象图",
        uml_type="Object Diagram",
        filename=filename(3, "assignment-grading-object"),
        description=(
            "作业批改对象图展示一次教师批改作业时的对象快照。"
            "教师、作业、提交记录、评分结果、成绩事件和通知对象之间形成一条从评价到反馈的链路。"
        ),
        evidence=[
            "assignment-service 的 AssignmentSubmissionRecord、GradeCommand 相关模型",
            "notification-service 中作业批改完成通知的持久化职责",
        ],
        plantuml="""@startuml
object "teacher:User" as teacher
object "assignment:Assignment" as assignment
object "submission:AssignmentSubmission" as sub
object "grade:AssignmentGrade" as grade
object "event:AssignmentGradedEvent" as event
object "notice:Notification" as notice
teacher --> grade
assignment --> sub
sub --> grade
grade --> event
event --> notice
@enduml""",
        nodes=[
            Node("teacher", "teacher:User", 120, 190, 300, 140, fill=BLUE, members=["id = 20011", "role = TEACHER", "realName = 李老师"]),
            Node("assignment", "assignment:Assignment", 560, 190, 330, 150, fill=GREEN, members=["title = Spring 作业", "courseId = 12", "deadline = 2026-06-10"]),
            Node("sub", "submission:AssignmentSubmission", 1000, 190, 380, 170, fill=YELLOW, members=["studentId = 10023", "status = SUBMITTED", "contentUrl = /uploads/a1.md"]),
            Node("grade", "grade:AssignmentGrade", 1000, 520, 380, 160, fill=ORANGE, members=["score = 88", "comment = 完成度较好", "gradedBy = 20011"]),
            Node("event", "event:AssignmentGradedEvent", 560, 760, 360, 150, fill=PURPLE, members=["aggregateId = A-501-S10023", "eventType = GRADED"]),
            Node("notice", "notice:Notification", 120, 760, 330, 150, fill=RED, members=["type = ASSIGNMENT_GRADED", "read = false", "receiver = 10023"]),
        ],
        edges=[
            Edge("teacher", "grade", "批改人"), Edge("assignment", "sub", "所属作业"), Edge("sub", "grade", "产生评分"),
            Edge("grade", "event", "发布成绩事件"), Edge("event", "notice", "生成通知"),
        ],
        metadata=meta("作业批改运行时对象"),
    ))

    login_lifelines = ["用户", "登录页面", "Gateway", "AuthController", "AuthService", "UserService", "Redis", "JWT"]
    login_messages = [
        {"from": "用户", "to": "登录页面", "label": "输入账号密码/验证码"},
        {"from": "登录页面", "to": "Gateway", "label": "POST /api/auth/login"},
        {"from": "Gateway", "to": "AuthController", "label": "转发登录请求"},
        {"from": "AuthController", "to": "AuthService", "label": "authenticate(command)"},
        {"from": "AuthService", "to": "UserService", "label": "查询用户、角色、状态"},
        {"from": "UserService", "to": "AuthService", "label": "返回用户资料", "return": True, "dashed": True},
        {"from": "AuthService", "to": "Redis", "label": "校验验证码/写 refresh JTI"},
        {"from": "AuthService", "to": "JWT", "label": "签发 access/refresh token"},
        {"from": "AuthService", "to": "AuthController", "label": "返回登录结果", "return": True, "dashed": True},
        {"from": "AuthController", "to": "Gateway", "label": "ResponseResult<TokenPair>", "return": True, "dashed": True},
        {"from": "Gateway", "to": "登录页面", "label": "保存 token 与角色", "return": True, "dashed": True},
    ]
    nodes, edges = sequence_nodes_and_edges(login_lifelines, login_messages)
    diagrams.append(Diagram(
        key="login_sequence",
        title="图14 登录认证顺序图",
        uml_type="Sequence Diagram",
        filename=filename(4, "login-sequence"),
        description=(
            "登录认证顺序图展示用户登录时 Gateway、auth-service、user-service、Redis 和 JWT 支撑组件之间的消息顺序。"
            "该图突出认证链路中的验证码校验、用户状态查询、角色装配和双 token 签发。"
        ),
        evidence=[
            "auth-service 的 AuthController、JwtTokenService、RsaKeySupport",
            "common 安全上下文与 Gateway JWT 校验逻辑",
        ],
        plantuml="""@startuml
actor 用户
participant 登录页面
participant Gateway
participant AuthController
participant AuthService
participant UserService
database Redis
participant JWT
用户 -> 登录页面 : 输入账号密码/验证码
登录页面 -> Gateway : POST /api/auth/login
Gateway -> AuthController : 转发登录请求
AuthController -> AuthService : authenticate(command)
AuthService -> UserService : 查询用户、角色、状态
UserService --> AuthService : 返回用户资料
AuthService -> Redis : 校验验证码/写 refresh JTI
AuthService -> JWT : 签发 token pair
AuthService --> AuthController : 登录结果
AuthController --> Gateway : ResponseResult
Gateway --> 登录页面 : 保存 token 与角色
@enduml""",
        nodes=nodes,
        edges=edges,
        metadata={**meta("登录认证消息顺序"), "lifelines": login_lifelines, "messages": login_messages},
    ))

    grade_lifelines = ["教师", "批改页面", "Gateway", "AssignmentController", "AssignmentService", "sc_assignment", "RabbitMQ", "Notification Service"]
    grade_messages = [
        {"from": "教师", "to": "批改页面", "label": "查看提交并填写评分"},
        {"from": "批改页面", "to": "Gateway", "label": "PUT /api/teacher/assignments/{id}/submissions/{sid}/grade"},
        {"from": "Gateway", "to": "AssignmentController", "label": "鉴权后转发"},
        {"from": "AssignmentController", "to": "AssignmentService", "label": "gradeSubmission(command)"},
        {"from": "AssignmentService", "to": "sc_assignment", "label": "校验教师授课和提交状态"},
        {"from": "AssignmentService", "to": "sc_assignment", "label": "保存分数、评语、批改人"},
        {"from": "AssignmentService", "to": "sc_assignment", "label": "写 AssignmentGradedEvent"},
        {"from": "AssignmentService", "to": "AssignmentController", "label": "返回批改结果", "return": True, "dashed": True},
        {"from": "sc_assignment", "to": "RabbitMQ", "label": "OutboxRelay 发布成绩事件"},
        {"from": "RabbitMQ", "to": "Notification Service", "label": "创建作业批改通知"},
    ]
    nodes, edges = sequence_nodes_and_edges(grade_lifelines, grade_messages)
    diagrams.append(Diagram(
        key="assignment_grading_sequence",
        title="图15 作业批改顺序图",
        uml_type="Sequence Diagram",
        filename=filename(5, "assignment-grading-sequence"),
        description=(
            "作业批改顺序图描述教师为学生提交记录打分的同步与异步链路。"
            "同步部分完成授课权限校验、提交状态校验和评分保存；异步部分通过 outbox 与 RabbitMQ 通知学生批改结果。"
        ),
        evidence=[
            "assignment-service 中教师批改接口与提交记录持久化逻辑",
            "common outbox 与 notification-service 的通知落库职责",
        ],
        plantuml="""@startuml
actor 教师
participant 批改页面
participant Gateway
participant AssignmentController
participant AssignmentService
database sc_assignment
queue RabbitMQ
participant NotificationService
教师 -> 批改页面 : 查看提交并填写评分
批改页面 -> Gateway : PUT grade
Gateway -> AssignmentController : 鉴权后转发
AssignmentController -> AssignmentService : gradeSubmission(command)
AssignmentService -> sc_assignment : 校验教师授课和提交状态
AssignmentService -> sc_assignment : 保存分数/评语/批改人
AssignmentService -> sc_assignment : 写 AssignmentGradedEvent
AssignmentService --> AssignmentController : 批改结果
sc_assignment -> RabbitMQ : 发布成绩事件
RabbitMQ -> NotificationService : 创建作业批改通知
@enduml""",
        nodes=nodes,
        edges=edges,
        metadata={**meta("作业批改消息顺序"), "lifelines": grade_lifelines, "messages": grade_messages},
    ))

    ai_lifelines = ["教师", "AI 组卷页面", "Gateway", "AiController", "AiService", "CourseService", "ExamService", "外部大模型"]
    ai_messages = [
        {"from": "教师", "to": "AI 组卷页面", "label": "选择课程、知识点、题型和难度"},
        {"from": "AI 组卷页面", "to": "Gateway", "label": "POST /api/ai/papers/generate"},
        {"from": "Gateway", "to": "AiController", "label": "鉴权后转发"},
        {"from": "AiController", "to": "AiService", "label": "generatePaper(command)"},
        {"from": "AiService", "to": "CourseService", "label": "查询课程知识点和班级上下文"},
        {"from": "AiService", "to": "外部大模型", "label": "提交 prompt 与约束"},
        {"from": "外部大模型", "to": "AiService", "label": "返回试卷草稿", "return": True, "dashed": True},
        {"from": "AiService", "to": "ExamService", "label": "保存试卷/题目草稿"},
        {"from": "AiService", "to": "AiController", "label": "返回生成记录", "return": True, "dashed": True},
        {"from": "AiController", "to": "AI 组卷页面", "label": "展示草稿并允许编辑", "return": True, "dashed": True},
    ]
    nodes, edges = sequence_nodes_and_edges(ai_lifelines, ai_messages)
    diagrams.append(Diagram(
        key="ai_paper_sequence",
        title="图16 AI 组卷顺序图",
        uml_type="Sequence Diagram",
        filename=filename(6, "ai-paper-sequence"),
        description=(
            "AI 组卷顺序图展示教师使用 AI 生成试卷草稿的交互过程。"
            "ai-service 汇总课程知识点、题型难度和教师约束，调用外部大模型生成草稿，再交由 exam-service 保存并供教师二次编辑。"
        ),
        evidence=[
            "ai-service-api 中题目生成、组卷生成请求 DTO",
            "course-service-api 与 exam-service-api 为 AI 生成提供上下文和落库能力",
        ],
        plantuml="""@startuml
actor 教师
participant AI组卷页面
participant Gateway
participant AiController
participant AiService
participant CourseService
participant ExamService
participant 外部大模型
教师 -> AI组卷页面 : 选择课程/知识点/题型/难度
AI组卷页面 -> Gateway : POST /api/ai/papers/generate
Gateway -> AiController : 鉴权后转发
AiController -> AiService : generatePaper(command)
AiService -> CourseService : 查询知识点上下文
AiService -> 外部大模型 : prompt + constraints
外部大模型 --> AiService : 试卷草稿
AiService -> ExamService : 保存试卷/题目草稿
AiService --> AiController : 生成记录
AiController --> AI组卷页面 : 展示草稿
@enduml""",
        nodes=nodes,
        edges=edges,
        metadata={**meta("AI 组卷消息顺序"), "lifelines": ai_lifelines, "messages": ai_messages},
    ))

    diagrams.append(Diagram(
        key="analysis_notification_communication",
        title="图17 学情分析通知协作图",
        uml_type="Communication Diagram",
        filename=filename(7, "analysis-notification-communication"),
        description=(
            "学情分析通知协作图说明考试和作业评价事件如何驱动分析服务更新 read model，并在达到阈值时通知服务生成预警通知。"
            "该图突出事件总线、分析规则、预警实体和消息投递之间的对象协作。"
        ),
        evidence=[
            "analysis-service 维护成绩趋势、知识点掌握度和 EarlyWarning",
            "notification-service 消费预警事件并创建站内信/通知记录",
        ],
        plantuml="""@startuml
object RabbitMQ
object AnalysisConsumer
object MasteryModel
object WarningRule
object EarlyWarning
object NotificationConsumer
object Notification
RabbitMQ -- AnalysisConsumer : 1 消费成绩事件
AnalysisConsumer -- MasteryModel : 2 更新掌握度
AnalysisConsumer -- WarningRule : 3 计算预警
WarningRule -- EarlyWarning : 4 保存预警
EarlyWarning -- RabbitMQ : 5 发布预警事件
RabbitMQ -- NotificationConsumer : 6 消费预警事件
NotificationConsumer -- Notification : 7 创建通知
@enduml""",
        nodes=[
            Node("mq", "RabbitMQ", 120, 250, 250, 92, fill=PURPLE),
            Node("consumer", "AnalysisConsumer", 520, 160, 280, 92, fill=PURPLE),
            Node("mastery", "MasteryModel", 960, 160, 260, 92, fill=GREEN),
            Node("rule", "WarningRule", 960, 390, 260, 92, fill=YELLOW),
            Node("warning", "EarlyWarning", 520, 570, 280, 92, fill=RED),
            Node("noticeConsumer", "NotificationConsumer", 520, 820, 320, 92, fill=ORANGE),
            Node("notice", "Notification", 980, 820, 260, 92, fill=ORANGE),
        ],
        edges=[
            Edge("mq", "consumer", "1: 消费成绩事件"), Edge("consumer", "mastery", "2: 更新掌握度"),
            Edge("consumer", "rule", "3: 计算预警"), Edge("rule", "warning", "4: 保存预警"),
            Edge("warning", "mq", "5: 发布预警事件"), Edge("mq", "noticeConsumer", "6: 消费预警事件"),
            Edge("noticeConsumer", "notice", "7: 创建通知"),
        ],
        metadata=meta("分析预警与通知协作"),
    ))

    diagrams.append(Diagram(
        key="ai_generation_communication",
        title="图18 AI 生成协作图",
        uml_type="Communication Diagram",
        filename=filename(8, "ai-generation-communication"),
        description=(
            "AI 生成协作图从对象链接角度说明教师请求、AI 编排服务、课程上下文、外部模型、生成记录和业务草稿之间的协作关系。"
            "该图用于补充 AI 组卷顺序图，强调参与对象和责任分配。"
        ),
        evidence=[
            "ai-service-api 中 AiGenerateQuestionRequest、AiGeneratePaperRequest 等契约",
            "ai-service 对外部 AI 能力的编排以及 exam-service 草稿落库职责",
        ],
        plantuml="""@startuml
object 教师
object AiGenerationRequest
object AiService
object CourseContext
object PromptBuilder
object ExternalModel
object AiGenerationRecord
object ExamDraft
教师 -- AiGenerationRequest : 1 输入约束
AiGenerationRequest -- AiService : 2 生成命令
AiService -- CourseContext : 3 查询上下文
AiService -- PromptBuilder : 4 构造 prompt
PromptBuilder -- ExternalModel : 5 调用模型
ExternalModel -- AiGenerationRecord : 6 返回结果
AiGenerationRecord -- ExamDraft : 7 保存草稿
@enduml""",
        nodes=[
            Node("teacher", "教师", 120, 250, 220, 92, fill=BLUE),
            Node("request", "AiGenerationRequest", 480, 140, 300, 92, fill=YELLOW),
            Node("service", "AiService", 880, 250, 250, 92, fill=GREEN),
            Node("context", "CourseContext", 1250, 130, 280, 92, fill=PURPLE),
            Node("prompt", "PromptBuilder", 1250, 390, 280, 92, fill=PURPLE),
            Node("model", "ExternalModel", 880, 640, 280, 92, fill=RED),
            Node("record", "AiGenerationRecord", 480, 760, 320, 92, fill=ORANGE),
            Node("draft", "ExamDraft", 120, 640, 250, 92, fill=BLUE),
        ],
        edges=[
            Edge("teacher", "request", "1: 输入约束"), Edge("request", "service", "2: 生成命令"),
            Edge("service", "context", "3: 查询上下文"), Edge("service", "prompt", "4: 构造 prompt"),
            Edge("prompt", "model", "5: 调用模型"), Edge("model", "record", "6: 返回结果"),
            Edge("record", "draft", "7: 保存草稿"),
        ],
        metadata=meta("AI 生成对象协作"),
    ))

    diagrams.append(Diagram(
        key="assignment_submission_state",
        title="图19 作业提交状态图",
        uml_type="Statechart Diagram",
        filename=filename(9, "assignment-submission-state"),
        description=(
            "作业提交状态图选择 AssignmentSubmission 作为生命周期对象。"
            "它从未提交开始，经历草稿、已提交、待批改、已批改、需修改和归档等状态，反映学生提交与教师反馈的闭环。"
        ),
        evidence=[
            "assignment-service 中作业提交、教师批改、重新提交与成绩发布路径",
            "AssignmentSubmissionRecord 中提交内容、批改状态、得分与评语字段",
        ],
        plantuml="""@startuml
[*] --> 未提交
未提交 --> 草稿 : 保存草稿
草稿 --> 已提交 : 提交作业
未提交 --> 已提交 : 直接提交
已提交 --> 待批改 : 截止前有效
待批改 --> 已批改 : 教师评分
已批改 --> 需修改 : 教师退回
需修改 --> 已提交 : 学生重新提交
已批改 --> 已归档 : 发布成绩
已归档 --> [*]
@enduml""",
        nodes=[
            Node("start", "", 150, 230, 34, 34, "initial"),
            Node("none", "未提交", 270, 205, 210, 84, "state", BLUE),
            Node("draft", "草稿", 590, 120, 210, 84, "state", YELLOW),
            Node("submitted", "已提交", 590, 320, 210, 84, "state", GREEN),
            Node("pending", "待批改", 910, 320, 210, 84, "state", ORANGE),
            Node("graded", "已批改", 1230, 320, 210, 84, "state", GREEN),
            Node("revise", "需修改", 910, 545, 210, 84, "state", RED),
            Node("archived", "已归档", 1540, 320, 210, 84, "state", PURPLE),
            Node("end", "", 1660, 555, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "none"), Edge("none", "draft", "保存草稿"), Edge("draft", "submitted", "提交作业"),
            Edge("none", "submitted", "直接提交"), Edge("submitted", "pending", "截止前有效"),
            Edge("pending", "graded", "教师评分"), Edge("graded", "revise", "教师退回"),
            Edge("revise", "submitted", "重新提交"), Edge("graded", "archived", "发布成绩"), Edge("archived", "end"),
        ],
        metadata=meta("作业提交生命周期"),
    ))

    diagrams.append(Diagram(
        key="notification_state",
        title="图20 通知消息状态图",
        uml_type="Statechart Diagram",
        filename=filename(10, "notification-state"),
        description=(
            "通知消息状态图描述 NotificationEntity 从创建、待投递、已投递、已读、处理完成到归档的生命周期。"
            "对于预警类通知，还存在用户处理或教师确认后关闭的状态迁移。"
        ),
        evidence=[
            "notification-service 保存站内信、阅读状态和通知投递记录",
            "analysis-service 预警事件触发通知生成并等待用户处理",
        ],
        plantuml="""@startuml
[*] --> 已创建
已创建 --> 待投递 : 保存通知
待投递 --> 投递失败 : MQ/通道异常
投递失败 --> 待投递 : 重试
待投递 --> 已投递 : 通道成功
已投递 --> 已读 : 用户打开
已读 --> 待处理 : 预警类通知
待处理 --> 已处理 : 教师/学生确认
已读 --> 已归档 : 普通通知
已处理 --> 已归档 : 关闭预警
已归档 --> [*]
@enduml""",
        nodes=[
            Node("start", "", 130, 260, 34, 34, "initial"),
            Node("created", "已创建", 260, 235, 210, 84, "state", BLUE),
            Node("pending", "待投递", 570, 235, 210, 84, "state", YELLOW),
            Node("failed", "投递失败", 570, 470, 210, 84, "state", RED),
            Node("delivered", "已投递", 880, 235, 210, 84, "state", GREEN),
            Node("read", "已读", 1190, 235, 210, 84, "state", GREEN),
            Node("handling", "待处理", 1190, 470, 210, 84, "state", ORANGE),
            Node("handled", "已处理", 1490, 470, 210, 84, "state", PURPLE),
            Node("archived", "已归档", 1490, 235, 210, 84, "state", PURPLE),
            Node("end", "", 1660, 250, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "created"), Edge("created", "pending", "保存通知"), Edge("pending", "failed", "通道异常"),
            Edge("failed", "pending", "重试"), Edge("pending", "delivered", "投递成功"),
            Edge("delivered", "read", "用户打开"), Edge("read", "handling", "预警类通知"),
            Edge("handling", "handled", "确认处理"), Edge("read", "archived", "普通通知"),
            Edge("handled", "archived", "关闭预警"), Edge("archived", "end"),
        ],
        metadata=meta("通知消息生命周期"),
    ))

    diagrams.append(Diagram(
        key="login_activity",
        title="图21 登录认证活动图",
        uml_type="Activity Diagram",
        filename=filename(11, "login-activity"),
        description=(
            "登录认证活动图从控制流角度描述账号密码、验证码、用户状态、角色选择和 token 签发。"
            "它适合说明认证服务如何处理失败分支，并与 Gateway 后续鉴权形成闭环。"
        ),
        evidence=[
            "auth-service 登录、刷新 token、角色切换相关接口",
            "Redis 存储验证码、refresh token JTI 和限流辅助信息",
        ],
        plantuml="""@startuml
start
:打开登录页;
:输入账号、密码、验证码;
if (验证码有效?) then (是)
  :查询用户和角色;
else (否)
  :返回验证码错误;
  stop
endif
if (账号启用且密码正确?) then (是)
  :选择默认角色;
else (否)
  :返回认证失败;
  stop
endif
:签发 access token;
:写入 refresh JTI;
:返回用户资料和菜单权限;
stop
@enduml""",
        nodes=[
            Node("start", "", 900, 145, 34, 34, "initial"),
            Node("open", "打开登录页", 760, 235, 340, 78, "activity", BLUE),
            Node("input", "输入账号、密码、验证码", 720, 365, 420, 78, "activity", BLUE),
            Node("captcha", "验证码有效?", 805, 515, 250, 118, "decision", YELLOW),
            Node("captchaFail", "返回验证码错误", 1220, 535, 310, 78, "activity", RED),
            Node("query", "查询用户和角色", 760, 705, 340, 78, "activity", GREEN),
            Node("valid", "账号启用且密码正确?", 790, 855, 280, 126, "decision", YELLOW),
            Node("authFail", "返回认证失败", 1220, 880, 310, 78, "activity", RED),
            Node("role", "选择默认角色", 760, 1045, 340, 78, "activity", GREEN),
            Node("token", "签发 access token", 760, 1170, 340, 78, "activity", PURPLE),
            Node("refresh", "写入 refresh JTI", 760, 1295, 340, 78, "activity", PURPLE),
            Node("menu", "返回用户资料和菜单权限", 720, 1420, 420, 78, "activity", ORANGE),
            Node("end", "", 900, 1570, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "open"), Edge("open", "input"), Edge("input", "captcha"),
            Edge("captcha", "captchaFail", "否"), Edge("captcha", "query", "是"),
            Edge("query", "valid"), Edge("valid", "authFail", "否"), Edge("valid", "role", "是"),
            Edge("role", "token"), Edge("token", "refresh"), Edge("refresh", "menu"), Edge("menu", "end"),
        ],
        metadata=meta("登录认证控制流", (1900, 1750)),
    ))

    diagrams.append(Diagram(
        key="assignment_activity",
        title="图22 作业发布批改活动图",
        uml_type="Activity Diagram",
        filename=filename(12, "assignment-activity"),
        description=(
            "作业发布批改活动图覆盖教师发布作业、学生提交、教师批改、成绩发布和通知反馈的完整闭环。"
            "该图与作业批改顺序图互补，强调流程分支和业务状态变化。"
        ),
        evidence=[
            "assignment-service 的作业发布、班级发布范围、学生提交和教师批改接口",
            "notification-service 在批改完成后创建作业结果通知",
        ],
        plantuml="""@startuml
start
:教师选择课程和班级;
:填写作业要求与截止时间;
:发布作业;
:学生查看作业;
if (截止前提交?) then (是)
  :保存提交记录;
else (否)
  :标记逾期或拒绝提交;
endif
:教师查看提交列表;
if (需要退回修改?) then (是)
  :退回并填写修改意见;
  :学生重新提交;
else (否)
  :评分并填写评语;
endif
:发布成绩事件;
:通知学生查看结果;
stop
@enduml""",
        nodes=[
            Node("start", "", 900, 145, 34, 34, "initial"),
            Node("select", "教师选择课程和班级", 730, 235, 420, 78, "activity", BLUE),
            Node("fill", "填写作业要求与截止时间", 705, 365, 470, 78, "activity", BLUE),
            Node("publish", "发布作业", 780, 495, 320, 78, "activity", GREEN),
            Node("view", "学生查看作业", 780, 625, 320, 78, "activity", GREEN),
            Node("deadline", "截止前提交?", 815, 765, 250, 118, "decision", YELLOW),
            Node("save", "保存提交记录", 455, 920, 320, 78, "activity", GREEN),
            Node("late", "标记逾期或拒绝提交", 1080, 920, 380, 78, "activity", RED),
            Node("list", "教师查看提交列表", 730, 1065, 420, 78, "activity", BLUE),
            Node("revise", "需要退回修改?", 815, 1210, 250, 118, "decision", YELLOW),
            Node("return", "退回并填写修改意见", 430, 1380, 390, 78, "activity", ORANGE),
            Node("resubmit", "学生重新提交", 430, 1510, 390, 78, "activity", ORANGE),
            Node("grade", "评分并填写评语", 1080, 1380, 360, 78, "activity", GREEN),
            Node("event", "发布成绩事件", 780, 1640, 320, 78, "activity", PURPLE),
            Node("notice", "通知学生查看结果", 730, 1765, 420, 78, "activity", ORANGE),
            Node("end", "", 910, 1910, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "select"), Edge("select", "fill"), Edge("fill", "publish"), Edge("publish", "view"),
            Edge("view", "deadline"), Edge("deadline", "save", "是"), Edge("deadline", "late", "否"),
            Edge("save", "list"), Edge("late", "list"), Edge("list", "revise"),
            Edge("revise", "return", "是"), Edge("return", "resubmit"), Edge("resubmit", "list"),
            Edge("revise", "grade", "否"), Edge("grade", "event"), Edge("event", "notice"), Edge("notice", "end"),
        ],
        metadata=meta("作业发布批改流程", (1900, 2050)),
    ))

    diagrams.append(Diagram(
        key="ai_analysis_activity",
        title="图23 AI 分析推荐活动图",
        uml_type="Activity Diagram",
        filename=filename(13, "ai-analysis-activity"),
        description=(
            "AI 分析推荐活动图展示平台如何基于成绩、知识点掌握度和教师约束生成学习建议或教学资源。"
            "流程包含数据聚合、风险识别、prompt 构造、模型调用、结果审核和通知反馈。"
        ),
        evidence=[
            "analysis-service 提供成绩趋势、知识点掌握和预警 read model",
            "ai-service-api 支持学习建议、题目和组卷生成请求",
        ],
        plantuml="""@startuml
start
:收集考试/作业成绩事件;
:更新成绩趋势和知识点掌握;
if (达到预警条件?) then (是)
  :生成预警对象;
endif
:教师或学生发起 AI 建议请求;
:汇总课程、知识点和薄弱项;
:构造 prompt 和安全约束;
:调用外部模型;
if (生成结果可用?) then (是)
  :保存生成记录;
  :返回学习建议/题目草稿;
else (否)
  :返回降级模板建议;
endif
:通知用户查看结果;
stop
@enduml""",
        nodes=[
            Node("start", "", 900, 145, 34, 34, "initial"),
            Node("collect", "收集考试/作业成绩事件", 700, 235, 480, 78, "activity", PURPLE),
            Node("update", "更新成绩趋势和知识点掌握", 690, 365, 500, 78, "activity", PURPLE),
            Node("warn", "达到预警条件?", 815, 515, 250, 118, "decision", YELLOW),
            Node("warning", "生成预警对象", 1180, 535, 320, 78, "activity", RED),
            Node("request", "教师或学生发起 AI 建议请求", 665, 700, 550, 78, "activity", BLUE),
            Node("context", "汇总课程、知识点和薄弱项", 665, 830, 550, 78, "activity", GREEN),
            Node("prompt", "构造 prompt 和安全约束", 700, 960, 480, 78, "activity", GREEN),
            Node("model", "调用外部模型", 760, 1090, 360, 78, "activity", ORANGE),
            Node("usable", "生成结果可用?", 815, 1235, 250, 118, "decision", YELLOW),
            Node("save", "保存生成记录并返回建议/草稿", 425, 1410, 500, 78, "activity", GREEN),
            Node("fallback", "返回降级模板建议", 1080, 1410, 380, 78, "activity", RED),
            Node("notice", "通知用户查看结果", 730, 1600, 420, 78, "activity", ORANGE),
            Node("end", "", 910, 1750, 42, 42, "final"),
        ],
        edges=[
            Edge("start", "collect"), Edge("collect", "update"), Edge("update", "warn"),
            Edge("warn", "warning", "是"), Edge("warn", "request", "否"), Edge("warning", "request"),
            Edge("request", "context"), Edge("context", "prompt"), Edge("prompt", "model"),
            Edge("model", "usable"), Edge("usable", "save", "是"), Edge("usable", "fallback", "否"),
            Edge("save", "notice"), Edge("fallback", "notice"), Edge("notice", "end"),
        ],
        metadata=meta("AI 分析推荐流程", (1900, 1900)),
    ))

    diagrams.append(Diagram(
        key="core_microservice_component",
        title="图24 核心微服务组件细化图",
        uml_type="Component Diagram",
        filename=filename(14, "core-microservice-component"),
        description=(
            "核心微服务组件细化图聚焦业务服务之间的 API 与事件关系。"
            "Auth/User/Course 提供身份、资料和教学基础数据，Assignment/Exam 产生评价事实，Analysis/Notification/AI 围绕评价事实提供分析、反馈和智能生成能力。"
        ),
        evidence=[
            "docs/service-boundary-and-monolith-shrink-plan.md 中各服务边界与迁移计划",
            "common-events 模块定义评价、预警和通知相关领域事件",
        ],
        plantuml="""@startuml
component Gateway
component AuthService
component UserService
component CourseService
component AssignmentService
component ExamService
component AnalysisService
component NotificationService
component AiService
component CommonEvents
queue RabbitMQ
database MySQL
Gateway --> AuthService
Gateway --> UserService
Gateway --> CourseService
Gateway --> AssignmentService
Gateway --> ExamService
Gateway --> AnalysisService
Gateway --> AiService
AuthService ..> UserService
AssignmentService ..> CourseService
ExamService ..> CourseService
AiService ..> CourseService
AiService ..> ExamService
AssignmentService ..> CommonEvents
ExamService ..> CommonEvents
AnalysisService ..> CommonEvents
CommonEvents ..> RabbitMQ
AnalysisService --> NotificationService
AuthService --> MySQL
UserService --> MySQL
CourseService --> MySQL
AssignmentService --> MySQL
ExamService --> MySQL
AnalysisService --> MySQL
NotificationService --> MySQL
AiService --> MySQL
@enduml""",
        nodes=[
            Node("gateway", "Gateway", 80, 210, 260, 105, "component", GREEN),
            Node("auth", "AuthService", 455, 140, 260, 105, "component", BLUE),
            Node("user", "UserService", 455, 320, 260, 105, "component", BLUE),
            Node("course", "CourseService", 825, 230, 280, 105, "component", BLUE),
            Node("assignment", "AssignmentService", 1200, 120, 320, 105, "component", YELLOW),
            Node("exam", "ExamService", 1200, 310, 320, 105, "component", YELLOW),
            Node("ai", "AiService", 1200, 500, 320, 105, "component", ORANGE),
            Node("events", "CommonEvents", 825, 620, 300, 105, "component", PURPLE),
            Node("rabbit", "RabbitMQ", 455, 690, 260, 105, "component", PURPLE),
            Node("analysis", "AnalysisService", 1200, 720, 320, 105, "component", RED),
            Node("notice", "NotificationService", 1200, 930, 320, 105, "component", ORANGE),
            Node("mysql", "MySQL sc_* schemas", 610, 980, 420, 105, "component", GREEN),
        ],
        edges=[
            Edge("gateway", "auth"), Edge("gateway", "user"), Edge("gateway", "course"),
            Edge("gateway", "assignment"), Edge("gateway", "exam"), Edge("gateway", "ai"),
            Edge("auth", "user", "用户角色", "dependency", True), Edge("assignment", "course", "课程班级", "dependency", True),
            Edge("exam", "course", "课程题目", "dependency", True), Edge("ai", "course", "知识点上下文", "dependency", True),
            Edge("ai", "exam", "试卷草稿", "dependency", True), Edge("assignment", "events", "作业事件", "dependency", True),
            Edge("exam", "events", "考试事件", "dependency", True), Edge("events", "rabbit", "发布/消费", "dependency", True),
            Edge("rabbit", "analysis", "评价事件"), Edge("analysis", "notice", "预警通知"),
            Edge("auth", "mysql"), Edge("user", "mysql"), Edge("course", "mysql"), Edge("assignment", "mysql"),
            Edge("exam", "mysql"), Edge("ai", "mysql"), Edge("analysis", "mysql"), Edge("notice", "mysql"),
        ],
        metadata=meta("核心微服务依赖"),
    ))

    return diagrams


def module_plantuml(module_name: str, title: str, nodes: list[Node], edges: list[Edge]) -> str:
    lines = [
        "@startuml",
        "skinparam classAttributeIconSize 0",
        f'title {title}',
        f'package "{module_name}" {{',
    ]
    for node in nodes:
        if node.kind in {"initial", "final", "actor"}:
            continue
        stereotype = f" <<{node.stereotype}>>" if node.stereotype else ""
        safe_name = node.id.replace("-", "_")
        lines.append(f'  class "{node.label.replace(chr(10), " ")}" as {safe_name}{stereotype} {{')
        for member in node.members:
            lines.append(f"    {member}")
        lines.append("  }")
    lines.append("}")
    for edge in edges:
        src = edge.source.replace("-", "_")
        dst = edge.target.replace("-", "_")
        arrow_text = "..>" if edge.dashed or edge.kind == "dependency" else "-->"
        label = f" : {edge.label}" if edge.label else ""
        lines.append(f"{src} {arrow_text} {dst}{label}")
    lines.append("@enduml")
    return "\n".join(lines)


def create_module_diagram(index: int,
                          key: str,
                          title: str,
                          module_name: str,
                          description: str,
                          evidence: list[str],
                          components: list[dict[str, Any]],
                          relations: list[tuple[str, str, str, str | None]]) -> Diagram:
    lane_x = [90, 500, 910, 1320, 1730]
    lane_y = [190, 540, 890, 1240]
    fill_by_layer = {
        "controller": BLUE,
        "service": GREEN,
        "repository": GRAY,
        "entity": YELLOW,
        "domain": YELLOW,
        "api": PURPLE,
        "infra": ORANGE,
        "external": RED,
        "config": GRAY,
    }
    nodes: list[Node] = []
    for i, comp in enumerate(components):
        col = comp.get("col", i % 5)
        row = comp.get("row", i // 5)
        layer = comp.get("layer", "service")
        nodes.append(Node(
            comp["id"],
            comp["label"],
            lane_x[col],
            lane_y[row],
            comp.get("w", 320),
            comp.get("h", 240),
            fill=comp.get("fill", fill_by_layer.get(layer, WHITE)),
            stereotype=comp.get("stereotype", layer),
            members=comp.get("members", []),
        ))
    component_pos = {comp["id"]: (comp.get("col", i % 5), comp.get("row", i // 5))
                     for i, comp in enumerate(components)}
    edges: list[Edge] = []
    for src, dst, label, kind in relations:
        src_pos = component_pos.get(src)
        dst_pos = component_pos.get(dst)
        if src_pos and dst_pos:
            src_col, src_row = src_pos
            dst_col, dst_row = dst_pos
            if dst_col < src_col:
                continue
            if src_row != dst_row:
                continue
            if dst_col - src_col > 1:
                continue
        # Module detail diagrams are explained in report text; labels on dense edges make the figures unreadable.
        edges.append(Edge(src, dst, "", kind or "association", dashed=(kind == "dependency")))
    return Diagram(
        key=key,
        title=f"图{index} {title}",
        uml_type="Module Detail Class Diagram",
        filename=f"{index:02d}-{key}.png",
        description=description,
        evidence=evidence,
        plantuml=module_plantuml(module_name, title, nodes, edges),
        nodes=nodes,
        edges=edges,
        metadata={"module": module_name, "detail": True},
    )


def make_missing_maven_module_specs() -> list[dict[str, Any]]:
    return [
        {
            "key": "module-parent-pom",
            "title": "Parent POM 聚合与依赖治理模块详细图",
            "module": "parent-pom",
            "description": "parent-pom 模块承担 Maven 多模块工程的统一版本、依赖和插件治理职责。根 pom.xml 负责声明模块顺序，parent-pom 负责 Spring Boot、Spring Cloud、MyBatis、测试、镜像构建等公共构建约束，保证各服务使用一致的技术栈。",
            "evidence": ["pom.xml modules", "parent-pom/pom.xml"],
            "components": [
                {"id": "root-pom", "label": "Root pom.xml", "layer": "config", "col": 0, "row": 0, "members": ["packaging=pom", "23 modules", "构建入口"]},
                {"id": "parent-pom", "label": "parent-pom", "layer": "config", "col": 1, "row": 0, "members": ["公共 parent", "版本收敛"]},
                {"id": "dependency-mgmt", "label": "Dependency Management", "layer": "config", "col": 2, "row": 0, "members": ["Spring Boot", "Spring Cloud", "MyBatis", "Testcontainers"]},
                {"id": "plugin-mgmt", "label": "Plugin Management", "layer": "config", "col": 2, "row": 1, "members": ["maven-surefire", "spring-boot-maven", "docker/coverage"]},
                {"id": "service-modules", "label": "Business Service Modules", "layer": "external", "col": 3, "row": 0, "members": ["auth/user/course", "assignment/exam/analysis", "notification/ai"]},
                {"id": "api-modules", "label": "*-service-api Modules", "layer": "external", "col": 3, "row": 1, "members": ["DTO", "Feign", "fallback"]},
                {"id": "infra-modules", "label": "Infrastructure Modules", "layer": "external", "col": 4, "row": 0, "members": ["common", "common-events", "gateway", "registry-server"]},
                {"id": "legacy-modules", "label": "Migration Modules", "layer": "external", "col": 4, "row": 1, "members": ["major_assignment", "legacy-adapter"]},
            ],
            "relations": [
                ("root-pom", "parent-pom", "继承/聚合", None), ("parent-pom", "dependency-mgmt", "管理依赖版本", None),
                ("parent-pom", "plugin-mgmt", "管理插件", None), ("service-modules", "parent-pom", "继承构建规范", "dependency"),
                ("api-modules", "parent-pom", "继承构建规范", "dependency"), ("infra-modules", "parent-pom", "继承构建规范", "dependency"),
                ("legacy-modules", "parent-pom", "继承构建规范", "dependency"), ("service-modules", "api-modules", "依赖契约", "dependency"),
            ],
        },
        {
            "key": "module-major-assignment",
            "title": "Major Assignment 过渡单体模块详细图",
            "module": "major_assignment",
            "description": "major_assignment 模块是迁移期保留的过渡单体，包含静态学生/教师页面、兼容 Controller、Service、MyBatis Mapper、实体 DTO、会话与网关 JWT 桥接过滤器。它支撑本地验证和旧页面兼容，同时通过客户端逐步接入拆分后的微服务。",
            "evidence": ["major_assignment/src/main/java", "major_assignment/src/main/resources/static", "docs/service-boundary-and-monolith-shrink-plan.md"],
            "components": [
                {"id": "app", "label": "MajorAssignmentApplication", "layer": "service", "col": 0, "row": 0, "members": ["Spring Boot", "过渡单体入口"]},
                {"id": "static", "label": "Static Pages", "layer": "infra", "col": 0, "row": 1, "members": ["student-*.html", "teacher-*.html", "admin pages"]},
                {"id": "controllers", "label": "Compatibility Controllers", "layer": "controller", "col": 1, "row": 0, "members": ["Auth/Course/Assignment", "Exam/Analysis/Notification", "Student/Teacher Dashboard"]},
                {"id": "session", "label": "Session & Gateway Bridge", "layer": "infra", "col": 1, "row": 1, "members": ["GatewayJwtBridgeFilter", "MultiRoleSessionFilter", "TraceIdFilter"]},
                {"id": "services", "label": "Legacy Services", "layer": "service", "col": 2, "row": 0, "members": ["User/Course/Assignment", "Exam/Analysis/Notification", "AI"]},
                {"id": "clients", "label": "Microservice Clients", "layer": "api", "col": 2, "row": 1, "members": ["AuthServiceClient", "UserServiceProfileClient"]},
                {"id": "mappers", "label": "MyBatis Mappers", "layer": "repository", "col": 3, "row": 0, "members": ["UserMapper", "CourseMapper", "ExamMapper", "AssignmentMapper"]},
                {"id": "entities", "label": "Entities & DTOs", "layer": "entity", "col": 4, "row": 0, "members": ["User/Course/Exam", "Assignment/Warning", "ResponseResult"]},
                {"id": "db", "label": "major_assignment DB", "layer": "external", "col": 4, "row": 1, "members": ["MySQL", "Redis cache/session"]},
            ],
            "relations": [
                ("static", "controllers", "页面请求", "dependency"), ("app", "controllers", "注册控制器", None),
                ("controllers", "session", "登录/角色上下文", None), ("controllers", "services", "兼容业务调用", None),
                ("services", "clients", "迁移期远程调用", "dependency"), ("services", "mappers", "旧表读写", None),
                ("mappers", "entities", "对象映射", None), ("mappers", "db", "持久化", "dependency"),
                ("session", "db", "会话/缓存", "dependency"),
            ],
        },
        {
            "key": "module-auth-service-api",
            "title": "Auth Service API 契约模块详细图",
            "module": "auth-service-api",
            "description": "auth-service-api 模块定义认证服务对其它模块开放的内部契约。AuthFeignClient 负责令牌解析和认证上下文获取，AuthContextDTO 与 TokenIntrospectionDTO 形成调用方可依赖的最小认证数据模型。",
            "evidence": ["auth-service-api/src/main/java"],
            "components": [
                {"id": "client", "label": "AuthFeignClient", "layer": "api", "col": 0, "row": 0, "members": ["@FeignClient auth-service", "/internal/auth", "introspect()", "context()"]},
                {"id": "context", "label": "AuthContextDTO", "layer": "api", "col": 1, "row": 0, "members": ["userId", "username", "roles", "activeRole"]},
                {"id": "token", "label": "TokenIntrospectionDTO", "layer": "api", "col": 2, "row": 0, "members": ["valid", "subject", "expiresAt", "claims"]},
                {"id": "consumer", "label": "Gateway / Business Services", "layer": "external", "col": 3, "row": 0, "members": ["鉴权过滤", "角色校验", "用户上下文"]},
                {"id": "provider", "label": "auth-service Internal API", "layer": "external", "col": 4, "row": 0, "members": ["令牌校验", "认证事实源"]},
            ],
            "relations": [
                ("consumer", "client", "依赖 Feign 契约", "dependency"), ("client", "provider", "HTTP 调用", "dependency"),
                ("client", "context", "返回认证上下文", None), ("client", "token", "返回令牌解析", None),
            ],
        },
        {
            "key": "module-user-service-api",
            "title": "User Service API 契约模块详细图",
            "module": "user-service-api",
            "description": "user-service-api 模块封装用户与学生资料的跨服务查询契约。它通过 UserFeignClient 暴露用户画像、角色、学生资料和更新请求 DTO，并提供 fallback factory 让调用方在远程失败时获得可控错误。",
            "evidence": ["user-service-api/src/main/java"],
            "components": [
                {"id": "client", "label": "UserFeignClient", "layer": "api", "col": 0, "row": 0, "members": ["@FeignClient user-service", "/internal/users", "profile()", "roles()"]},
                {"id": "fallback", "label": "UserFeignClientFallbackFactory", "layer": "infra", "col": 0, "row": 1, "members": ["远程失败降级", "统一异常"]},
                {"id": "profile", "label": "UserProfileDTO", "layer": "api", "col": 1, "row": 0, "members": ["id", "username", "realName", "avatar"]},
                {"id": "student", "label": "StudentProfileDTO", "layer": "api", "col": 2, "row": 0, "members": ["studentNo", "major", "className"]},
                {"id": "roles", "label": "UserRolesDTO", "layer": "api", "col": 3, "row": 0, "members": ["roles", "activeRole"]},
                {"id": "update", "label": "Update Profile DTOs", "layer": "api", "col": 2, "row": 1, "members": ["UpdateUserProfileDTO", "UpdateStudentProfileDTO"]},
                {"id": "consumers", "label": "Auth/Course/Analysis", "layer": "external", "col": 4, "row": 0, "members": ["角色查询", "学生维度", "展示资料"]},
            ],
            "relations": [
                ("consumers", "client", "读取用户资料", "dependency"), ("client", "fallback", "失败降级", "dependency"),
                ("client", "profile", "用户画像", None), ("client", "student", "学生资料", None),
                ("client", "roles", "角色列表", None), ("client", "update", "更新命令", None),
            ],
        },
        {
            "key": "module-course-service-api",
            "title": "Course Service API 契约模块详细图",
            "module": "course-service-api",
            "description": "course-service-api 模块定义课程、班级、专业、教师授课和分页查询的跨服务契约。作业、考试和分析服务通过 CourseFeignClient 校验课程班级范围，避免直接读取课程服务数据库。",
            "evidence": ["course-service-api/src/main/java"],
            "components": [
                {"id": "client", "label": "CourseFeignClient", "layer": "api", "col": 0, "row": 0, "members": ["@FeignClient course-service", "/internal/courses", "course()", "classCourses()"]},
                {"id": "fallback", "label": "CourseFeignClientFallbackFactory", "layer": "infra", "col": 0, "row": 1, "members": ["远程失败降级"]},
                {"id": "course", "label": "CourseDTO", "layer": "api", "col": 1, "row": 0, "members": ["courseId", "courseName", "teacherId"]},
                {"id": "class", "label": "Class/Teacher DTOs", "layer": "api", "col": 2, "row": 0, "members": ["ClassCourseDTO", "TeacherClassDTO", "MajorDTO"]},
                {"id": "requests", "label": "Upsert/Assignment DTOs", "layer": "api", "col": 3, "row": 0, "members": ["CourseUpsertRequestDTO", "ClassUpsertRequestDTO", "CourseAssignmentRequestDTO"]},
                {"id": "page", "label": "PageResultDTO<T>", "layer": "api", "col": 2, "row": 1, "members": ["records", "total", "page", "size"]},
                {"id": "consumers", "label": "Assignment/Exam/Analysis", "layer": "external", "col": 4, "row": 0, "members": ["范围校验", "课程维度查询"]},
            ],
            "relations": [
                ("consumers", "client", "课程范围查询", "dependency"), ("client", "fallback", "失败降级", "dependency"),
                ("client", "course", "课程 DTO", None), ("client", "class", "班级/专业 DTO", None),
                ("client", "requests", "管理命令 DTO", None), ("client", "page", "分页响应", None),
            ],
        },
        {
            "key": "module-assignment-service-api",
            "title": "Assignment Service API 契约模块详细图",
            "module": "assignment-service-api",
            "description": "assignment-service-api 模块定义作业、提交、成绩和教师批改命令的跨服务契约。分析服务和通知服务依赖这些 DTO 获取作业维度数据，教师端则通过请求 DTO 维护发布和批改命令。",
            "evidence": ["assignment-service-api/src/main/java"],
            "components": [
                {"id": "client", "label": "AssignmentFeignClient", "layer": "api", "col": 0, "row": 0, "members": ["@FeignClient assignment-service", "/internal/assignments"]},
                {"id": "assignment", "label": "AssignmentDTO", "layer": "api", "col": 1, "row": 0, "members": ["assignmentId", "courseId", "deadline", "status"]},
                {"id": "submission", "label": "AssignmentSubmissionDTO", "layer": "api", "col": 2, "row": 0, "members": ["submissionId", "studentId", "content", "score"]},
                {"id": "score", "label": "AssignmentStudentScoreDTO", "layer": "api", "col": 3, "row": 0, "members": ["studentId", "score", "graded"]},
                {"id": "submit", "label": "AssignmentSubmitRequestDTO", "layer": "api", "col": 1, "row": 1, "members": ["content", "attachments"]},
                {"id": "teacher", "label": "Teacher Assignment Requests", "layer": "api", "col": 2, "row": 1, "members": ["TeacherAssignmentUpsertRequestDTO", "TeacherAssignmentGradeRequestDTO"]},
                {"id": "consumers", "label": "Analysis/Notification/Teacher UI", "layer": "external", "col": 4, "row": 0, "members": ["成绩分析", "提交提醒", "批改"]},
            ],
            "relations": [
                ("consumers", "client", "读取作业/成绩", "dependency"), ("client", "assignment", "作业 DTO", None),
                ("client", "submission", "提交 DTO", None), ("client", "score", "成绩 DTO", None),
                ("client", "submit", "学生提交命令", None), ("client", "teacher", "教师命令", None),
            ],
        },
        {
            "key": "module-exam-service-api",
            "title": "Exam Service API 契约模块详细图",
            "module": "exam-service-api",
            "description": "exam-service-api 模块定义考试发布、考试提交、学生成绩和教师阅卷命令的跨服务契约。Gateway、分析和通知链路可通过这些 DTO 理解考试完成结果，而无需访问 exam-service 内部实体。",
            "evidence": ["exam-service-api/src/main/java"],
            "components": [
                {"id": "client", "label": "ExamFeignClient", "layer": "api", "col": 0, "row": 0, "members": ["@FeignClient exam-service", "/internal/exams"]},
                {"id": "exam", "label": "ExamDTO", "layer": "api", "col": 1, "row": 0, "members": ["examId", "courseId", "startTime", "endTime"]},
                {"id": "submission", "label": "ExamSubmissionDTO", "layer": "api", "col": 2, "row": 0, "members": ["submissionId", "studentId", "score", "graded"]},
                {"id": "submit", "label": "ExamSubmitRequestDTO", "layer": "api", "col": 1, "row": 1, "members": ["answers", "duration"]},
                {"id": "score", "label": "Student Score DTOs", "layer": "api", "col": 3, "row": 0, "members": ["StudentScoreDTO", "StudentScoreListItemDTO"]},
                {"id": "teacher", "label": "Teacher Exam Requests", "layer": "api", "col": 2, "row": 1, "members": ["TeacherExamUpsertRequestDTO", "TeacherExamGradeRequestDTO"]},
                {"id": "consumers", "label": "Analysis/Notification/Gateway", "layer": "external", "col": 4, "row": 0, "members": ["考试结果", "阅卷状态", "提交路由"]},
            ],
            "relations": [
                ("consumers", "client", "读取考试结果", "dependency"), ("client", "exam", "考试 DTO", None),
                ("client", "submission", "提交 DTO", None), ("client", "submit", "学生提交命令", None),
                ("client", "score", "成绩 DTO", None), ("client", "teacher", "教师命令", None),
            ],
        },
        {
            "key": "module-analysis-service-api",
            "title": "Analysis Service API 契约模块详细图",
            "module": "analysis-service-api",
            "description": "analysis-service-api 模块定义学情分析 read model 的最小跨服务 DTO，包括成绩趋势和知识点掌握度。其它模块和前端只依赖这些稳定输出，不需要理解分析服务内部聚合与事件处理细节。",
            "evidence": ["analysis-service-api/src/main/java"],
            "components": [
                {"id": "api", "label": "analysis-service-api", "layer": "api", "col": 0, "row": 0, "members": ["read model contract", "DTO only"]},
                {"id": "trend", "label": "ScoreTrendDTO", "layer": "api", "col": 1, "row": 0, "members": ["studentId", "courseId", "score", "assessmentType", "occurredAt"]},
                {"id": "mastery", "label": "KnowledgeMasteryDTO", "layer": "api", "col": 2, "row": 0, "members": ["studentId", "knowledgePointId", "masteryLevel", "source"]},
                {"id": "provider", "label": "analysis-service Query API", "layer": "external", "col": 3, "row": 0, "members": ["trend()", "mastery()", "dashboard()"]},
                {"id": "consumers", "label": "Teacher/Student UI", "layer": "external", "col": 4, "row": 0, "members": ["趋势图", "知识点掌握", "预警展示"]},
            ],
            "relations": [
                ("api", "trend", "定义输出", None), ("api", "mastery", "定义输出", None),
                ("provider", "trend", "返回趋势", None), ("provider", "mastery", "返回掌握度", None),
                ("consumers", "provider", "查询分析", "dependency"),
            ],
        },
        {
            "key": "module-notification-service-api",
            "title": "Notification Service API 契约模块详细图",
            "module": "notification-service-api",
            "description": "notification-service-api 模块承载通知服务对外命令契约。当前核心 DTO 是 TeacherSendNotificationRequestDTO，用于教师主动发送站内通知；系统事件生成通知则由 common-events 事件链路承载。",
            "evidence": ["notification-service-api/src/main/java"],
            "components": [
                {"id": "api", "label": "notification-service-api", "layer": "api", "col": 0, "row": 0, "members": ["notification contract", "command DTO"]},
                {"id": "teacher-send", "label": "TeacherSendNotificationRequestDTO", "layer": "api", "col": 1, "row": 0, "members": ["studentIds", "title", "content", "type"]},
                {"id": "provider", "label": "notification-service Command API", "layer": "external", "col": 2, "row": 0, "members": ["send()", "markRead()", "list()"]},
                {"id": "teacher-ui", "label": "Teacher UI", "layer": "external", "col": 3, "row": 0, "members": ["主动通知"]},
                {"id": "event-path", "label": "common-events", "layer": "external", "col": 4, "row": 0, "members": ["ExamFinished", "AssignmentGraded", "WarningRaised"]},
            ],
            "relations": [
                ("teacher-ui", "teacher-send", "填写通知命令", "dependency"), ("teacher-send", "provider", "发送请求", None),
                ("event-path", "provider", "系统通知另走事件链", "dependency"), ("api", "teacher-send", "定义 DTO", None),
            ],
        },
        {
            "key": "module-ai-service-api",
            "title": "AI Service API 契约模块详细图",
            "module": "ai-service-api",
            "description": "ai-service-api 模块定义 AI 生成请求的输入契约，包括题目生成、组卷生成和学习建议生成。它只表达请求参数，生成记录和业务落库分别由 ai-service 与目标业务服务负责。",
            "evidence": ["ai-service-api/src/main/java"],
            "components": [
                {"id": "api", "label": "ai-service-api", "layer": "api", "col": 0, "row": 0, "members": ["AI request contract", "DTO only"]},
                {"id": "questions", "label": "GenerateQuestionsRequestDTO", "layer": "api", "col": 1, "row": 0, "members": ["courseId", "knowledgePointIds", "questionType", "count"]},
                {"id": "exam", "label": "GenerateExamRequestDTO", "layer": "api", "col": 2, "row": 0, "members": ["courseId", "duration", "difficulty", "questionPlan"]},
                {"id": "suggestion", "label": "LearningSuggestionRequestDTO", "layer": "api", "col": 3, "row": 0, "members": ["studentId", "courseId", "weakPoints"]},
                {"id": "provider", "label": "ai-service Generation API", "layer": "external", "col": 4, "row": 0, "members": ["generateQuestions()", "generateExam()", "learningSuggestion()"]},
            ],
            "relations": [
                ("api", "questions", "定义题目请求", None), ("api", "exam", "定义组卷请求", None),
                ("api", "suggestion", "定义建议请求", None), ("questions", "provider", "生成题目", "dependency"),
                ("exam", "provider", "生成试卷", "dependency"), ("suggestion", "provider", "生成建议", "dependency"),
            ],
        },
    ]


def make_module_detail_diagrams(start_index: int) -> list[Diagram]:
    specs: list[dict[str, Any]] = [
        {
            "key": "module-gateway",
            "title": "Gateway 模块详细图",
            "module": "gateway",
            "description": "Gateway 模块详细图展开统一入口的路由、安全、限流、熔断与前端能力兜底职责。它不持有业务数据，而是通过 Eureka 发现服务，通过 Redis 维护限流与 JWT 黑名单状态，并把请求转发到各业务微服务或 legacy-monolith。",
            "evidence": ["gateway/src/main/resources/application.yml", "docs/service-boundary-and-monolith-shrink-plan.md"],
            "components": [
                {"id": "gw-route", "label": "RouteLocator\napplication.yml routes", "layer": "config", "col": 0, "row": 0, "members": ["auth-route", "student-exam-submit-route", "ai-route"]},
                {"id": "gw-security", "label": "Gateway JWT Security", "layer": "service", "col": 1, "row": 0, "members": ["校验 JWT", "透传用户上下文", "白名单路径"]},
                {"id": "gw-rate", "label": "Rate Limit Filter", "layer": "service", "col": 2, "row": 0, "members": ["routeId 分桶", "Redis 计数", "Retry-After"]},
                {"id": "gw-fallback", "label": "FallbackController", "layer": "controller", "col": 3, "row": 0, "members": ["服务熔断响应", "统一错误结构"]},
                {"id": "gw-edge", "label": "FrontendCapabilityEdgeController", "layer": "controller", "col": 4, "row": 0, "members": ["前端能力声明", "兼容提示"]},
                {"id": "eureka", "label": "Eureka Registry", "layer": "external", "col": 1, "row": 1, "members": ["服务发现"]},
                {"id": "redis", "label": "Redis", "layer": "external", "col": 2, "row": 1, "members": ["限流桶", "JTI blacklist"]},
                {"id": "services", "label": "Business Services", "layer": "external", "col": 3, "row": 1, "members": ["auth/user/course", "assignment/exam/analysis", "notification/ai"]},
                {"id": "legacy", "label": "legacy-monolith", "layer": "external", "col": 4, "row": 1, "members": ["迁移期兜底"]},
            ],
            "relations": [
                ("gw-route", "gw-security", "匹配路由后鉴权", None), ("gw-security", "gw-rate", "通过后限流", None),
                ("gw-rate", "services", "负载均衡转发", "dependency"), ("gw-rate", "redis", "读写限流状态", None),
                ("gw-security", "redis", "查询 JTI 黑名单", None), ("gw-route", "eureka", "lb:// 服务发现", "dependency"),
                ("gw-route", "legacy", "legacy-route 兜底", "dependency"), ("gw-fallback", "services", "熔断 fallback", "dependency"),
            ],
        },
        {
            "key": "module-auth",
            "title": "Auth 认证模块详细图",
            "module": "auth-service",
            "description": "Auth 模块负责登录、验证码、JWT 签发、刷新、注销和角色切换。认证事实保存在 sc_auth 与 Redis，用户画像和角色业务属性通过 user-service API 查询，保证认证与用户资料边界清晰。",
            "evidence": ["auth-service/src/main/java", "docs/data-ownership.md 中 Auth_Service 数据归属"],
            "components": [
                {"id": "auth-controller", "label": "AuthController", "layer": "controller", "col": 0, "row": 0, "members": ["login()", "refresh()", "switchRole()", "logout()"]},
                {"id": "captcha-controller", "label": "PublicCaptchaCompatibilityController", "layer": "controller", "col": 0, "row": 1, "members": ["captcha()"]},
                {"id": "auth-service", "label": "AuthApplicationService", "layer": "service", "col": 1, "row": 0, "members": ["authenticate()", "refreshToken()", "changePassword()"]},
                {"id": "jwt-service", "label": "JwtTokenService", "layer": "service", "col": 2, "row": 0, "members": ["issueTokenPair()", "parseClaims()", "blacklistJti()"]},
                {"id": "captcha-service", "label": "CaptchaService", "layer": "service", "col": 2, "row": 1, "members": ["createImage()", "validate()"]},
                {"id": "credential-repo", "label": "AuthCredentialRepository", "layer": "repository", "col": 3, "row": 0, "members": ["findByUsername()", "updatePassword()"]},
                {"id": "credential", "label": "AuthCredentialEntity", "layer": "entity", "col": 4, "row": 0, "members": ["userId", "username", "passwordHash", "enabled"]},
                {"id": "redis", "label": "Redis", "layer": "external", "col": 3, "row": 1, "members": ["Refresh token", "Captcha", "JTI blacklist"]},
                {"id": "user-api", "label": "UserFeignClient", "layer": "api", "col": 4, "row": 1, "members": ["roles()", "profile()"]},
            ],
            "relations": [
                ("auth-controller", "auth-service", "认证命令", None), ("captcha-controller", "captcha-service", "验证码", None),
                ("auth-service", "jwt-service", "签发/解析令牌", None), ("auth-service", "credential-repo", "读取凭证", None),
                ("credential-repo", "credential", "JPA 映射", None), ("jwt-service", "redis", "刷新令牌/JTI", None),
                ("captcha-service", "redis", "验证码状态", None), ("auth-service", "user-api", "查询角色", "dependency"),
            ],
        },
        {
            "key": "module-user",
            "title": "User 用户模块详细图",
            "module": "user-service",
            "description": "User 模块拥有用户档案、角色和学生资料更新职责。它向认证、课程、考试和分析等服务提供只读用户信息，避免其它服务直接访问用户表。",
            "evidence": ["user-service/src/main/java", "user-service-api/src/main/java"],
            "components": [
                {"id": "user-controller", "label": "UserController", "layer": "controller", "col": 0, "row": 0, "members": ["me()", "updateProfile()"]},
                {"id": "student-controller", "label": "StudentCompatibilityController", "layer": "controller", "col": 0, "row": 1, "members": ["studentProfile()", "uploadAvatar()"]},
                {"id": "internal-controller", "label": "UserInternalController", "layer": "controller", "col": 1, "row": 1, "members": ["roles()", "lookup()"]},
                {"id": "user-service", "label": "UserApplicationService", "layer": "service", "col": 1, "row": 0, "members": ["getProfile()", "updateProfile()", "studentProfile()"]},
                {"id": "user-repo", "label": "UserRepository", "layer": "repository", "col": 2, "row": 0, "members": ["findById()", "findByUsername()", "updateProfile()"]},
                {"id": "jpa-user", "label": "JpaUserRepository", "layer": "repository", "col": 3, "row": 0, "members": ["toRecord()", "roles()"]},
                {"id": "user-entity", "label": "UserEntity", "layer": "entity", "col": 4, "row": 0, "members": ["id", "username", "realName", "majorId"]},
                {"id": "role-entity", "label": "RoleEntity", "layer": "entity", "col": 4, "row": 1, "members": ["id", "roleName", "description"]},
                {"id": "user-api", "label": "User service API DTO", "layer": "api", "col": 2, "row": 1, "members": ["UserProfileDTO", "StudentProfileDTO", "UserRolesDTO"]},
            ],
            "relations": [
                ("user-controller", "user-service", "用户资料", None), ("student-controller", "user-service", "学生兼容接口", None),
                ("internal-controller", "user-service", "内部查询", None), ("user-service", "user-repo", "仓储接口", None),
                ("user-repo", "jpa-user", "实现", "dependency"), ("jpa-user", "user-entity", "映射", None),
                ("jpa-user", "role-entity", "角色查询", None), ("user-service", "user-api", "DTO 输出", "dependency"),
            ],
        },
        {
            "key": "module-course",
            "title": "Course 课程模块详细图",
            "module": "course-service",
            "description": "Course 模块拥有课程、班级、专业、班级学生、教师授课关系和知识点主数据。它为作业、考试和分析服务提供课程与班级范围校验，是教学组织结构的事实源。",
            "evidence": ["course-service/src/main/java", "docs/data-ownership.md 中 Course_Service 数据归属"],
            "components": [
                {"id": "teacher-course", "label": "TeacherCourseAdminController", "layer": "controller", "col": 0, "row": 0, "members": ["课程 CRUD", "班级分配", "学生加入"]},
                {"id": "student-course", "label": "StudentCourseController", "layer": "controller", "col": 0, "row": 1, "members": ["学生课程列表", "课程详情"]},
                {"id": "kp-controller", "label": "TeacherKnowledgePointCompatibilityController", "layer": "controller", "col": 1, "row": 1, "members": ["知识点 CRUD"]},
                {"id": "course-service", "label": "CourseApplicationService", "layer": "service", "col": 1, "row": 0, "members": ["upsertCourse()", "assignClass()", "studentCourses()"]},
                {"id": "kp-service", "label": "TeacherKnowledgePointService", "layer": "service", "col": 2, "row": 1, "members": ["listByCourse()", "savePoint()"]},
                {"id": "course-repo", "label": "CourseRepository", "layer": "repository", "col": 2, "row": 0, "members": ["findCourses()", "findClasses()", "assignStudent()"]},
                {"id": "course-entity", "label": "CourseEntity", "layer": "entity", "col": 3, "row": 0, "members": ["courseName", "teacherId", "semester"]},
                {"id": "class-entity", "label": "CourseClassEntity", "layer": "entity", "col": 4, "row": 0, "members": ["className", "majorId", "studentCount"]},
                {"id": "kp-entity", "label": "TeacherKnowledgePointEntity", "layer": "entity", "col": 3, "row": 1, "members": ["courseId", "name", "difficulty"]},
                {"id": "major-entity", "label": "MajorEntity", "layer": "entity", "col": 4, "row": 1, "members": ["majorName"]},
            ],
            "relations": [
                ("teacher-course", "course-service", "教师管理", None), ("student-course", "course-service", "学生查询", None),
                ("kp-controller", "kp-service", "知识点管理", None), ("course-service", "course-repo", "课程仓储", None),
                ("kp-service", "kp-entity", "知识点持久化", None), ("course-repo", "course-entity", "课程", None),
                ("course-repo", "class-entity", "班级", None), ("course-repo", "major-entity", "专业", None),
            ],
        },
        {
            "key": "module-assignment",
            "title": "Assignment 作业模块详细图",
            "module": "assignment-service",
            "description": "Assignment 模块负责作业发布、班级发布范围、学生提交、教师批改和作业知识点关联。提交和批改会通过 outbox 生成事件，供分析和通知服务消费。",
            "evidence": ["assignment-service/src/main/java", "common-events assignment 事件"],
            "components": [
                {"id": "teacher-assignment", "label": "TeacherAssignmentController", "layer": "controller", "col": 0, "row": 0, "members": ["发布/编辑/删除作业", "查询作业"]},
                {"id": "student-assignment", "label": "StudentAssignmentController", "layer": "controller", "col": 0, "row": 1, "members": ["作业列表", "submitAssignment()"]},
                {"id": "teacher-submission", "label": "TeacherSubmissionController", "layer": "controller", "col": 1, "row": 1, "members": ["查看提交", "批改"]},
                {"id": "assignment-service", "label": "AssignmentApplicationService", "layer": "service", "col": 1, "row": 0, "members": ["submitAssignment()", "queryStudentList()"]},
                {"id": "cmd-service", "label": "TeacherAssignmentCommandService", "layer": "service", "col": 2, "row": 0, "members": ["publish()", "update()", "delete()"]},
                {"id": "query-service", "label": "TeacherAssignmentQueryService", "layer": "service", "col": 2, "row": 1, "members": ["teacherList()", "submissionList()"]},
                {"id": "assignment-repo", "label": "AssignmentRepository", "layer": "repository", "col": 3, "row": 0, "members": ["insertAssignment()", "insertSubmission()", "updateSubmission()"]},
                {"id": "assignment", "label": "AssignmentRecord", "layer": "domain", "col": 4, "row": 0, "members": ["courseId", "classId", "deadline", "status"]},
                {"id": "submission", "label": "AssignmentSubmissionRecord", "layer": "domain", "col": 4, "row": 1, "members": ["studentId", "content", "graded", "score", "isLate"]},
                {"id": "outbox", "label": "OutboxEventRepository", "layer": "infra", "col": 3, "row": 1, "members": ["AssignmentSubmittedEvent", "AssignmentGradedEvent"]},
            ],
            "relations": [
                ("teacher-assignment", "cmd-service", "发布命令", None), ("teacher-assignment", "query-service", "教师查询", None),
                ("student-assignment", "assignment-service", "学生提交", None), ("teacher-submission", "assignment-service", "批改提交", None),
                ("assignment-service", "assignment-repo", "读写提交", None), ("cmd-service", "assignment-repo", "读写作业", None),
                ("query-service", "assignment-repo", "查询", None), ("assignment-repo", "assignment", "作业聚合", None),
                ("assignment-repo", "submission", "提交记录", None), ("assignment-service", "outbox", "提交/批改事件", "dependency"),
            ],
        },
        {
            "key": "module-exam",
            "title": "Exam 考试模块详细图",
            "module": "exam-service",
            "description": "Exam 模块负责考试发布、题目、学生考试提交、自动阅卷、教师复核和考试知识点关联。考试完成事件是分析和通知链路的重要输入。",
            "evidence": ["exam-service/src/main/java", "ExamApplicationService persistExamFinishedEvent"],
            "components": [
                {"id": "teacher-exam", "label": "TeacherExamController", "layer": "controller", "col": 0, "row": 0, "members": ["考试 CRUD", "教师阅卷"]},
                {"id": "student-exam", "label": "StudentExamController", "layer": "controller", "col": 0, "row": 1, "members": ["考试列表", "submitExam()", "scores()"]},
                {"id": "exam-service", "label": "ExamApplicationService", "layer": "service", "col": 1, "row": 0, "members": ["submitExam()", "autoGradeIfPossible()", "gradeTeacherSubmission()"]},
                {"id": "kp-service", "label": "ExamKnowledgePointService", "layer": "service", "col": 1, "row": 1, "members": ["bindKnowledgePoints()"]},
                {"id": "exam-repo", "label": "ExamRepository", "layer": "repository", "col": 2, "row": 0, "members": ["findExam()", "findQuestions()", "insertSubmission()"]},
                {"id": "exam", "label": "ExamRecord", "layer": "domain", "col": 3, "row": 0, "members": ["courseId", "classId", "startTime", "endTime"]},
                {"id": "question", "label": "ExamQuestionRecord", "layer": "domain", "col": 4, "row": 0, "members": ["type", "answer", "score"]},
                {"id": "submission", "label": "ExamSubmissionRecord", "layer": "domain", "col": 3, "row": 1, "members": ["studentId", "timeTaken", "graded", "score"]},
                {"id": "outbox", "label": "OutboxEventRepository", "layer": "infra", "col": 4, "row": 1, "members": ["ExamFinishedEvent"]},
            ],
            "relations": [
                ("teacher-exam", "exam-service", "教师管理", None), ("student-exam", "exam-service", "学生提交", None),
                ("kp-service", "exam-repo", "知识点绑定", None), ("exam-service", "exam-repo", "考试读写", None),
                ("exam-repo", "exam", "考试", None), ("exam-repo", "question", "试题", None),
                ("exam-repo", "submission", "提交记录", None), ("exam-service", "outbox", "完成事件", "dependency"),
            ],
        },
        {
            "key": "module-analysis",
            "title": "Analysis 学情分析模块详细图",
            "module": "analysis-service",
            "description": "Analysis 模块消费作业和考试事件，生成成绩趋势、知识点掌握度、学情预警和分析触发任务。它是典型 read model 服务，不直接改写作业、考试或课程事实表。",
            "evidence": ["analysis-service/src/main/java", "docs/data-ownership.md 中 Analysis_Service 数据归属"],
            "components": [
                {"id": "query-controller", "label": "AnalysisQueryController", "layer": "controller", "col": 0, "row": 0, "members": ["scoreTrend()", "mastery()", "dashboard()"]},
                {"id": "warning-controller", "label": "EarlyWarningCompatibilityController", "layer": "controller", "col": 0, "row": 1, "members": ["warnings()", "resolve()"]},
                {"id": "trigger-controller", "label": "AnalysisTriggerCompatibilityController", "layer": "controller", "col": 1, "row": 1, "members": ["manualAnalyze()"]},
                {"id": "query-service", "label": "AnalysisQueryService", "layer": "service", "col": 1, "row": 0, "members": ["queryTrend()", "queryMastery()", "teacherDashboard()"]},
                {"id": "trigger-service", "label": "AnalysisTriggerCompatibilityService", "layer": "service", "col": 2, "row": 1, "members": ["createTriggerJob()", "runAnalysis()"]},
                {"id": "assignment-handler", "label": "AssignmentSubmittedAnalysisHandler", "layer": "service", "col": 2, "row": 0, "members": ["handleSubmitted()", "handleGraded()"]},
                {"id": "exam-handler", "label": "ExamFinishedAnalysisHandler", "layer": "service", "col": 3, "row": 0, "members": ["handleExamFinished()"]},
                {"id": "analysis-repo", "label": "AnalysisRepository", "layer": "repository", "col": 3, "row": 1, "members": ["upsertScoreTrend()", "upsertKnowledgeMastery()"]},
                {"id": "trend", "label": "ScoreTrendEntity", "layer": "entity", "col": 4, "row": 0, "members": ["studentId", "courseId", "score", "assessmentType"]},
                {"id": "warning", "label": "EarlyWarningEntity", "layer": "entity", "col": 4, "row": 1, "members": ["warningLevel", "resolved", "triggerDate"]},
            ],
            "relations": [
                ("query-controller", "query-service", "查询 read model", None), ("warning-controller", "query-service", "预警查询", None),
                ("trigger-controller", "trigger-service", "手动触发", None), ("assignment-handler", "analysis-repo", "作业事件聚合", None),
                ("exam-handler", "analysis-repo", "考试事件聚合", None), ("query-service", "analysis-repo", "查询", None),
                ("analysis-repo", "trend", "趋势", None), ("analysis-repo", "warning", "预警", None),
                ("trigger-service", "analysis-repo", "回填任务", None),
            ],
        },
        {
            "key": "module-notification",
            "title": "Notification 通知模块详细图",
            "module": "notification-service",
            "description": "Notification 模块保存站内信、通知投递记录和阅读状态。它可由教师主动发送，也可消费作业、考试、预警事件生成系统通知。",
            "evidence": ["notification-service/src/main/java", "common-events notification/warning 事件"],
            "components": [
                {"id": "notification-controller", "label": "NotificationController", "layer": "controller", "col": 0, "row": 0, "members": ["list()", "markRead()", "teacherSend()"]},
                {"id": "cmd-service", "label": "NotificationCommandService", "layer": "service", "col": 1, "row": 0, "members": ["sendToStudent()", "markRead()"]},
                {"id": "query-service", "label": "NotificationQueryService", "layer": "service", "col": 1, "row": 1, "members": ["studentList()", "unreadCount()"]},
                {"id": "submitted-handler", "label": "AssignmentSubmittedNotificationHandler", "layer": "service", "col": 2, "row": 0, "members": ["作业提交通知"]},
                {"id": "warning-handler", "label": "EarlyWarningRaisedNotificationHandler", "layer": "service", "col": 2, "row": 1, "members": ["预警通知"]},
                {"id": "exam-handler", "label": "ExamFinishedNotificationHandler", "layer": "service", "col": 3, "row": 0, "members": ["考试结果通知"]},
                {"id": "notification-repo", "label": "NotificationRepository", "layer": "repository", "col": 3, "row": 1, "members": ["save()", "findByStudent()", "markRead()"]},
                {"id": "notification-entity", "label": "NotificationEntity", "layer": "entity", "col": 4, "row": 0, "members": ["studentId", "teacherId", "type", "title", "read"]},
                {"id": "rabbit", "label": "RabbitMQ Events", "layer": "external", "col": 4, "row": 1, "members": ["AssignmentGraded", "ExamFinished", "EarlyWarningRaised"]},
            ],
            "relations": [
                ("notification-controller", "cmd-service", "命令", None), ("notification-controller", "query-service", "查询", None),
                ("submitted-handler", "cmd-service", "生成通知", None), ("warning-handler", "cmd-service", "生成通知", None),
                ("exam-handler", "cmd-service", "生成通知", None), ("cmd-service", "notification-repo", "写入", None),
                ("query-service", "notification-repo", "读取", None), ("notification-repo", "notification-entity", "JPA", None),
                ("rabbit", "submitted-handler", "事件消费", "dependency"), ("rabbit", "warning-handler", "事件消费", "dependency"),
                ("rabbit", "exam-handler", "事件消费", "dependency"),
            ],
        },
        {
            "key": "module-ai",
            "title": "AI 智能生成模块详细图",
            "module": "ai-service",
            "description": "AI 模块负责题目生成、组卷生成和学习建议生成请求的编排与记录。AI 生成结果先保存在 sc_ai，业务落库仍由考试或作业等拥有者服务完成。",
            "evidence": ["ai-service/src/main/java", "ai-service-api/src/main/java"],
            "components": [
                {"id": "ai-controller", "label": "AiController", "layer": "controller", "col": 0, "row": 0, "members": ["generateQuestions()", "generateExam()", "learningSuggestion()"]},
                {"id": "ai-service", "label": "AiGenerationService", "layer": "service", "col": 1, "row": 0, "members": ["generateQuestions()", "generateExam()", "saveRecord()"]},
                {"id": "model-client", "label": "AiModelClient", "layer": "api", "col": 2, "row": 0, "members": ["generate(prompt)", "modelName()"]},
                {"id": "mock-client", "label": "LocalMockAiModelClient", "layer": "infra", "col": 2, "row": 1, "members": ["本地模拟生成"]},
                {"id": "ai-repo", "label": "AiGenerationRepository", "layer": "repository", "col": 3, "row": 0, "members": ["save(AiGenerationRecord)"]},
                {"id": "ai-record", "label": "AiGenerationRecord", "layer": "domain", "col": 3, "row": 1, "members": ["requestType", "inputSummary", "outputSummary"]},
                {"id": "ai-entity", "label": "AiGenerationEntity", "layer": "entity", "col": 4, "row": 0, "members": ["modelName", "createdAt", "status"]},
                {"id": "business", "label": "Exam/Assignment Service", "layer": "external", "col": 4, "row": 1, "members": ["生成结果确认后落库"]},
            ],
            "relations": [
                ("ai-controller", "ai-service", "AI 请求", None), ("ai-service", "model-client", "调用模型", "dependency"),
                ("model-client", "mock-client", "本地实现", "dependency"), ("ai-service", "ai-repo", "保存生成记录", None),
                ("ai-repo", "ai-record", "记录对象", None), ("ai-repo", "ai-entity", "JPA 映射", None),
                ("ai-service", "business", "生成结果返回", "dependency"),
            ],
        },
        {
            "key": "module-common",
            "title": "Common 公共基础设施模块详细图",
            "module": "common",
            "description": "common 模块沉淀所有微服务复用的响应结构、异常体系、Feign 上下文传播、TraceId/MDC、管理端点访问控制、outbox relay 和幂等消费能力。它是跨服务技术规范的集中承载点，不拥有具体教学业务数据。",
            "evidence": ["common/src/main/java", "common/src/test/java"],
            "components": [
                {"id": "response", "label": "ResponseResult", "layer": "infra", "col": 0, "row": 0, "members": ["success()", "failure()", "traceId"]},
                {"id": "exception", "label": "BaseGlobalExceptionHandler", "layer": "infra", "col": 0, "row": 1, "members": ["统一异常响应"]},
                {"id": "feign", "label": "FeignRequestInterceptor", "layer": "infra", "col": 1, "row": 0, "members": ["TraceId", "UserId", "ActiveRole"]},
                {"id": "mdc", "label": "CommonServletMdcFilter", "layer": "infra", "col": 1, "row": 1, "members": ["MDC 注入", "响应 TraceId"]},
                {"id": "outbox", "label": "OutboxRelayJob", "layer": "infra", "col": 2, "row": 0, "members": ["扫描 outbox_event", "StreamBridge 发布"]},
                {"id": "idempotent", "label": "IdempotentEventHandler", "layer": "infra", "col": 2, "row": 1, "members": ["processed_event 去重"]},
                {"id": "auto-common", "label": "CommonAutoConfiguration", "layer": "config", "col": 3, "row": 0, "members": ["filter beans", "error decoder", "management access"]},
                {"id": "auto-outbox", "label": "OutboxAutoConfiguration", "layer": "config", "col": 3, "row": 1, "members": ["relay job", "publisher", "repositories"]},
                {"id": "services", "label": "All Microservices", "layer": "external", "col": 4, "row": 0, "members": ["引入 common", "统一横切能力"]},
            ],
            "relations": [
                ("response", "exception", "错误结构", None), ("feign", "mdc", "上下文传播", None),
                ("auto-common", "response", "自动装配", "dependency"), ("auto-common", "feign", "自动装配", "dependency"),
                ("auto-outbox", "outbox", "自动装配", "dependency"), ("auto-outbox", "idempotent", "自动装配", "dependency"),
                ("services", "auto-common", "复用", "dependency"), ("services", "auto-outbox", "事件基础设施", "dependency"),
            ],
        },
        {
            "key": "module-common-events",
            "title": "Common Events 事件契约模块详细图",
            "module": "common-events",
            "description": "common-events 模块定义跨服务异步通信的稳定契约。DomainEvent 统一事件元数据，Assignment、Exam、Notification 和 Warning 事件按聚合划分载荷，供 outbox relay 发布并由分析、通知等消费者幂等处理。",
            "evidence": ["common-events/src/main/java", "common-events/src/test/java"],
            "components": [
                {"id": "domain-event", "label": "DomainEvent<T>", "layer": "api", "col": 0, "row": 0, "members": ["eventId", "aggregate", "occurredAt", "payload"]},
                {"id": "aggregate", "label": "EventAggregate", "layer": "api", "col": 0, "row": 1, "members": ["ASSIGNMENT", "EXAM", "NOTIFICATION", "WARNING"]},
                {"id": "assign-submitted", "label": "AssignmentSubmittedEvent", "layer": "api", "col": 1, "row": 0, "members": ["assignmentId", "submissionId", "studentId"]},
                {"id": "assign-graded", "label": "AssignmentGradedEvent", "layer": "api", "col": 1, "row": 1, "members": ["score", "gradedAt"]},
                {"id": "exam-finished", "label": "ExamFinishedEvent", "layer": "api", "col": 2, "row": 0, "members": ["examId", "submissionId", "score"]},
                {"id": "notification", "label": "NotificationPushedEvent", "layer": "api", "col": 2, "row": 1, "members": ["notificationId", "studentId", "channel"]},
                {"id": "warning-raised", "label": "EarlyWarningRaisedEvent", "layer": "api", "col": 3, "row": 0, "members": ["warningId", "level", "reason"]},
                {"id": "warning-rollback", "label": "EarlyWarningRollbackEvent", "layer": "api", "col": 3, "row": 1, "members": ["warningId", "rollbackReason"]},
                {"id": "consumers", "label": "Event Consumers", "layer": "external", "col": 4, "row": 0, "members": ["analysis-service", "notification-service"]},
            ],
            "relations": [
                ("domain-event", "assign-submitted", "泛型承载", None), ("domain-event", "assign-graded", "泛型承载", None),
                ("domain-event", "exam-finished", "泛型承载", None), ("domain-event", "notification", "泛型承载", None),
                ("domain-event", "warning-raised", "泛型承载", None), ("domain-event", "warning-rollback", "泛型承载", None),
                ("aggregate", "domain-event", "聚合分类", None), ("consumers", "domain-event", "按契约消费", "dependency"),
            ],
        },
        {
            "key": "module-registry-server",
            "title": "Registry Server 注册中心模块详细图",
            "module": "registry-server",
            "description": "registry-server 模块作为 Eureka 服务注册中心，为 Gateway 和所有业务服务提供服务注册、实例续约、服务发现和健康检查入口。它不承载业务数据，但决定微服务按服务名解耦调用的基础能力。",
            "evidence": ["registry-server/src/main/resources/application.yml", "docker-compose.yml", "docs/deployment.md"],
            "components": [
                {"id": "registry-app", "label": "RegistryServerApplication", "layer": "service", "col": 0, "row": 0, "members": ["@EnableEurekaServer"]},
                {"id": "eureka-config", "label": "Eureka Server Config", "layer": "config", "col": 1, "row": 0, "members": ["8761", "服务注册", "服务发现"]},
                {"id": "gateway", "label": "Gateway Eureka Client", "layer": "external", "col": 2, "row": 0, "members": ["fetchRegistry", "lb:// routes"]},
                {"id": "services", "label": "Business Eureka Clients", "layer": "external", "col": 3, "row": 0, "members": ["auth/user/course", "assignment/exam/..."]},
                {"id": "actuator", "label": "Spring Boot Actuator", "layer": "infra", "col": 1, "row": 1, "members": ["health", "metrics", "prometheus"]},
                {"id": "prom", "label": "Prometheus", "layer": "external", "col": 2, "row": 1, "members": ["targets", "alerts"]},
                {"id": "grafana", "label": "Grafana", "layer": "external", "col": 3, "row": 1, "members": ["dashboards"]},
                {"id": "helm", "label": "Helm Charts", "layer": "config", "col": 4, "row": 0, "members": ["target-platform", "values-dev/test/prod"]},
            ],
            "relations": [
                ("registry-app", "eureka-config", "启动配置", None), ("gateway", "eureka-config", "发现服务", "dependency"),
                ("services", "eureka-config", "注册实例", "dependency"), ("services", "actuator", "暴露指标", "dependency"),
                ("prom", "actuator", "抓取指标", "dependency"), ("grafana", "prom", "数据源", "dependency"),
                ("helm", "services", "部署模板", "dependency"),
            ],
        },
        {
            "key": "module-legacy",
            "title": "Legacy Adapter 迁移适配模块详细图",
            "module": "legacy-adapter",
            "description": "legacy-adapter 模块专门承接迁移期仍需保留的少量旧读接口，当前以 LegacyKnowledgeReadController 和 LegacyKnowledgeReadService 暴露知识点兼容查询。它通过 Gateway legacy-route 进入，边界目标是只读、可替换、避免新业务继续回流到单体。",
            "evidence": ["legacy-adapter/src/main/java", "legacy-adapter/src/test/java", "docs/service-boundary-and-monolith-shrink-plan.md"],
            "components": [
                {"id": "legacy-app", "label": "LegacyAdapterApplication", "layer": "service", "col": 0, "row": 0, "members": ["Spring Boot app", "legacy-adapter"]},
                {"id": "legacy-controller", "label": "LegacyKnowledgeReadController", "layer": "controller", "col": 1, "row": 0, "members": ["旧知识点读接口", "兼容 API"]},
                {"id": "legacy-service", "label": "LegacyKnowledgeReadService", "layer": "service", "col": 2, "row": 0, "members": ["兼容查询", "只读边界"]},
                {"id": "legacy-db", "label": "major_assignment schema", "layer": "external", "col": 3, "row": 0, "members": ["迁移期读取", "旧知识点表"]},
                {"id": "gateway", "label": "Gateway legacy-route", "layer": "external", "col": 0, "row": 1, "members": ["LEGACY_BASE_URL", "路由兜底"]},
                {"id": "tests", "label": "LegacyKnowledgeReadControllerTest", "layer": "infra", "col": 2, "row": 1, "members": ["兼容契约验证"]},
                {"id": "target", "label": "Target Course Service", "layer": "external", "col": 4, "row": 0, "members": ["逐步替代旧接口"]},
            ],
            "relations": [
                ("gateway", "legacy-controller", "旧读接口", "dependency"), ("legacy-app", "legacy-controller", "注册控制器", None),
                ("legacy-controller", "legacy-service", "查询", None), ("legacy-service", "legacy-db", "只读兼容", "dependency"),
                ("tests", "legacy-controller", "验证 API", "dependency"), ("target", "legacy-controller", "替代后下线", "dependency"),
            ],
        },
    ]

    specs.extend(make_missing_maven_module_specs())
    module_order = {name: i for i, name in enumerate(MAVEN_MODULES)}
    specs.sort(key=lambda item: module_order.get(item["module"], len(module_order)))

    diagrams: list[Diagram] = []
    for offset, spec in enumerate(specs):
        diagrams.append(create_module_diagram(
            start_index + offset,
            spec["key"],
            spec["title"],
            spec["module"],
            spec["description"],
            spec["evidence"],
            spec["components"],
            spec["relations"],
        ))
    return diagrams


def clear_generated_image_files() -> None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    for file in IMAGE_DIR.glob("[0-9][0-9]-*.png"):
        file.unlink()
    contact_sheet = IMAGE_DIR / "contact-sheet.png"
    if contact_sheet.exists():
        contact_sheet.unlink()


def render_all(diagrams: list[Diagram]) -> list[Path]:
    clear_generated_image_files()
    paths = []
    for diagram in diagrams:
        if diagram.uml_type == "Sequence Diagram":
            paths.append(render_sequence(diagram))
        elif diagram.uml_type == "Communication Diagram":
            paths.append(render_communication(diagram))
        elif diagram.uml_type == "Object Diagram":
            paths.append(render_object(diagram))
        else:
            if diagram.metadata.get("size"):
                size = tuple(diagram.metadata["size"])
            elif diagram.metadata.get("detail"):
                size = (2200, 1540)
            else:
                if diagram.key == "activity":
                    size = (1900, 2000)
                elif diagram.key in {"class", "deployment", "component"}:
                    size = (2100, 1420)
                else:
                    size = (1900, 1220)
            paths.append(render_graph_diagram(diagram, size))
    return paths


def make_contact_sheet(diagrams: list[Diagram]) -> Path:
    images = [Image.open(IMAGE_DIR / d.filename).convert("RGB") for d in diagrams]
    thumb_w = 560
    thumbs = []
    for img in images:
        ratio = thumb_w / img.width
        thumbs.append(img.resize((thumb_w, int(img.height * ratio))))
    label_font = font(24)
    cols = 3
    rows = math.ceil(len(thumbs) / cols)
    cell_h = max(t.height for t in thumbs) + 54
    sheet = Image.new("RGB", (thumb_w * cols, cell_h * rows), WHITE)
    draw = ImageDraw.Draw(sheet)
    for i, thumb in enumerate(thumbs):
        x = (i % cols) * thumb_w
        y = (i // cols) * cell_h
        sheet.paste(thumb, (x, y + 40))
        draw.text((x + 12, y + 8), f"{i + 1:02d} {diagrams[i].uml_type}", font=label_font, fill="#0f172a")
    out = IMAGE_DIR / "contact-sheet.png"
    sheet.save(out)
    return out


def write_sources(diagrams: list[Diagram]) -> None:
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    for file in SOURCE_DIR.glob("[0-9][0-9]-*.puml"):
        file.unlink()
    for idx, diagram in enumerate(diagrams, 1):
        (SOURCE_DIR / f"{idx:02d}-{diagram.key}.puml").write_text(diagram.plantuml, encoding="utf-8")
    spec = {
        "systemName": "智能学习辅助系统",
        "project": "D:/111/Distributed framework technology/JavaCode/majorassignment",
        "generatedFor": "分布式框架技术课程 UML 实验报告",
        "diagrams": [
            {
                "key": d.key,
                "title": d.title,
                "umlType": d.uml_type,
                "image": f"images/{d.filename}",
                "description": d.description,
                "evidence": d.evidence,
                "nodes": [n.__dict__ for n in d.nodes],
                "edges": [e.__dict__ for e in d.edges],
                "metadata": d.metadata,
            }
            for d in diagrams
        ],
    }
    SPEC_PATH.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")


def sid(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:18]}"


def mdj_ref(element_id: str) -> dict[str, str]:
    return {"$ref": element_id}


def make_mdj(diagrams: list[Diagram]) -> None:
    project_id = sid("project")
    model_id = sid("model")
    model: dict[str, Any] = {
        "_type": "UMLModel",
        "_id": model_id,
        "_parent": mdj_ref(project_id),
        "name": "智能学习辅助系统 UML 模型",
        "ownedElements": [],
    }
    project: dict[str, Any] = {
        "_type": "Project",
        "_id": project_id,
        "name": "智能学习辅助系统 UML 模型",
        "ownedElements": [model],
    }

    type_map = {
        "Use Case Diagram": "UMLUseCaseDiagram",
        "Class Diagram": "UMLClassDiagram",
        "Object Diagram": "UMLObjectDiagram",
        "Sequence Diagram": "UMLSequenceDiagram",
        "Communication Diagram": "UMLCommunicationDiagram",
        "Statechart Diagram": "UMLStatechartDiagram",
        "Activity Diagram": "UMLActivityDiagram",
        "Component Diagram": "UMLComponentDiagram",
        "Deployment Diagram": "UMLDeploymentDiagram",
    }
    element_map = {
        "actor": "UMLActor",
        "usecase": "UMLUseCase",
        "class": "UMLClass",
        "lifeline": "UMLClass",
        "state": "UMLState",
        "activity": "UMLAction",
        "decision": "UMLDecisionNode",
        "component": "UMLComponent",
        "node3d": "UMLNode",
        "initial": "UMLState",
        "final": "UMLState",
    }

    for diagram in diagrams:
        pkg_id = sid(diagram.key)
        pkg: dict[str, Any] = {
            "_type": "UMLPackage",
            "_id": pkg_id,
            "_parent": mdj_ref(model_id),
            "name": diagram.title,
            "documentation": diagram.description,
            "ownedElements": [],
        }
        dgm_id = sid(f"{diagram.key}-diagram")
        dgm: dict[str, Any] = {
            "_type": type_map.get(diagram.uml_type, "UMLClassDiagram"),
            "_id": dgm_id,
            "_parent": mdj_ref(pkg_id),
            "name": diagram.title,
            "visible": True,
            "ownedViews": [],
        }
        pkg["ownedElements"].append(dgm)
        node_ids: dict[str, str] = {}
        for node in diagram.nodes:
            if node.kind in {"initial", "final"}:
                continue
            elem_id = sid(node.id)
            node_ids[node.id] = elem_id
            elem_type = element_map.get(node.kind, "UMLClass")
            elem: dict[str, Any] = {
                "_type": elem_type,
                "_id": elem_id,
                "_parent": mdj_ref(pkg_id),
                "name": node.label.replace("\\n", " "),
                "documentation": "\n".join(node.members),
            }
            pkg["ownedElements"].append(elem)
        for edge in diagram.edges:
            rel_id = sid("rel")
            rel_type = "UMLDependency" if edge.kind == "dependency" or edge.dashed else "UMLAssociation"
            if edge.kind == "message":
                rel_type = "UMLAssociation"
            rel: dict[str, Any] = {
                "_type": rel_type,
                "_id": rel_id,
                "_parent": mdj_ref(pkg_id),
                "name": edge.label,
            }
            if edge.source in node_ids:
                rel["source"] = mdj_ref(node_ids[edge.source])
            if edge.target in node_ids:
                rel["target"] = mdj_ref(node_ids[edge.target])
            pkg["ownedElements"].append(rel)
        model["ownedElements"].append(pkg)
    MODEL_PATH.write_text(json.dumps(project, ensure_ascii=False, indent=2), encoding="utf-8")


def write_staruml_extension() -> None:
    spec_text = SPEC_PATH.read_text(encoding="utf-8")
    spec = json.loads(spec_text)
    diagram_total = len(spec.get("diagrams", []))
    core_total = sum(1 for item in spec.get("diagrams", []) if item.get("metadata", {}).get("group") == "core")
    extension_total = diagram_total - core_total
    package_json = {
        "name": "smart-learning-uml-generator",
        "title": "智能学习辅助系统 UML 一键生成",
        "description": f"为 majorassignment 项目一键生成 {core_total} 张课程核心 UML 图和 {extension_total} 张业务扩展 UML 图。",
        "version": "1.0.0",
        "keywords": ["uml", "coursework", "smart-learning", "staruml"],
        "author": {"name": "Codex"},
        "license": "MIT",
        "engines": {"staruml": ">=6.0.0"},
        "main": "main.js",
    }
    (STARUML_DIR / "package.json").write_text(json.dumps(package_json, ensure_ascii=False, indent=2), encoding="utf-8")
    menu = [
        {
            "label": "生成课程核心 9 张 UML 图",
            "id": "tools.smart-learning-uml.generate-core",
            "command": "smart-learning-uml:generate-core",
        },
        {
            "label": "生成推荐 24 张 UML 图",
            "id": "tools.smart-learning-uml.generate-all",
            "command": "smart-learning-uml:generate-all",
        },
        {
            "label": "智能学习 UML 使用说明",
            "id": "tools.smart-learning-uml.help",
            "command": "smart-learning-uml:show-help",
        },
    ]
    (STARUML_DIR / "menus" / "smart-learning-uml.json").write_text(
        json.dumps(menu, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    main_js = f"""/*
 * StarUML extension generated for the intelligent learning assistance system.
 * StarUML v6 extension entry: main.js exports init(), registers commands,
 * then attaches menu items to Tools with app.menu.add().
 */
const MODEL_SPEC = {spec_text};
const type = app.type || (typeof global !== 'undefined' && global.type) || (typeof window !== 'undefined' && window.type);

function createModel(parent, id, name, documentation) {{
  return app.factory.createModel({{
    id,
    parent,
    modelInitializer: function (elem) {{
      elem.name = name || id;
      if (documentation) {{
        elem.documentation = documentation;
      }}
    }}
  }});
}}

function createDiagram(parent, id, name) {{
  const diagram = app.factory.createDiagram({{
    id,
    parent,
    diagramInitializer: function (dgm) {{
      dgm.name = name;
    }}
  }});
  return {{ model: diagram._parent || parent, diagram }};
}}

function createBox(diagram, parent, id, label, x, y, w, h) {{
  return app.factory.createModelAndView({{
    id,
    parent,
    diagram,
    x1: x,
    y1: y,
    x2: x + w,
    y2: y + h,
    viewInitializer: function (view) {{
      view.width = w;
      view.height = h;
    }},
    modelInitializer: function (elem) {{
      elem.name = label;
    }}
  }});
}}

function createPlainBoxModel(diagram, parent, id, label, x, y, w, h) {{
  return app.factory.createModelAndView({{
    id,
    parent,
    diagram,
    x1: x,
    y1: y,
    x2: x + w,
    y2: y + h,
    viewInitializer: function (view) {{
      view.width = w;
      view.height = h;
    }},
    modelInitializer: function (elem) {{
      elem.name = label;
    }}
  }});
}}

function createRelation(diagram, parent, id, label, tailView, headView) {{
  try {{
    return app.factory.createModelAndView({{
      id,
      parent,
      diagram,
      tailView,
      headView,
      tailModel: tailView.model,
      headModel: headView.model,
      x1: tailView.left + Math.floor(tailView.width / 2),
      y1: tailView.top + Math.floor(tailView.height / 2),
      modelInitializer: function (elem) {{
        elem.name = label || '';
      }}
    }});
  }} catch (err) {{
    console.warn('Relation skipped:', id, label, err);
    return null;
  }}
}}

function createActivityFlow(diagram, parent, label, tailView, headView) {{
  return createRelation(diagram, parent, 'UMLControlFlow', label, tailView, headView);
}}

function activityNodeId(kind) {{
  const map = {{
    action: 'UMLAction',
    decision: 'UMLDecisionNode',
    merge: 'UMLMergeNode',
    fork: 'UMLForkNode',
    join: 'UMLJoinNode',
    initial: 'UMLInitialNode',
    final: 'UMLActivityFinalNode'
  }};
  return map[kind] || 'UMLAction';
}}

function stateBounds(node) {{
  const w = Math.max(40, Math.round((node.w || 120) / 2));
  const h = Math.max(34, Math.round((node.h || 70) / 2));
  const x = Math.round((node.x || 0) / 2);
  const y = Math.round((node.y || 0) / 2);
  return {{ x1: x, y1: y, x2: x + w, y2: y + h, width: w, height: h }};
}}

function viewCenter(view) {{
  const left = typeof view.left === 'number' ? view.left : 0;
  const top = typeof view.top === 'number' ? view.top : 0;
  const width = typeof view.width === 'number' ? view.width : 0;
  const height = typeof view.height === 'number' ? view.height : 0;
  return [left + width / 2, top + height / 2];
}}

function insertNodeView(builder, diagram, model, ViewType, node) {{
  const b = stateBounds(node);
  const view = new ViewType();
  view._parent = diagram;
  view.model = model;
  view.initialize(null, b.x1, b.y1, b.x2, b.y2);
  builder.insert(view);
  builder.fieldInsert(diagram, 'ownedViews', view);
  return view;
}}

function stateNodeName(node) {{
  if (node.kind === 'initial') return '开始';
  if (node.kind === 'final') return '结束';
  return node.label || node.id;
}}

function buildStateNode(builder, region, diagram, node) {{
  const model = new type.UMLState();
  const ViewType = type.UMLStateView;
  model._parent = region;
  model.name = stateNodeName(node);
  builder.insert(model);
  builder.fieldInsert(region, 'vertices', model);
  return insertNodeView(builder, diagram, model, ViewType, node);
}}

function buildStateTransition(builder, region, diagram, sourceView, targetView, label) {{
  const model = new type.UMLTransition();
  model._parent = region;
  model.name = label || '';
  model.source = sourceView.model;
  model.target = targetView.model;
  if (type.UMLTransition && type.UMLTransition.TK_EXTERNAL) {{
    model.kind = type.UMLTransition.TK_EXTERNAL;
  }}
  builder.insert(model);
  builder.fieldInsert(region, 'transitions', model);
  const view = new type.UMLTransitionView();
  view._parent = diagram;
  view.model = model;
  view.tail = sourceView;
  view.head = targetView;
  const start = viewCenter(sourceView);
  const end = viewCenter(targetView);
  view.initialize(null, start[0], start[1], end[0], end[1]);
  builder.insert(view);
  builder.fieldInsert(diagram, 'ownedViews', view);
  return view;
}}

function createSequenceDiagram(pkg, item) {{
  const diagram = createDiagram(pkg, 'UMLCommunicationDiagram', item.title).diagram;
  const lifelines = (item.metadata && item.metadata.lifelines) || [];
  const messages = (item.metadata && item.metadata.messages) || [];
  const views = {{}};
  const top = 90;
  const spacing = 150;
  lifelines.forEach(function (label, idx) {{
    const id = 'seq-' + idx;
    const x = 40 + idx * spacing;
    views[label] = createPlainBoxModel(diagram, pkg, 'UMLClass', label, x, top, 126, 460);
  }});
  messages.forEach(function (msg, idx) {{
    const fromView = views[msg.from];
    const toView = views[msg.to];
    if (fromView && toView) {{
      createRelation(diagram, pkg, msg.dashed ? 'UMLDependency' : 'UMLAssociation', (idx + 1) + '. ' + msg.label, fromView, toView);
    }}
  }});
}}

function createStatechartDiagram(parent, item) {{
  const builder = app.repository.getOperationBuilder();
  builder.begin('generate ' + item.title);
  const stateMachine = new type.UMLStateMachine();
  stateMachine._parent = parent;
  stateMachine.name = item.title;
  builder.insert(stateMachine);
  builder.fieldInsert(parent, 'ownedElements', stateMachine);
  const region = new type.UMLRegion();
  region._parent = stateMachine;
  builder.insert(region);
  builder.fieldInsert(stateMachine, 'regions', region);
  const diagram = new type.UMLStatechartDiagram();
  diagram._parent = stateMachine;
  diagram.name = item.title;
  builder.insert(diagram);
  builder.fieldInsert(stateMachine, 'ownedElements', diagram);
  const views = {{}};
  (item.nodes || []).forEach(function (node) {{
    views[node.id] = buildStateNode(builder, region, diagram, node);
  }});
  (item.edges || []).forEach(function (edge) {{
    if (views[edge.source] && views[edge.target]) {{
      buildStateTransition(builder, region, diagram, views[edge.source], views[edge.target], edge.label || '');
    }}
  }});
  builder.end();
  app.repository.doOperation(builder.getOperation());
}}

function diagramId(typeName) {{
  const map = {{
    'Use Case Diagram': 'UMLUseCaseDiagram',
    'Class Diagram': 'UMLClassDiagram',
    'Object Diagram': 'UMLObjectDiagram',
    'Sequence Diagram': 'UMLSequenceDiagram',
    'Communication Diagram': 'UMLCommunicationDiagram',
    'Statechart Diagram': 'UMLStatechartDiagram',
    'Activity Diagram': 'UMLActivityDiagram',
    'Component Diagram': 'UMLComponentDiagram',
    'Deployment Diagram': 'UMLDeploymentDiagram'
  }};
  return map[typeName] || 'UMLClassDiagram';
}}

function modelId(kind, typeName) {{
  if (typeName === 'Use Case Diagram') {{
    return kind === 'actor' ? 'UMLActor' : 'UMLUseCase';
  }}
  if (typeName === 'Component Diagram') return 'UMLComponent';
  if (typeName === 'Deployment Diagram') return 'UMLNode';
  if (typeName === 'Statechart Diagram') return 'UMLState';
  if (typeName === 'Activity Diagram') return activityNodeId(kind);
  if (typeName === 'Sequence Diagram') return 'UMLLifeline';
  return 'UMLClass';
}}

function relationId(edge, typeName) {{
  if (typeName === 'Activity Diagram') return 'UMLControlFlow';
  if (typeName === 'Statechart Diagram') return 'UMLTransition';
  if (edge.kind === 'dependency' || edge.dashed) return 'UMLDependency';
  return 'UMLAssociation';
}}

function createOneDiagram(root, item) {{
  const pkg = createModel(root, 'UMLPackage', item.title, item.description);
  if (item.umlType === 'Sequence Diagram') {{
    createSequenceDiagram(pkg, item);
    return;
  }}
  if (item.umlType === 'Statechart Diagram') {{
    createStatechartDiagram(pkg, item);
    return;
  }}
  const created = createDiagram(pkg, diagramId(item.umlType), item.title);
  const diagram = created.diagram;
  const owner = (item.umlType === 'Activity Diagram' || item.umlType === 'Component Diagram' || item.umlType === 'Deployment Diagram') ? created.model : pkg;
  const views = {{}};
  (item.nodes || []).forEach(function (node) {{
    if (item.umlType === 'Statechart Diagram' && (node.kind === 'initial' || node.kind === 'final')) return;
    const view = createBox(
      diagram,
      owner,
      modelId(node.kind, item.umlType),
      String(node.label || node.id).replace(/\\\\n/g, ' '),
      Math.round((node.x || 0) / 2),
      Math.round((node.y || 0) / 2),
      Math.max(80, Math.round((node.w || 200) / 2)),
      Math.max(40, Math.round((node.h || 90) / 2))
    );
    views[node.id] = view;
    if (node.members && node.members.length && view.model) {{
      view.model.documentation = node.members.join('\\n');
    }}
  }});
  (item.edges || []).forEach(function (edge) {{
    if (views[edge.source] && views[edge.target]) {{
      createRelation(diagram, owner, relationId(edge, item.umlType), edge.label || '', views[edge.source], views[edge.target]);
    }}
  }});
}}

function selectedDiagrams(scope) {{
  if (scope === 'core') {{
    return MODEL_SPEC.diagrams.filter(function (item) {{
      return item.metadata && item.metadata.group === 'core';
    }});
  }}
  return MODEL_SPEC.diagrams;
}}

function handleGenerate(scope) {{
  const project = app.repository.select('@Project')[0];
  const items = selectedDiagrams(scope);
  const suffix = scope === 'core' ? '课程核心 9 张' : '推荐 24 张业务主线';
  const root = createModel(project, 'UMLModel', MODEL_SPEC.systemName + ' UML 模型-' + suffix, MODEL_SPEC.generatedFor);
  const failures = [];
  items.forEach(function (item) {{
    try {{
      createOneDiagram(root, item);
    }} catch (err) {{
      failures.push(item.title + ': ' + (err && err.message ? err.message : err));
      console.error('Diagram generation failed:', item.title, err);
    }}
  }});
  if (failures.length) {{
    app.dialogs.showInfoDialog('生成完成，但有 ' + failures.length + ' 张图失败。已成功生成 ' + (items.length - failures.length) + ' 张。失败项：\\n' + failures.join('\\n'));
  }} else {{
    app.dialogs.showInfoDialog('已生成智能学习辅助系统 ' + items.length + ' 张 UML 模型图。请在 Model Explorer 中展开“智能学习辅助系统 UML 模型”查看；需要图片时运行交付目录中的 export-staruml-images.ps1。');
  }}
}}

function handleHelp() {{
  app.dialogs.showInfoDialog('使用流程：1. 运行“一键打开StarUML并安装扩展.bat”；2. StarUML 打开后进入 Tools 菜单；3. 点击“生成课程核心 9 张 UML 图”可生成课程必需图；4. 点击“生成推荐 24 张 UML 图”可生成完整业务主线版。');
}}

function init() {{
  app.commands.register('smart-learning-uml:generate-core', function () {{ handleGenerate('core'); }});
  app.commands.register('smart-learning-uml:generate-all', function () {{ handleGenerate('all'); }});
  app.commands.register('smart-learning-uml:show-help', handleHelp);
  app.menu.add([{{ id: 'tools', submenu: require('./menus/smart-learning-uml.json') }}]);
}}

exports.init = init;
"""
    (STARUML_DIR / "main.js").write_text(main_js, encoding="utf-8")
    zip_path = ROOT / "smart-learning-uml-generator-staruml-extension.zip"
    if zip_path.exists():
        zip_path.unlink()
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for file in STARUML_DIR.rglob("*"):
            if file.is_file():
                zf.write(file, file.relative_to(STARUML_DIR.parent))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_doc_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size, color in [
        ("Title", 22, "0F172A"),
        ("Heading 1", 16, "1D4ED8"),
        ("Heading 2", 13, "334155"),
        ("Heading 3", 11.5, "334155"),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    with Image.open(image_path) as img:
        width_px, height_px = img.size
    max_width = 6.8
    max_height = 8.0
    ratio = width_px / max(1, height_px)
    width_in = min(max_width, max_height * ratio)
    height_in = width_in / max(0.1, ratio)
    if height_in > max_height:
        height_in = max_height
        width_in = height_in * ratio
    run.add_picture(str(image_path), width=Inches(width_in))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(75, 85, 99)


def make_report(diagrams: list[Diagram]) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    diagram_total = len(diagrams)
    core_diagrams = [d for d in diagrams if d.metadata.get("group") == "core"]
    extension_diagrams = [d for d in diagrams if d.metadata.get("group") == "extension"]
    doc = Document()
    set_doc_styles(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("智能学习辅助系统 UML 建模实验报告")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("分布式框架技术课程项目设计文档配套 UML 图").font.size = Pt(11)

    doc.add_heading("一、实验目的", level=1)
    add_para(doc, f"本实验围绕 majorassignment 项目的智能学习辅助系统进行 UML 建模。报告采用“课程核心图 + 业务扩展图”的组织方式：先输出用例图、类图、对象图、顺序图、协作图、状态图、活动图、组件图和部署图 9 类课程核心 UML 图，再围绕登录认证、作业批改、考试提交、AI 组卷、学情分析、通知推送和微服务依赖补充 {len(extension_diagrams)} 张业务扩展图，共 {diagram_total} 张 UML 图片。")
    add_para(doc, "项目采用 Spring Cloud 微服务迁移架构，包含 Gateway、Eureka 注册中心、认证、用户、课程、作业、考试、分析、通知、AI 服务以及 legacy-adapter 和过渡单体 major_assignment。")

    doc.add_heading("二、实验环境与生成方式", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "说明"
    for cell in hdr:
        set_cell_shading(cell, "DBEAFE")
    rows = [
        ("系统名称", "智能学习辅助系统"),
        ("项目目录", "D:/111/Distributed framework technology/JavaCode/majorassignment"),
        ("建模范围", f"{len(core_diagrams)} 张课程核心 UML 图 + {len(extension_diagrams)} 张智能学习业务扩展 UML 图"),
        ("图类型", "用例图、类图、对象图、顺序图、协作图、状态图、活动图、组件图、部署图"),
        ("图片数量", f"{diagram_total} 张 UML 图片，外加 1 张 contact-sheet 总览图"),
        ("可编辑产物", "sources/smart-learning-system.mdj 与 staruml-extension/smart-learning-uml-generator"),
        ("一键生成入口", "运行 tools/generate_uml_delivery.py 重新生成全部图片、模型源、扩展包和报告"),
    ]
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = v
        set_cell_width(cells[0], 4.2)
        set_cell_width(cells[1], 11.8)

    doc.add_heading("三、UML 图清单", level=1)
    list_table = doc.add_table(rows=1, cols=5)
    list_table.style = "Table Grid"
    headers = ["序号", "分组", "图名", "UML 类型", "建模重点"]
    for idx, text in enumerate(headers):
        list_table.rows[0].cells[idx].text = text
        set_cell_shading(list_table.rows[0].cells[idx], "DBEAFE")
    focus_map = {
        "use_case": "参与者与用户目标",
        "class": "领域对象与分层结构",
        "object": "考试提交运行时快照",
        "sequence": "考试提交消息顺序",
        "communication": "对象协作与消息编号",
        "state": "考试提交记录生命周期",
        "activity": "提交、阅卷、分析流程",
        "component": "微服务组件依赖",
        "deployment": "Docker Compose 部署拓扑",
    }
    for i, diagram in enumerate(diagrams, 1):
        cells = list_table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = "课程核心" if diagram.metadata.get("group") == "core" else "业务扩展"
        cells[2].text = diagram.title
        cells[3].text = diagram.uml_type
        cells[4].text = diagram.metadata.get("focus", focus_map.get(diagram.key, diagram.description[:30]))
        set_cell_width(cells[0], 1.2)
        set_cell_width(cells[1], 2.5)
        set_cell_width(cells[2], 5.0)
        set_cell_width(cells[3], 3.8)
        set_cell_width(cells[4], 4.5)

    def add_diagram_section(section_title: str, items: list[Diagram], start_new_page: bool) -> None:
        if start_new_page:
            doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading(section_title, level=1)
        for index, diagram in enumerate(items, 1):
            if index > 1:
                doc.add_section(WD_SECTION_START.NEW_PAGE)
            doc.add_heading(diagram.title, level=2)
            add_figure(doc, IMAGE_DIR / diagram.filename, f"{diagram.title}（{diagram.uml_type}）")
            add_para(doc, diagram.description)
            doc.add_heading("建模依据", level=3)
            for item in diagram.evidence:
                p = doc.add_paragraph(style=None)
                p.style = doc.styles["Normal"]
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.first_line_indent = Cm(-0.3)
                p.add_run("• ").bold = True
                p.add_run(item)

    add_diagram_section("四、课程核心 UML 图及说明", core_diagrams, False)
    add_diagram_section("五、智能学习业务扩展 UML 图及说明", extension_diagrams, True)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("六、StarUML 导入与扩展使用说明", level=1)
    add_para(doc, "方式一：双击交付目录中的“一键打开StarUML并安装扩展.bat”。脚本会安装 StarUML 扩展并打开 StarUML，打开后在顶部 Tools 菜单点击“生成课程核心 9 张 UML 图”或“生成推荐 24 张 UML 图”，即可在 Model Explorer 中生成并查看对应模型图。")
    add_para(doc, f"方式二：打开 StarUML，直接选择 File > Open，打开 sources/smart-learning-system.mdj，可查看模型包和 {diagram_total} 张图的模型元素。")
    add_para(doc, "方式三：将 staruml-extension/smart-learning-uml-generator 文件夹复制到 StarUML 用户扩展目录，重启 StarUML 后在 Tools 菜单选择需要的生成入口。")
    add_para(doc, "Windows 用户扩展目录通常是 C:/Users/<用户名>/AppData/Roaming/StarUML/extensions/user。扩展包 smart-learning-uml-generator-staruml-extension.zip 可直接解压到该目录。")
    add_para(doc, "经查 StarUML 官方文档，扩展 API 可注册命令、菜单并创建模型；批量导出全部图为 PNG 的官方路径是 StarUML CLI 的 image 命令，默认选择 @Diagram，可一次导出模型中的全部图。本交付提供 export-staruml-images.ps1 调用该能力；本地生成器也会同步生成全部 PNG，确保没有 StarUML CLI 时仍可提交报告。")
    add_para(doc, "图片与报告由本地生成器确定性生成，StarUML 中的图形位置可根据教师审美要求继续手工微调。")

    doc.add_heading("七、实验总结", level=1)
    add_para(doc, "本次建模从功能、静态结构、动态交互、生命周期、流程、组件和部署多个视角描述了智能学习辅助系统。模型反映了项目当前“微服务为主体、单体为过渡层”的架构状态，也体现了同步 Gateway 路由与异步事件驱动分析通知链路的组合。")
    out = REPORT_DIR / "智能学习辅助系统-UML建模实验报告.docx"
    doc.save(out)


def write_readme(diagrams: list[Diagram]) -> None:
    diagram_total = len(diagrams)
    core_total = sum(1 for d in diagrams if d.metadata.get("group") == "core")
    extension_total = diagram_total - core_total
    lines = [
        "# 智能学习辅助系统 UML 交付目录",
        "",
        f"本目录包含 {core_total} 张课程核心 UML 图、{extension_total} 张智能学习业务扩展 UML 图、StarUML 可导入模型、StarUML 一键生成扩展和实验报告。",
        "",
        "## 一键重新生成",
        "",
        "```powershell",
        "powershell -ExecutionPolicy Bypass -File docs\\uml-staruml-delivery\\generate-all.ps1",
        "```",
        "",
        "## 一键安装 StarUML 扩展",
        "",
        "```powershell",
        "powershell -ExecutionPolicy Bypass -File docs\\uml-staruml-delivery\\install-staruml-extension.ps1",
        "```",
        "",
        "## 一键打开 StarUML 并点击扩展生成图",
        "",
        "最简单方式：双击交付目录里的 `一键打开StarUML并安装扩展.bat`。脚本会把扩展安装到 StarUML 用户扩展目录，然后打开 StarUML。StarUML 打开后，在顶部菜单选择：",
        "",
        "`Tools > 生成课程核心 9 张 UML 图` 或 `Tools > 生成推荐 24 张 UML 图`",
        "",
        f"点击后会在 Model Explorer 中生成对应的 {core_total} 张或 {diagram_total} 张 UML 模型图。",
        "",
        "也可以用 PowerShell 运行同一个流程：",
        "",
        "```powershell",
        "powershell -ExecutionPolicy Bypass -File docs\\uml-staruml-delivery\\install-and-open-staruml.ps1",
        "```",
        "",
        "## 备用：直接打开 StarUML 模型文件",
        "",
        "```powershell",
        "powershell -ExecutionPolicy Bypass -File docs\\uml-staruml-delivery\\open-in-staruml.ps1",
        "```",
        "",
        "该脚本会查找 `staruml` 命令或常见 Windows 安装路径，并直接用 StarUML 打开 `sources/smart-learning-system.mdj`。",
        "",
        "## 一键使用 StarUML CLI 导出全部图片",
        "",
        "```powershell",
        "powershell -ExecutionPolicy Bypass -File docs\\uml-staruml-delivery\\export-staruml-images.ps1",
        "```",
        "",
        "StarUML 官方 CLI 的 `image` 命令默认选择器是 `@Diagram`，因此该脚本会从 `sources/smart-learning-system.mdj` 一次导出全部图。",
        "",
        "## 主要产物",
        "",
        f"- `images/`：{diagram_total} 张 UML 图片 PNG。",
        f"- `images/contact-sheet.png`：{diagram_total} 张图的接触表，便于快速检查整体可读性。",
        "- `sources/*.puml`：每张图的 PlantUML 骨架，便于课程答辩说明建模逻辑。",
        "- `sources/uml-model-spec.json`：StarUML 扩展和图片生成共同使用的结构化模型源。",
        "- `sources/smart-learning-system.mdj`：可直接在 StarUML 中打开的模型项目。",
        "- `staruml-extension/smart-learning-uml-generator/`：StarUML 扩展源码。",
        "- `smart-learning-uml-generator-staruml-extension.zip`：可解压安装的 StarUML 扩展包。",
        "- `一键打开StarUML并安装扩展.bat`：双击后安装扩展并打开 StarUML，然后在 Tools 菜单选择核心 9 张或推荐 24 张生成入口。",
        "- `open-in-staruml.ps1`：直接启动 StarUML 并打开 `.mdj` 模型。",
        "- `install-and-open-staruml.ps1`：安装 StarUML 扩展后启动 StarUML。",
        "- `export-staruml-images.ps1`：调用 StarUML CLI `image` 命令，把 `.mdj` 中全部图导出为 PNG。",
        "- `report/智能学习辅助系统-UML建模实验报告.docx`：含每张 UML 图片和说明的实验报告。",
        "- `report/智能学习辅助系统-UML建模实验报告.pdf`：若本机安装 Word，`generate-all.ps1` 会自动导出 PDF。",
        "",
        "## StarUML 使用",
        "",
        "1. 扩展生成：双击 `一键打开StarUML并安装扩展.bat`，等待 StarUML 打开。",
        "2. 在 StarUML 顶部菜单点击 `Tools > 生成课程核心 9 张 UML 图` 或 `Tools > 生成推荐 24 张 UML 图`。",
        f"3. 生成后，在左侧 `Model Explorer` 展开 `智能学习辅助系统 UML 模型`，即可查看对应的 {core_total} 张或 {diagram_total} 张 UML 模型图。",
        "4. 备用模型：StarUML 选择 `File > Open`，打开 `sources/smart-learning-system.mdj`。",
        "5. 批量导图：安装 StarUML CLI 后运行 `export-staruml-images.ps1`，会按官方 `image` 命令导出全部图。",
        "6. Windows 扩展目录通常是 `C:\\Users\\<用户名>\\AppData\\Roaming\\StarUML\\extensions\\user`。",
        "",
        "## 图清单",
        "",
    ]
    for d in diagrams:
        group = "课程核心" if d.metadata.get("group") == "core" else "业务扩展"
        lines.append(f"- `{d.filename}`：{group}，{d.title}，{d.uml_type}")
    lines.append("")
    lines.append("## 文档依据")
    lines.append("")
    lines.append("- StarUML 官方扩展文档：`main.js` 作为入口，`app.commands.register` 注册命令，`menus/*.json` 声明菜单。")
    lines.append("- StarUML 官方 CLI 文档：`staruml image <file> -f png -o \"out/<%=filenamify(element.name)%>.png\"` 可批量导出全部图，默认选择器为 `@Diagram`。")
    lines.append("- 项目本地文档：`docs/architecture.md`、`docs/deployment.md`、`docs/data-ownership.md`、`docs/service-boundary-and-monolith-shrink-plan.md`。")
    (ROOT / "README.md").write_text("\n".join(lines), encoding="utf-8")


def write_staruml_open_scripts() -> None:
    finder = r'''
function Resolve-StarUmlExecutable {
    $Command = Get-Command staruml -ErrorAction SilentlyContinue | Select-Object -First 1 -ExpandProperty Source
    if ($Command) {
        return $Command
    }

    $Candidates = @(
        "C:\Program Files\StarUML\StarUML.exe",
        "C:\Program Files (x86)\StarUML\StarUML.exe",
        "$env:LOCALAPPDATA\Programs\StarUML\StarUML.exe"
    )
    foreach ($Candidate in $Candidates) {
        if ($Candidate -and (Test-Path $Candidate)) {
            return $Candidate
        }
    }
    return $null
}
'''.strip()

    open_script = f'''$ErrorActionPreference = "Stop"

$Root = Split-Path -Parent $MyInvocation.MyCommand.Path
$Model = Join-Path $Root "sources\\smart-learning-system.mdj"

if (-not (Test-Path $Model)) {{
    throw "StarUML model not found: $Model. Run generate-all.ps1 first."
}}

{finder}

$Staruml = Resolve-StarUmlExecutable
if (-not $Staruml) {{
    throw "StarUML was not found. Install StarUML or add the 'staruml' command to PATH, then rerun this script."
}}

Start-Process -FilePath $Staruml -ArgumentList @($Model) -WindowStyle Normal
Write-Host "Opened StarUML model:" -ForegroundColor Green
Write-Host "  $Model"
'''
    (ROOT / "open-in-staruml.ps1").write_text(open_script, encoding="utf-8")

    install_open_script = f'''$ErrorActionPreference = "Stop"

$Root = Split-Path -Parent $MyInvocation.MyCommand.Path
$Source = Join-Path $Root "staruml-extension\\smart-learning-uml-generator"
$TargetRoot = Join-Path $env:APPDATA "StarUML\\extensions\\user"
$Target = Join-Path $TargetRoot "smart-learning-uml-generator"

if (-not (Test-Path $Source)) {{
    throw "Extension folder not found: $Source. Run generate-all.ps1 first."
}}

New-Item -ItemType Directory -Force -Path $TargetRoot | Out-Null
$BackupRoot = Join-Path $env:APPDATA "StarUML-disabled-extension-backups"
New-Item -ItemType Directory -Force -Path $BackupRoot | Out-Null
$LegacyBackupRoot = Join-Path (Split-Path -Parent $TargetRoot) "disabled-extension-backups"
if (Test-Path $LegacyBackupRoot) {{
    Get-ChildItem -LiteralPath $LegacyBackupRoot -Directory -Filter "smart-learning-uml-generator.bak-*" -ErrorAction SilentlyContinue | ForEach-Object {{
        $Destination = Join-Path $BackupRoot $_.Name
        if (Test-Path $Destination) {{
            Remove-Item -LiteralPath $Destination -Recurse -Force
        }}
        Move-Item -LiteralPath $_.FullName -Destination $Destination
    }}
}}
Get-ChildItem -LiteralPath $TargetRoot -Directory -Filter "smart-learning-uml-generator.bak-*" -ErrorAction SilentlyContinue | ForEach-Object {{
    $Destination = Join-Path $BackupRoot $_.Name
    if (Test-Path $Destination) {{
        Remove-Item -LiteralPath $Destination -Recurse -Force
    }}
    Move-Item -LiteralPath $_.FullName -Destination $Destination
}}
if (Test-Path $Target) {{
    $Backup = Join-Path $BackupRoot "smart-learning-uml-generator.bak-$(Get-Date -Format 'yyyyMMdd-HHmmss')"
    Move-Item -LiteralPath $Target -Destination $Backup
    Write-Host "Existing extension backed up to:"
    Write-Host "  $Backup"
}}
Copy-Item -LiteralPath $Source -Destination $Target -Recurse

{finder}

$Staruml = Resolve-StarUmlExecutable
if (-not $Staruml) {{
    Write-Host "StarUML extension installed to:" -ForegroundColor Green
    Write-Host "  $Target"
    throw "StarUML was not found. Install StarUML, then restart it and use Tools > 生成推荐 24 张 UML 图."
}}

$Running = Get-Process -Name "StarUML" -ErrorAction SilentlyContinue
if ($Running) {{
    Write-Host "StarUML is already running. If the new menu is not visible, close StarUML and run this script again." -ForegroundColor Yellow
}}

Start-Process -FilePath $Staruml -WindowStyle Normal
Write-Host "StarUML extension installed and StarUML opened:" -ForegroundColor Green
Write-Host "  $Target"
Write-Host "In StarUML, use Tools > 生成课程核心 9 张 UML 图 or Tools > 生成推荐 24 张 UML 图."
'''
    (ROOT / "install-and-open-staruml.ps1").write_text(install_open_script, encoding="utf-8")

    launcher_bat = r'''@echo off
setlocal
cd /d "%~dp0"
echo Installing StarUML extension and opening StarUML...
powershell -NoProfile -ExecutionPolicy Bypass -File "%~dp0install-and-open-staruml.ps1"
if errorlevel 1 (
  echo.
  echo Failed. Please install StarUML first, then run this file again.
  pause
  exit /b 1
)
echo.
echo StarUML has been opened. In StarUML, click:
echo   Tools ^> 生成课程核心 9 张 UML 图
echo   or
echo   Tools ^> 生成推荐 24 张 UML 图
echo If the menu is not visible, close StarUML and run this file again.
echo.
pause
'''
    (ROOT / "一键打开StarUML并安装扩展.bat").write_text(launcher_bat, encoding="utf-8")


def main() -> None:
    reset_dirs()
    diagrams = make_diagrams()
    render_all(diagrams)
    make_contact_sheet(diagrams)
    write_sources(diagrams)
    make_mdj(diagrams)
    write_staruml_extension()
    make_report(diagrams)
    write_staruml_open_scripts()
    write_readme(diagrams)
    print(f"Generated UML delivery under: {ROOT}")
    print(f"Diagrams: {len(diagrams)}")
    print("Report: report/*.docx")
    print("StarUML project: sources/smart-learning-system.mdj")


if __name__ == "__main__":
    main()
